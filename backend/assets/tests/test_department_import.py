from io import BytesIO
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase

from assets.models import AuditLog, Department
from assets.operation_limits import MAX_ASSET_MODEL_IMPORT_BYTES, MAX_ASSET_MODEL_IMPORT_ROWS


class DepartmentImportApiTests(APITestCase):
    HEADERS = ("code", "name", "parent_code")

    def setUp(self):
        self.user = User.objects.create_superuser(
            username="department-import-admin",
            first_name="Department",
            last_name="Importer",
            password="test-password-123",
        )
        self.client.force_authenticate(user=self.user)

    def workbook_upload(self, rows, *, headers=None, filename="department-import.xlsx"):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "部门导入"
        selected_headers = tuple(headers or self.HEADERS)
        sheet.append(list(selected_headers))
        for row in rows:
            sheet.append([row.get(header, "") for header in selected_headers])
        output = BytesIO()
        workbook.save(output)
        return SimpleUploadedFile(
            filename,
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def preview(self, rows, **kwargs):
        return self.client.post(
            "/api/v1/departments/import/preview/",
            {"file": self.workbook_upload(rows, **kwargs)},
            format="multipart",
        )

    def commit(self, rows, **kwargs):
        return self.client.post(
            "/api/v1/departments/import/",
            {"file": self.workbook_upload(rows, **kwargs)},
            format="multipart",
        )

    @staticmethod
    def row(code="001", name="平台部", parent_code=""):
        return {"code": code, "name": name, "parent_code": parent_code}

    def test_template_exposes_stable_machine_fields(self):
        response = self.client.get("/api/v1/departments/import/template/")

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=False)
        self.assertEqual(tuple(next(workbook.active.iter_rows(values_only=True))), self.HEADERS)
        self.assertIn("填写说明", workbook.sheetnames)
        workbook.close()

    def test_template_formats_code_and_parent_code_through_import_row_limit(self):
        response = self.client.get("/api/v1/departments/import/template/")

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook["部门导入"]
        self.assertEqual([sheet.cell(1, column).value for column in (1, 2, 3)], list(self.HEADERS))
        for row in (2, 3, MAX_ASSET_MODEL_IMPORT_ROWS + 1):
            with self.subTest(row=row):
                self.assertEqual(sheet.cell(row, 1).number_format, "@")
                self.assertEqual(sheet.cell(row, 3).number_format, "@")
        workbook.close()

    def test_official_template_roundtrip_preserves_textual_hierarchy_codes(self):
        response = self.client.get("/api/v1/departments/import/template/")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook["部门导入"]
        for row, values in enumerate((
            ("001", "根部门", ""),
            ("010", "子部门", "001"),
            ("0007", "末级部门", "010"),
        ), start=2):
            for column, value in enumerate(values, start=1):
                sheet.cell(row, column, value)
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        upload = SimpleUploadedFile("official-departments.xlsx", output.getvalue())
        preview = self.client.post("/api/v1/departments/import/preview/", {"file": upload}, format="multipart")
        self.assertEqual(preview.status_code, 200, preview.content)
        self.assertEqual(preview.json()["error"], 0)
        self.assertEqual(
            [(row["code"], row["parent_code"]) for row in preview.json()["rows"]],
            [("001", ""), ("010", "001"), ("0007", "010")],
        )
        self.assertFalse(Department.objects.exists())

        upload = SimpleUploadedFile("official-departments.xlsx", output.getvalue())
        commit = self.client.post("/api/v1/departments/import/", {"file": upload}, format="multipart")
        self.assertEqual(commit.status_code, 200, commit.content)
        root = Department.objects.get(code="001")
        child = Department.objects.get(code="010")
        leaf = Department.objects.get(code="0007")
        self.assertEqual(child.parent_id, root.pk)
        self.assertEqual(leaf.parent_id, child.pk)

    def test_preview_create_does_not_mutate_and_preserves_leading_zero_code(self):
        response = self.preview([self.row("001", "平台部"), self.row("010", "网络部"), self.row("0007", "机房部")])

        self.assertEqual(response.status_code, 200, response.content)
        payload = response.json()
        self.assertEqual((payload["create"], payload["update"], payload["unchanged"], payload["error"]), (3, 0, 0, 0))
        self.assertEqual([row["code"] for row in payload["rows"]], ["001", "010", "0007"])
        self.assertFalse(Department.objects.exists())

    def test_commit_preserves_leading_zero_codes(self):
        response = self.commit([self.row("001", "平台部"), self.row("010", "网络部"), self.row("0007", "机房部")])

        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(set(Department.objects.values_list("code", flat=True)), {"001", "010", "0007"})

    def test_preview_mixed_create_update_and_unchanged_uses_code_identity(self):
        existing = Department.objects.create(code="001", name="旧平台")
        unchanged = Department.objects.create(code="010", name="网络部")

        response = self.preview([
            self.row("001", "新平台"),
            self.row("010", "网络部"),
            self.row("0007", "机房部"),
        ])

        self.assertEqual(response.status_code, 200, response.content)
        payload = response.json()
        self.assertEqual((payload["create"], payload["update"], payload["unchanged"], payload["error"]), (1, 1, 1, 0))
        self.assertEqual(payload["rows"][0]["operation"], "update")
        self.assertEqual(payload["rows"][1]["operation"], "unchanged")
        self.assertEqual(payload["rows"][2]["operation"], "create")
        existing.refresh_from_db()
        unchanged.refresh_from_db()
        self.assertEqual(existing.name, "旧平台")
        self.assertEqual(unchanged.name, "网络部")

    def test_commit_mixed_rows_writes_one_summary_audit(self):
        Department.objects.create(code="001", name="旧平台")
        Department.objects.create(code="010", name="网络部")

        response = self.commit([
            self.row("001", "新平台"),
            self.row("010", "网络部"),
            self.row("0007", "机房部"),
        ])

        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["created"], 1)
        self.assertEqual(response.json()["updated"], 1)
        self.assertEqual(response.json()["unchanged"], 1)
        self.assertEqual(AuditLog.objects.filter(resource_type="department", action="import").count(), 1)
        audit = AuditLog.objects.get(resource_type="department", action="import")
        self.assertEqual(audit.payload["extra"]["source"], "department_import")
        self.assertEqual(audit.payload["extra"]["created_count"], 1)
        self.assertEqual(audit.payload["extra"]["updated_count"], 1)
        self.assertEqual(audit.payload["extra"]["unchanged_count"], 1)

    def test_child_before_parent_is_created_in_dependency_order(self):
        response = self.commit([
            self.row("C", "子部门", "P"),
            self.row("P", "父部门"),
        ])

        self.assertEqual(response.status_code, 200, response.content)
        parent = Department.objects.get(code="P")
        child = Department.objects.get(code="C")
        self.assertEqual(child.parent_id, parent.id)

    def test_parent_can_reference_existing_department(self):
        parent = Department.objects.create(code="P", name="父部门")

        response = self.commit([self.row("C", "子部门", "P")])

        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(Department.objects.get(code="C").parent_id, parent.id)

    def test_duplicate_code_marks_all_rows_as_errors(self):
        response = self.preview([self.row("001", "平台部"), self.row("001", "另一个部门")])

        self.assertEqual(response.status_code, 200, response.content)
        payload = response.json()
        self.assertEqual(payload["error"], 2)
        self.assertTrue(all(row["operation"] == "error" for row in payload["rows"]))
        self.assertTrue(all(row["errors"][0]["code"] == "duplicate_code" for row in payload["rows"]))

    def test_duplicate_headers_are_file_errors(self):
        response = self.preview(
            [self.row("001", "平台部")],
            headers=("code", "name", "name"),
        )

        self.assertEqual(response.status_code, 400, response.content)
        self.assertEqual(response.json()["code"], "file_error")

    def test_duplicate_names_follow_department_model_constraint(self):
        response = self.preview([
            self.row("001", "同名部门"),
            self.row("002", "同名部门"),
        ])

        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["error"], 2)
        self.assertTrue(all(row["errors"][0]["code"] == "duplicate_name" for row in response.json()["rows"]))

    def test_missing_parent_and_self_parent_are_row_errors(self):
        response = self.preview([
            self.row("C", "子部门", "MISSING"),
            self.row("SELF", "自循环", "SELF"),
        ])

        self.assertEqual(response.status_code, 200, response.content)
        rows = response.json()["rows"]
        self.assertEqual(rows[0]["errors"][0]["code"], "missing_parent")
        self.assertEqual(rows[1]["errors"][0]["code"], "self_parent")

    def test_two_and_three_node_cycles_are_rejected(self):
        for rows in (
            [self.row("A", "甲", "B"), self.row("B", "乙", "A")],
            [self.row("A", "甲", "B"), self.row("B", "乙", "C"), self.row("C", "丙", "A")],
        ):
            with self.subTest(rows=rows):
                response = self.preview(rows)
                self.assertEqual(response.status_code, 200, response.content)
                payload = response.json()
                self.assertEqual(payload["error"], len(rows))
                self.assertTrue(all(row["errors"][0]["code"] == "hierarchy_cycle" for row in payload["rows"]))
                self.assertFalse(Department.objects.exists())

    def test_existing_cycle_and_descendant_reparent_are_rejected(self):
        first = Department.objects.create(code="A", name="甲")
        second = Department.objects.create(code="B", name="乙", parent=first)
        first.parent = second
        first.save(update_fields=["parent", "updated_at"])

        response = self.preview([self.row("A", "甲")])
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["rows"][0]["errors"][0]["code"], "hierarchy_cycle")

        unrelated = self.preview([self.row("UNRELATED", "无关部门")])
        self.assertEqual(unrelated.status_code, 200, unrelated.content)
        self.assertEqual(unrelated.json()["rows"][0]["errors"][0]["code"], "hierarchy_cycle")

        Department.objects.all().update(parent=None)
        Department.objects.all().delete()
        first = Department.objects.create(code="A", name="甲")
        second = Department.objects.create(code="B", name="乙", parent=first)
        response = self.preview([self.row("A", "甲", "B")])
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["rows"][0]["errors"][0]["code"], "hierarchy_cycle")
        self.assertEqual(Department.objects.get(code="B").parent_id, first.id)
        self.assertEqual(second.parent_id, first.id)

    def test_blank_name_and_parent_preserve_existing_values(self):
        parent = Department.objects.create(code="P", name="父部门")
        child = Department.objects.create(code="C", name="子部门", parent=parent)

        response = self.commit([self.row("C", "", "")])

        self.assertEqual(response.status_code, 200, response.content)
        child.refresh_from_db()
        self.assertEqual(child.name, "子部门")
        self.assertEqual(child.parent_id, parent.id)
        self.assertEqual(response.json()["unchanged"], 1)

    def test_blank_parent_on_create_creates_root(self):
        response = self.commit([self.row("ROOT", "根部门", "")])

        self.assertEqual(response.status_code, 200, response.content)
        self.assertIsNone(Department.objects.get(code="ROOT").parent_id)

    def test_missing_required_headers_empty_workbook_and_invalid_workbook_are_file_errors(self):
        missing = self.preview([self.row("A", "甲")], headers=("code", "parent_code"))
        self.assertEqual(missing.status_code, 400)
        self.assertEqual(missing.json()["code"], "file_error")
        self.assertIn("name", missing.json()["detail"])

        workbook = Workbook()
        workbook.active.title = "部门导入"
        output = BytesIO()
        workbook.save(output)
        empty = self.client.post(
            "/api/v1/departments/import/preview/",
            {"file": SimpleUploadedFile("empty.xlsx", output.getvalue())},
            format="multipart",
        )
        self.assertEqual(empty.status_code, 400)
        self.assertEqual(empty.json()["code"], "file_error")

        invalid = self.client.post(
            "/api/v1/departments/import/preview/",
            {"file": SimpleUploadedFile("broken.xlsx", b"not an xlsx")},
            format="multipart",
        )
        self.assertEqual(invalid.status_code, 400)
        self.assertEqual(invalid.json()["code"], "file_error")

    def test_oversized_and_too_many_rows_are_rejected(self):
        oversized = self.client.post(
            "/api/v1/departments/import/preview/",
            {"file": SimpleUploadedFile("department-import.xlsx", b"x" * (MAX_ASSET_MODEL_IMPORT_BYTES + 1))},
            format="multipart",
        )
        self.assertEqual(oversized.status_code, 400)
        self.assertIn(str(MAX_ASSET_MODEL_IMPORT_BYTES // (1024 * 1024)), oversized.json()["detail"])

        rows = [self.row(f"D-{index}", f"部门 {index}") for index in range(MAX_ASSET_MODEL_IMPORT_ROWS + 1)]
        response = self.preview(rows)
        self.assertEqual(response.status_code, 400, response.content)
        self.assertIn(str(MAX_ASSET_MODEL_IMPORT_ROWS), response.json()["detail"])

    def test_name_and_field_length_constraints_are_reported_without_mutation(self):
        response = self.preview([
            self.row("A" * 51, "正常名称"),
            self.row("B", "名称" * 51),
        ])

        self.assertEqual(response.status_code, 200, response.content)
        rows = response.json()["rows"]
        self.assertEqual(rows[0]["errors"][0]["field"], "code")
        self.assertEqual(rows[1]["errors"][0]["field"], "name")
        self.assertFalse(Department.objects.exists())

    def test_invalid_row_prevents_all_mutations(self):
        response = self.commit([
            self.row("GOOD", "应回滚"),
            self.row("BAD", "", "MISSING"),
        ])

        self.assertEqual(response.status_code, 400, response.content)
        self.assertFalse(Department.objects.exists())
        self.assertFalse(AuditLog.objects.filter(resource_type="department", action="import").exists())

    def test_audit_failure_rolls_back_all_department_mutations(self):
        with patch("assets.department_imports.write_audit_log", side_effect=RuntimeError("audit unavailable")):
            with self.assertRaises(RuntimeError):
                self.commit([self.row("ROLLBACK", "不应保存")])

        self.assertFalse(Department.objects.filter(code="ROLLBACK").exists())
        self.assertFalse(AuditLog.objects.filter(resource_type="department", action="import").exists())

    def test_integrity_race_returns_latest_preview_and_rolls_back(self):
        with patch("assets.department_imports.Department.save", side_effect=IntegrityError("race")):
            response = self.commit([self.row("RACE", "并发失败")])

        self.assertEqual(response.status_code, 409, response.content)
        self.assertTrue(response.json()["preview"]["error"] == 0)
        self.assertFalse(Department.objects.filter(code="RACE").exists())

    def test_view_only_user_cannot_use_department_import(self):
        reader = User.objects.create_user(username="department-import-reader", password="test-password-123")
        self.client.force_authenticate(user=reader)

        response = self.commit([self.row("FORBIDDEN", "不应保存")])

        self.assertEqual(response.status_code, 403)
        self.assertFalse(Department.objects.filter(code="FORBIDDEN").exists())

    def test_view_only_user_cannot_preview_or_download_department_import(self):
        reader = User.objects.create_user(username="department-import-reader-2", password="test-password-123")
        self.client.force_authenticate(user=reader)

        preview = self.preview([self.row("FORBIDDEN", "不应读取")])
        template = self.client.get("/api/v1/departments/import/template/")

        self.assertEqual(preview.status_code, 403)
        self.assertEqual(template.status_code, 403)
