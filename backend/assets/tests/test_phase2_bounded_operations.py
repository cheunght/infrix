import json
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from openpyxl import Workbook
from rest_framework.exceptions import ValidationError
from rest_framework.test import APITestCase

from assets.asset_model_imports import _validate_file_size
from assets.models import (
    Asset,
    AssetModel,
    AuditLog,
    DataCenter,
    DeviceType,
    InventoryItem,
    InventoryTask,
    Manufacturer,
)
from assets.operation_limits import MAX_ASSET_MODEL_IMPORT_BYTES


class InventoryTaskBoundedCreationTests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            username="inventory-admin",
            email="inventory-admin@example.com",
            password="test-password-123",
        )
        self.client.force_authenticate(user=self.user)
        self.manufacturer = Manufacturer.objects.create(name="Bounded Manufacturer", code="bounded")
        self.device_type = DeviceType.objects.create(name="Bounded Device")
        self.data_center = DataCenter.objects.create(name="Bounded DC")
        self.other_data_center = DataCenter.objects.create(name="Other DC")

    def create_assets(self, count, *, data_center=None, prefix="asset"):
        data_center = data_center or self.data_center
        return [
            Asset.objects.create(
                asset_no=f"{prefix}-{index}",
                name=f"Asset {index}",
                asset_data_center=data_center,
                standalone_manufacturer=self.manufacturer,
                standalone_device_type=self.device_type,
            )
            for index in range(count)
        ]

    def task_payload(self, *, scope="all_assets", data_center=None):
        return {
            "name": "Bounded inventory task",
            "scope": scope,
            "data_center": data_center,
            "server_room": None,
            "start_at": "2026-09-23T09:00:00+08:00",
            "end_at": "2026-09-23T18:00:00+08:00",
        }

    @patch("assets.views.MAX_INVENTORY_TASK_ASSETS", 2)
    def test_all_assets_below_limit_are_created(self):
        self.create_assets(1)

        response = self.client.post(
            "/api/v1/inventory-tasks/",
            self.task_payload(),
            format="json",
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(InventoryTask.objects.count(), 1)
        self.assertEqual(InventoryItem.objects.count(), 1)
        self.assertEqual(response.json()["summary"]["total"], 1)

    @patch("assets.views.MAX_INVENTORY_TASK_ASSETS", 2)
    def test_all_assets_exactly_at_limit_are_created(self):
        self.create_assets(2)

        response = self.client.post(
            "/api/v1/inventory-tasks/",
            self.task_payload(),
            format="json",
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(InventoryTask.objects.count(), 1)
        self.assertEqual(InventoryItem.objects.count(), 2)

    @patch("assets.views.MAX_INVENTORY_TASK_ASSETS", 2)
    def test_all_assets_above_limit_are_rejected_without_partial_creation(self):
        self.create_assets(3)

        response = self.client.post(
            "/api/v1/inventory-tasks/",
            self.task_payload(),
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("3", json.dumps(response.json(), ensure_ascii=False))
        self.assertIn("2", json.dumps(response.json(), ensure_ascii=False))
        self.assertEqual(InventoryTask.objects.count(), 0)
        self.assertEqual(InventoryItem.objects.count(), 0)

    @patch("assets.views.MAX_INVENTORY_TASK_ASSETS", 2)
    def test_filtered_scope_count_matches_created_items_and_preview_limit(self):
        self.create_assets(2, data_center=self.data_center, prefix="primary")
        self.create_assets(1, data_center=self.other_data_center, prefix="other")

        preview = self.client.get(
            "/api/v1/inventory-tasks/scope-preview/",
            {"scope": "data_center", "data_center": self.data_center.pk},
        )
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.json()["total"], 2)
        self.assertEqual(preview.json()["max_assets"], 2)

        response = self.client.post(
            "/api/v1/inventory-tasks/",
            self.task_payload(scope="data_center", data_center=self.data_center.pk),
            format="json",
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(InventoryItem.objects.count(), 2)
        self.assertSetEqual(
            set(InventoryItem.objects.values_list("asset__asset_no", flat=True)),
            {"primary-0", "primary-1"},
        )

    @patch("assets.views.MAX_INVENTORY_TASK_ASSETS", 2)
    @patch("assets.views.get_inventory_scope_assets")
    def test_over_limit_rejects_before_materialization(self, get_assets):
        queryset = MagicMock()
        queryset.count.return_value = 3
        get_assets.return_value = queryset

        response = self.client.post(
            "/api/v1/inventory-tasks/",
            self.task_payload(),
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        queryset.__getitem__.assert_not_called()
        self.assertEqual(InventoryTask.objects.count(), 0)

    @patch("assets.views.MAX_INVENTORY_TASK_ASSETS", 2)
    @patch("assets.views.get_inventory_scope_assets")
    def test_materialization_guard_rejects_race_growth_before_task_save(self, get_assets):
        queryset = MagicMock()
        queryset.count.return_value = 2
        queryset.__getitem__.return_value = [object(), object(), object()]
        get_assets.return_value = queryset

        response = self.client.post(
            "/api/v1/inventory-tasks/",
            self.task_payload(),
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(InventoryTask.objects.count(), 0)
        self.assertEqual(InventoryItem.objects.count(), 0)


class AssetModelImportBoundedProcessingTests(APITestCase):
    headers = [
        "name",
        "model_number",
        "manufacturer",
        "device_type",
        "fieldset",
        "default_warranty_months",
        "expected_life_months",
        "notes",
        "is_active",
    ]

    def setUp(self):
        self.user = User.objects.create_superuser(
            username="model-admin",
            email="model-admin@example.com",
            password="test-password-123",
        )
        self.client.force_authenticate(user=self.user)
        self.manufacturer = Manufacturer.objects.create(name="Import Manufacturer", code="import")
        self.device_type = DeviceType.objects.create(name="Import Device")

    def workbook_bytes(self, count):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "资产型号导入"
        sheet.append(self.headers)
        for index in range(count):
            sheet.append([
                f"Imported Model {index}",
                f"IM-{index}",
                self.manufacturer.name,
                self.device_type.name,
                "",
                12,
                60,
                "bounded import test",
                "true",
            ])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()

    @staticmethod
    def upload(data):
        return SimpleUploadedFile(
            "asset-models.xlsx",
            data,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def post_preview(self, data):
        return self.client.post(
            "/api/v1/asset-models/import/preview/",
            {"file": self.upload(data)},
            format="multipart",
        )

    def post_commit(self, data):
        return self.client.post(
            "/api/v1/asset-models/import/",
            {"file": self.upload(data)},
            format="multipart",
        )

    def test_valid_small_xlsx_is_previewed_and_committed(self):
        data = self.workbook_bytes(1)

        preview = self.post_preview(data)
        commit = self.post_commit(data)

        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.json()["total"], 1)
        self.assertEqual(preview.json()["valid"], 1)
        self.assertEqual(commit.status_code, 200)
        self.assertEqual(commit.json()["created"], 1)
        self.assertEqual(AssetModel.objects.count(), 1)
        self.assertEqual(AuditLog.objects.filter(resource_type="asset_model").count(), 1)

    @patch("assets.asset_model_imports.MAX_ASSET_MODEL_IMPORT_ROWS", 2)
    def test_rows_exactly_at_limit_are_accepted_by_preview_and_commit(self):
        data = self.workbook_bytes(2)

        preview = self.post_preview(data)
        commit = self.post_commit(data)

        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.json()["total"], 2)
        self.assertEqual(commit.status_code, 200)
        self.assertEqual(commit.json()["created"], 2)
        self.assertEqual(AssetModel.objects.count(), 2)

    @patch("assets.asset_model_imports.MAX_ASSET_MODEL_IMPORT_ROWS", 2)
    def test_rows_above_limit_are_rejected_by_preview_and_commit_without_mutation(self):
        data = self.workbook_bytes(3)

        preview = self.post_preview(data)
        commit = self.post_commit(data)

        self.assertEqual(preview.status_code, 400)
        self.assertEqual(commit.status_code, 400)
        self.assertEqual(AssetModel.objects.count(), 0)
        self.assertEqual(AuditLog.objects.filter(resource_type="asset_model").count(), 0)

    def test_file_size_exactly_at_limit_is_accepted(self):
        data = self.workbook_bytes(1)
        with patch("assets.asset_model_imports.MAX_ASSET_MODEL_IMPORT_BYTES", len(data)):
            response = self.post_preview(data)

        self.assertEqual(response.status_code, 200)

    def test_file_size_above_limit_is_rejected_before_workbook_loading_and_mutation(self):
        data = self.workbook_bytes(1)
        with patch("assets.asset_model_imports.MAX_ASSET_MODEL_IMPORT_BYTES", len(data) - 1), patch(
            "assets.asset_model_imports.load_workbook"
        ) as load_workbook:
            preview = self.post_preview(data)
            commit = self.post_commit(data)

        self.assertEqual(preview.status_code, 400)
        self.assertEqual(commit.status_code, 400)
        load_workbook.assert_not_called()
        self.assertEqual(AssetModel.objects.count(), 0)
        self.assertEqual(AuditLog.objects.filter(resource_type="asset_model").count(), 0)


class AssetModelImportFileSizeHelperTests(SimpleTestCase):
    def test_declared_size_boundary_is_inclusive(self):
        _validate_file_size(SimpleNamespace(name="asset-models.xlsx", size=MAX_ASSET_MODEL_IMPORT_BYTES))

        with self.assertRaises(ValidationError):
            _validate_file_size(SimpleNamespace(name="asset-models.xlsx", size=MAX_ASSET_MODEL_IMPORT_BYTES + 1))
