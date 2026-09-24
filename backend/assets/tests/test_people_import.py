from io import BytesIO
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError
from openpyxl import Workbook
from rest_framework.test import APITestCase

from assets.models import AuditLog, Department, Person
from assets.operation_limits import MAX_ASSET_MODEL_IMPORT_BYTES, MAX_ASSET_MODEL_IMPORT_ROWS


class PeopleImportApiTests(APITestCase):
    HEADERS = (
        "employee_no",
        "name",
        "department_code",
        "email",
        "organization",
        "contact",
        "is_active",
    )

    def setUp(self):
        self.user = User.objects.create_superuser(
            username="people-import-admin",
            first_name="People",
            last_name="Importer",
            password="test-password-123",
        )
        self.client.force_authenticate(user=self.user)
        self.department = Department.objects.create(name="平台部", code="PLATFORM")

    def workbook_upload(self, rows, *, headers=None, filename="people-import.xlsx"):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "人员导入"
        sheet.append(list(headers or self.HEADERS))
        for row in rows:
            selected_headers = headers or self.HEADERS
            sheet.append([row.get(header, "") for header in selected_headers])
        output = BytesIO()
        workbook.save(output)
        return SimpleUploadedFile(
            filename,
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def preview(self, rows, **kwargs):
        return self.client.post(
            "/api/v1/people/import/preview/",
            {"file": self.workbook_upload(rows, **kwargs)},
            format="multipart",
        )

    def commit(self, rows, **kwargs):
        return self.client.post(
            "/api/v1/people/import/",
            {"file": self.workbook_upload(rows, **kwargs)},
            format="multipart",
        )

    def row(self, employee_no="00123", **overrides):
        value = {
            "employee_no": employee_no,
            "name": "张三",
            "department_code": self.department.code,
            "email": "zhangsan@example.com",
            "organization": "信息技术部",
            "contact": "010-12345678",
            "is_active": "true",
        }
        value.update(overrides)
        return value

    def test_template_exposes_machine_fields_without_account_or_security_fields(self):
        response = self.client.get("/api/v1/people/import/template/")

        self.assertEqual(response.status_code, 200)
        # The response is an XLSX stream; parse it through openpyxl in read-only mode.
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=False)
        self.assertEqual(tuple(next(workbook.active.iter_rows(values_only=True))), self.HEADERS)
        self.assertNotIn("password", self.HEADERS)
        self.assertNotIn("account", self.HEADERS)
        self.assertNotIn("ldap_secret", self.HEADERS)
        workbook.close()

    def test_preview_create_preserves_text_employee_number_and_does_not_mutate(self):
        response = self.preview([self.row()])

        self.assertEqual(response.status_code, 200, response.content)
        payload = response.json()
        self.assertEqual(payload["create"], 1)
        self.assertEqual(payload["update"], 0)
        self.assertEqual(payload["unchanged"], 0)
        self.assertEqual(payload["error"], 0)
        self.assertEqual(payload["rows"][0]["operation"], "create")
        self.assertEqual(payload["rows"][0]["employee_no"], "00123")
        self.assertFalse(Person.objects.filter(employee_no="00123").exists())

    def test_preview_update_and_blank_optional_cells_preserve_existing_values(self):
        person = Person.objects.create(
            name="旧姓名",
            employee_no="00123",
            department=self.department,
            email="old@example.com",
            organization="旧单位",
            contact="旧联系方式",
            is_active=True,
        )

        response = self.preview([self.row(name="新姓名", email="", organization="", contact="", is_active="")])

        self.assertEqual(response.status_code, 200, response.content)
        payload = response.json()
        self.assertEqual(payload["update"], 1)
        self.assertEqual(payload["rows"][0]["operation"], "update")
        self.assertEqual(payload["rows"][0]["changes"], ["name"])
        person.refresh_from_db()
        self.assertEqual(person.name, "旧姓名")
        self.assertEqual(person.email, "old@example.com")
        self.assertEqual(person.organization, "旧单位")
        self.assertEqual(person.contact, "旧联系方式")

    def test_preview_unchanged_row_is_skipped(self):
        Person.objects.create(
            name="张三",
            employee_no="00123",
            department=self.department,
            email="zhangsan@example.com",
            organization="信息技术部",
            contact="010-12345678",
            is_active=True,
        )

        response = self.preview([self.row()])

        self.assertEqual(response.status_code, 200, response.content)
        payload = response.json()
        self.assertEqual(payload["unchanged"], 1)
        self.assertEqual(payload["rows"][0]["operation"], "unchanged")
        self.assertEqual(payload["rows"][0]["changes"], [])

    def test_unknown_department_is_row_error_without_auto_creation(self):
        response = self.preview([self.row(department_code="MISSING")])

        self.assertEqual(response.status_code, 200, response.content)
        payload = response.json()
        self.assertEqual(payload["error"], 1)
        self.assertEqual(payload["rows"][0]["operation"], "error")
        self.assertIn("部门", payload["rows"][0]["errors"][0]["message"])
        self.assertFalse(Department.objects.filter(code="MISSING").exists())
        self.assertFalse(Person.objects.filter(employee_no="00123").exists())

    def test_duplicate_employee_numbers_in_one_file_mark_all_rows_as_errors(self):
        response = self.preview([self.row(), self.row(name="另一个人")])

        self.assertEqual(response.status_code, 200, response.content)
        payload = response.json()
        self.assertEqual(payload["error"], 2)
        self.assertTrue(all(row["operation"] == "error" for row in payload["rows"]))
        self.assertTrue(all("重复" in row["errors"][0]["message"] for row in payload["rows"]))
        self.assertFalse(Person.objects.filter(employee_no__in=("00123",)).exists())

    def test_missing_employee_number_header_is_rejected(self):
        headers = tuple(header for header in self.HEADERS if header != "employee_no")
        response = self.preview([self.row()], headers=headers)

        self.assertEqual(response.status_code, 400)
        self.assertIn("employee_no", response.json()["detail"])

    def test_missing_employee_number_cell_is_a_row_error(self):
        response = self.preview([self.row(employee_no="")])

        self.assertEqual(response.status_code, 200, response.content)
        payload = response.json()
        self.assertEqual(payload["error"], 1)
        self.assertEqual(payload["rows"][0]["operation"], "error")
        self.assertIn("员工编号", payload["rows"][0]["errors"][0]["message"])

    def test_oversized_file_is_rejected_before_parsing(self):
        upload = SimpleUploadedFile(
            "people-import.xlsx",
            b"x" * (MAX_ASSET_MODEL_IMPORT_BYTES + 1),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = self.client.post("/api/v1/people/import/preview/", {"file": upload}, format="multipart")

        self.assertEqual(response.status_code, 400)
        self.assertIn("10 MB", response.json()["detail"])

    def test_business_row_limit_is_enforced(self):
        rows = [self.row(employee_no=f"P-{index:05d}", name=f"人员 {index}") for index in range(MAX_ASSET_MODEL_IMPORT_ROWS + 1)]

        response = self.preview(rows)

        self.assertEqual(response.status_code, 400, response.content)
        self.assertIn(str(MAX_ASSET_MODEL_IMPORT_ROWS), response.json()["detail"])
        self.assertFalse(Person.objects.filter(employee_no__startswith="P-").exists())

    def test_commit_creates_updates_skips_unchanged_and_writes_one_import_audit(self):
        Person.objects.create(
            name="旧姓名",
            employee_no="P-UPDATE",
            department=self.department,
            email="old@example.com",
            organization="旧单位",
            contact="旧联系方式",
            is_active=True,
        )
        Person.objects.create(
            name="不变人员",
            employee_no="P-SAME",
            department=self.department,
            email="same@example.com",
            organization="信息技术部",
            contact="010-12345678",
            is_active=True,
        )
        rows = [
            self.row(employee_no="P-CREATE", name="新建人员"),
            self.row(employee_no="P-UPDATE", name="更新姓名", email="new@example.com"),
            self.row(
                employee_no="P-SAME",
                name="不变人员",
                email="same@example.com",
            ),
        ]

        response = self.commit(rows)

        self.assertEqual(response.status_code, 200, response.content)
        payload = response.json()
        self.assertEqual(payload["created"], 1)
        self.assertEqual(payload["updated"], 1)
        self.assertEqual(payload["unchanged"], 1)
        self.assertEqual(payload["total"], 3)
        self.assertEqual(Person.objects.filter(employee_no="P-CREATE").count(), 1)
        self.assertEqual(Person.objects.get(employee_no="P-UPDATE").name, "更新姓名")
        self.assertEqual(Person.objects.get(employee_no="P-UPDATE").email, "new@example.com")
        self.assertEqual(Person.objects.get(employee_no="P-SAME").name, "不变人员")
        self.assertEqual(AuditLog.objects.filter(resource_type="person", action="import").count(), 1)
        audit = AuditLog.objects.get(resource_type="person", action="import")
        self.assertEqual(audit.payload["extra"]["source"], "people_import")
        self.assertEqual(audit.payload["extra"]["created_count"], 1)
        self.assertEqual(audit.payload["extra"]["updated_count"], 1)
        self.assertEqual(User.objects.count(), 1)

    def test_commit_re_resolves_identity_after_preview(self):
        upload_rows = [self.row(employee_no="P-DRIFT", name="导入姓名")]
        preview = self.preview(upload_rows)
        self.assertEqual(preview.status_code, 200)
        Person.objects.create(name="并发创建", employee_no="P-DRIFT")

        response = self.commit(upload_rows)

        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["updated"], 1)
        self.assertEqual(Person.objects.get(employee_no="P-DRIFT").name, "导入姓名")
        self.assertEqual(Person.objects.filter(employee_no="P-DRIFT").count(), 1)

    def test_audit_failure_rolls_back_all_people_mutations(self):
        rows = [self.row(employee_no="P-ROLLBACK", name="不应保存")]

        with patch("assets.people_imports.write_audit_log", side_effect=RuntimeError("audit unavailable")):
            with self.assertRaises(RuntimeError):
                self.commit(rows)

        self.assertFalse(Person.objects.filter(employee_no="P-ROLLBACK").exists())
        self.assertFalse(AuditLog.objects.filter(resource_type="person", action="import").exists())

    def test_integrity_race_returns_latest_preview_after_rollback(self):
        rows = [self.row(employee_no="P-RACE", name="并发失败")]

        with patch("assets.people_imports.PersonSerializer.save", side_effect=IntegrityError("race")):
            response = self.commit(rows)

        self.assertEqual(response.status_code, 409, response.content)
        self.assertEqual(response.json()["preview"]["create"], 1)
        self.assertFalse(Person.objects.filter(employee_no="P-RACE").exists())

    def test_view_only_user_cannot_commit_import(self):
        reader = User.objects.create_user(username="people-import-reader", password="test-password-123")
        self.client.force_authenticate(user=reader)

        response = self.commit([self.row(employee_no="P-FORBIDDEN")])

        self.assertEqual(response.status_code, 403)
        self.assertFalse(Person.objects.filter(employee_no="P-FORBIDDEN").exists())
