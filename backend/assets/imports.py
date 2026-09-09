"""Server-side asset import parsing, validation and commit orchestration.

The existing asset write serializer remains the source of truth for the actual
asset graph.  This module only normalizes workbook rows, resolves the
human-readable dictionary values used by the standard template, and coordinates
the preview/confirm transaction around that serializer.
"""

from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
import csv
import re

from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.validators import DecimalValidator
from django.db import IntegrityError, transaction
from django.db.models import Q
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import ValidationError as DRFValidationError

from .audit import asset_audit_snapshot, write_audit_log
from .custom_fields import validate_custom_field_value
from .enum_contracts import ASSET_STATUS_LABELS
from .lifecycle import validate_asset_status_transition
from .models import Asset, CustomField, DataCenter, Department, DeviceType, Manufacturer, Person, Rack, ServerRoom, Tag
from .serializers import AssetWriteSerializer
from .system_settings import get_system_settings


IMPORT_MAX_FILE_SIZE = 10 * 1024 * 1024
IMPORT_MAX_ROWS = 10000
IMPORT_TEMPLATE_FORMAT = "infrix-asset-import"

IMPORT_COLUMNS = (
    ("asset_no", "资产编号", True, "资产唯一编号；文件内不能重复，系统中已存在的编号不能再次导入。"),
    ("name", "资产名称", True, "资产显示名称。"),
    ("device_type", "设备类型", True, "填写启用中的设备类型名称。"),
    ("manufacturer", "厂商", False, "填写启用中的厂商名称或编码。"),
    ("model", "型号", False, "设备型号。"),
    ("manufacturer_model", "厂商/型号", False, "厂商和型号组合显示文本；通常填写型号即可。"),
    ("serial_number", "序列号", False, "留空表示没有序列号；序列号不能与其他资产重复。"),
    ("purpose", "用途", False, "资产用途。"),
    ("status", "状态", False, "可填 in_stock、in_use、idle、retired，或对应显示值在库、在用、闲置、已报废；维修中由故障流程维护。"),
    ("assigned_person_employee_no", "使用人员工编号", False, "优先按员工编号匹配启用人员；不能自动创建人员。"),
    ("assigned_person_name", "使用人姓名", False, "没有员工编号时必填，并结合使用人所属部门匹配；没有部门时姓名必须唯一。"),
    ("assigned_person_department", "使用人部门", False, "用于匹配人员资料，不代表资产归属部门；填写时必须与人员所属部门一致。"),
    ("assignment_reason", "使用人变更原因", False, "导入指定使用人的操作原因。"),
    ("notes", "备注", False, "资产备注。"),
    ("asset_data_center", "未上架所属数据中心", False, "未上架资产的所属数据中心；填写启用中的数据中心名称。"),
    ("data_center", "机柜所属数据中心", False, "上架时与机房、机柜、起止 U 一起填写；使用真实层级名称。"),
    ("server_room", "机房", False, "上架时填写所选数据中心下的机房名称。"),
    ("rack_code", "机柜编号", False, "上架时填写所选机房下的机柜编号。"),
    ("rack_start_u", "起始 U", False, "上架时填写大于等于 1 的整数。"),
    ("rack_end_u", "结束 U", False, "上架时填写不小于起始 U 且不超过机柜容量的整数。"),
    ("business_ip", "业务 IP", False, "合法 IPv4 或 IPv6 地址。"),
    ("management_ip", "管理 IP", False, "合法 IPv4 或 IPv6 地址。"),
    ("oob_ip", "带外 IP", False, "合法 IPv4 或 IPv6 地址。"),
    ("purchase_date", "采购日期", False, "Excel 日期单元格或 YYYY-MM-DD。填写采购信息时必填。"),
    ("supplier", "供应商", False, "供应商名称。"),
    ("purchase_order_no", "采购单号", False, "采购单号。"),
    ("purchase_amount", "采购金额", False, "非负金额；服务端使用 Decimal 校验。"),
    ("depreciation_start_date", "折旧起算日", False, "填写 YYYY-MM-DD；配置折旧时必须显式填写，不能自动使用采购日期。"),
    ("depreciation_years", "折旧年限", False, "填写大于等于 1 的整数，单位为年。"),
    ("residual_rate", "残值率", False, "填写 0~100 的百分数，例如 5 或 5% 表示 5%；不要填写 0.05。"),
    ("maintenance_provider", "维保厂商", False, "维保厂商名称。"),
    ("maintenance_contract_no", "维保合同号", False, "维保合同号。"),
    ("maintenance_start_date", "维保开始日", False, "Excel 日期单元格或 YYYY-MM-DD。"),
    ("maintenance_expiry_date", "维保到期日", False, "Excel 日期单元格或 YYYY-MM-DD，不能早于开始日。"),
    ("tags", "标签", False, "多个标签用英文分号、中文分号或逗号分隔；不会自动创建标签。"),
)

IMPORT_FIELD_LABELS = {key: label for key, label, _required, _description in IMPORT_COLUMNS}
IMPORT_FIELD_LABELS.update({"configuration": "机柜位置", "custom_values": "自定义字段"})
IMPORT_BASE_HEADERS = {key for key, _label, _required, _description in IMPORT_COLUMNS}
IMPORT_REQUIRED_HEADERS = {key for key, _label, required, _description in IMPORT_COLUMNS if required}
IMPORT_CUSTOM_HEADER_RE = re.compile(r"custom__[a-z][a-z0-9_]*$")
ASSET_STATUS_IMPORT_ALIASES = {
    **{value: value for value in ASSET_STATUS_LABELS},
    **{label: value for value, label in ASSET_STATUS_LABELS.items()},
}
class ImportFileError(ValueError):
    """A structural or parsing error that should be shown inside the dialog."""


class ImportValidationError(Exception):
    """Raised when preview or confirm validation finds invalid rows."""

    def __init__(self, preview, *, concurrent=False):
        self.preview = preview
        self.concurrent = concurrent
        super().__init__("导入文件校验未通过" if not concurrent else "确认导入前数据已发生变化")


@dataclass(frozen=True)
class ParsedImport:
    filename: str
    headers: tuple[str, ...]
    rows: tuple[tuple[int, dict[str, str]], ...]


def _error_items(detail):
    """Normalize nested Django/DRF errors into one predictable row shape."""

    result = []

    def visit(value, field="row"):
        if isinstance(value, dict):
            for key, nested in value.items():
                visit(nested, str(key))
            return
        if isinstance(value, (list, tuple)):
            for nested in value:
                visit(nested, field)
            return
        result.append({
            "field": field,
            "label": IMPORT_FIELD_LABELS.get(field, field),
            "message": str(value or "数据格式不正确"),
        })

    visit(detail)
    return result or [{"field": "row", "label": "整行", "message": "数据格式不正确"}]


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, Decimal)):
        return str(value)
    if isinstance(value, float):
        try:
            decimal_value = Decimal(str(value))
            return format(decimal_value, "f").rstrip("0").rstrip(".") or "0"
        except InvalidOperation:
            return str(value)
    return str(value).strip()


def _is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _normalize_headers(values):
    headers = [_cell_text(value).lstrip("\ufeff").strip() for value in values]
    while headers and not headers[-1]:
        headers.pop()
    if not headers or any(not header for header in headers):
        raise ImportFileError("导入文件第一行必须包含完整的字段名")
    if len(headers) != len(set(headers)):
        raise ImportFileError("导入文件包含重复字段名，请保留每个字段一列")
    unknown = [
        header for header in headers
        if header not in IMPORT_BASE_HEADERS and not IMPORT_CUSTOM_HEADER_RE.fullmatch(header)
    ]
    if unknown:
        raise ImportFileError(f"导入文件包含未知字段：{'、'.join(unknown)}")
    missing = sorted(IMPORT_REQUIRED_HEADERS.difference(headers))
    if missing:
        labels = "、".join(IMPORT_FIELD_LABELS.get(key, key) for key in missing)
        raise ImportFileError(f"导入文件缺少必填列：{labels}")
    return tuple(headers)


def _rows_from_csv(content):
    try:
        reader = csv.reader(StringIO(content))
        header_values = next(reader, None)
    except csv.Error as exc:
        raise ImportFileError("CSV 文件结构无法解析") from exc
    headers = _normalize_headers(header_values or [])
    rows = []
    for line, values in enumerate(reader, start=2):
        if not any(not _is_blank(value) for value in values):
            continue
        if len(values) > len(headers) and any(not _is_blank(value) for value in values[len(headers):]):
            raise ImportFileError(f"第 {line} 行包含多余列")
        values = list(values[:len(headers)]) + [""] * max(0, len(headers) - len(values))
        rows.append((line, {header: _cell_text(value) for header, value in zip(headers, values)}))
        if len(rows) > IMPORT_MAX_ROWS:
            raise ImportFileError(f"单次导入最多支持 {IMPORT_MAX_ROWS} 条资产")
    return headers, rows


def _rows_from_xlsx(content):
    try:
        workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:  # openpyxl exposes several parser-specific exceptions
        raise ImportFileError("Excel 文件无法解析，请确认文件为有效的 .xlsx 工作簿") from exc
    try:
        if not workbook.worksheets:
            raise ImportFileError("Excel 文件没有可读取的工作表")
        sheet = workbook.worksheets[0]
        if sheet.max_row and sheet.max_row > IMPORT_MAX_ROWS + 1:
            raise ImportFileError(f"单次导入最多支持 {IMPORT_MAX_ROWS} 条资产")
        iterator = sheet.iter_rows()
        header_cells = next(iterator, None)
        if not header_cells:
            raise ImportFileError("Excel 文件第一行必须包含字段名")
        headers = _normalize_headers([cell.value for cell in header_cells])
        rows = []
        for line, cells in enumerate(iterator, start=2):
            if not any(not _is_blank(cell.value) for cell in cells):
                continue
            for cell in cells:
                if cell.data_type == "f":
                    raise ImportFileError(f"第 {line} 行包含公式，请将公式结果复制为值后再导入")
            if len(cells) > len(headers) and any(not _is_blank(cell.value) for cell in cells[len(headers):]):
                raise ImportFileError(f"第 {line} 行包含多余列")
            values = [cell.value for cell in cells[:len(headers)]]
            values += [""] * max(0, len(headers) - len(values))
            rows.append((line, {header: _cell_text(value) for header, value in zip(headers, values)}))
            if len(rows) > IMPORT_MAX_ROWS:
                raise ImportFileError(f"单次导入最多支持 {IMPORT_MAX_ROWS} 条资产")
        return headers, rows
    finally:
        workbook.close()


def _parse_upload(upload):
    if not upload:
        raise ImportFileError("请上传 .xlsx 或 .csv 文件")
    if upload.size > IMPORT_MAX_FILE_SIZE:
        raise ImportFileError("导入文件不能超过 10 MB")
    filename = str(getattr(upload, "name", "") or "asset-import.xlsx")
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix not in {"xlsx", "csv"}:
        raise ImportFileError("仅支持 .xlsx 或 .csv 文件")
    content = upload.read()
    if suffix == "csv":
        try:
            headers, rows = _rows_from_csv(content.decode("utf-8-sig"))
        except UnicodeDecodeError as exc:
            raise ImportFileError("CSV 文件必须使用 UTF-8 编码") from exc
    else:
        headers, rows = _rows_from_xlsx(content)
    if not rows:
        raise ImportFileError("导入文件中没有资产数据")
    return ParsedImport(filename=filename, headers=tuple(headers), rows=tuple(rows))


def _named_active(queryset, value, field, label, *, allow_code=False):
    value = (value or "").strip()
    if not value:
        return None
    if allow_code:
        matches = list(queryset.filter(Q(name__iexact=value) | Q(code__iexact=value)))
    else:
        matches = list(queryset.filter(name__iexact=value))
    if not matches:
        raise DjangoValidationError({field: f"未找到启用的{label}“{value}”"})
    if len(matches) > 1:
        raise DjangoValidationError({field: f"{label}“{value}”匹配到多个结果，请使用唯一名称"})
    item = matches[0]
    if not getattr(item, "is_active", True):
        raise DjangoValidationError({field: f"停用的{label}不能用于新资产"})
    return item


def _named_location(queryset, value, field, label, *, lookup_field="name"):
    value = (value or "").strip()
    if not value:
        raise DjangoValidationError({field: f"{label}不能为空"})
    matches = list(queryset.filter(**{f"{lookup_field}__iexact": value}))
    if not matches:
        raise DjangoValidationError({field: f"未找到指定{label}“{value}”"})
    if len(matches) > 1:
        raise DjangoValidationError({field: f"{label}“{value}”定位不唯一，请补充正确的层级"})
    item = matches[0]
    if not item.is_active:
        raise DjangoValidationError({field: f"停用的{label}不能用于资产"})
    return item


def _split_values(value):
    return [item.strip() for item in re.split(r"[;,，；]", value or "") if item.strip()]


def _normalize_custom_option_value(field, value, header):
    """Accept an exact, unambiguous active option label as an import alias."""
    options = list(field.options.all())
    if value in {option.value for option in options}:
        return value
    label_matches = [option for option in options if option.label == value]
    active_matches = [option for option in label_matches if option.is_active]
    if len(label_matches) == 1 and len(active_matches) == 1:
        return active_matches[0].value
    if len(label_matches) > 1:
        raise DjangoValidationError({header: f"自定义字段选项标签“{value}”匹配到多个选项，必须填写唯一选项值"})
    return value


def _decimal_text(value, field, label, *, non_negative=False, max_decimal_places=None, max_digits=None):
    if not value:
        return ""
    try:
        parsed = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise DjangoValidationError({field: f"{label}格式不正确"}) from exc
    if not parsed.is_finite():
        raise DjangoValidationError({field: f"{label}必须是有限数字"})
    if non_negative and parsed < 0:
        raise DjangoValidationError({field: f"{label}不能小于 0"})
    if max_decimal_places is not None and max(0, -parsed.as_tuple().exponent) > max_decimal_places:
        raise DjangoValidationError({field: f"{label}最多支持 {max_decimal_places} 位小数"})
    if max_digits is not None:
        try:
            DecimalValidator(max_digits=max_digits, decimal_places=max_decimal_places or 0)(parsed)
        except DjangoValidationError as exc:
            raise DjangoValidationError({field: f"{label}超出数值范围"}) from exc
    text = format(parsed, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def _percentage_rate(value, field, label):
    value = (value or "").strip()
    if value.endswith("%"):
        value = value[:-1].strip()
    if not value:
        return ""
    try:
        percentage = Decimal(value)
    except (InvalidOperation, ValueError) as exc:
        raise DjangoValidationError({field: f"{label}格式不正确"}) from exc
    if not percentage.is_finite():
        raise DjangoValidationError({field: f"{label}必须是有限数字"})
    if percentage < 0 or percentage > 100:
        raise DjangoValidationError({field: f"{label}必须在 0 和 100 之间"})
    rate = percentage / Decimal("100")
    if max(0, -rate.normalize().as_tuple().exponent) > 4:
        raise DjangoValidationError({field: f"{label}最多支持 2 位小数"})
    return format(rate, "f").rstrip("0").rstrip(".") or "0"


def _iso_date(value, field, label):
    value = (value or "").strip()
    if not value:
        return ""
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError as exc:
        raise DjangoValidationError({field: f"{label}格式应为 YYYY-MM-DD"}) from exc


def _positive_integer(value, field, label):
    value = (value or "").strip()
    if not re.fullmatch(r"\d+", value):
        raise DjangoValidationError({field: f"{label}必须是整数"})
    parsed = int(value)
    if parsed < 1:
        raise DjangoValidationError({field: f"{label}必须大于等于 1"})
    return str(parsed)


def _location_text(row):
    values = [row.get(key, "").strip() for key in ("data_center", "server_room", "rack_code")]
    units = [row.get(key, "").strip() for key in ("rack_start_u", "rack_end_u")]
    if not any(values + units):
        return "未上架"
    location = " / ".join(item for item in values if item) or "位置未完整"
    if any(units):
        location = f"{location} · U{'-'.join(item for item in units if item)}"
    return location


def _depreciation_preview_text(row):
    start_date = row.get("depreciation_start_date", "").strip()
    years = row.get("depreciation_years", "").strip()
    residual_rate = row.get("residual_rate", "").strip()
    if not any((start_date, years, residual_rate)):
        return ""
    return f"{years or '—'} 年 · 残值率 {residual_rate or '—'}% · {start_date or '—'} 起"


def _resolve_assigned_person(row):
    employee_no = row.get("assigned_person_employee_no", "").strip()
    name = row.get("assigned_person_name", "").strip()
    department_name = row.get("assigned_person_department", "").strip()
    reason = row.get("assignment_reason", "").strip()
    if not any((employee_no, name, department_name, reason)):
        return None
    if reason and not (employee_no or name):
        raise DjangoValidationError({"assignment_reason": "填写使用人变更原因时必须同时指定使用人"})

    department = None
    if department_name:
        department = _named_active(
            Department.objects,
            department_name,
            "assigned_person_department",
            "使用人部门",
            allow_code=True,
        )

    if employee_no:
        matches = list(Person.objects.filter(employee_no__iexact=employee_no).select_related("department"))
        if not matches:
            raise DjangoValidationError({"assigned_person_employee_no": "未找到对应的人员"})
        person = matches[0]
        if name and person.name.casefold() != name.casefold():
            raise DjangoValidationError({"assigned_person_name": "姓名与员工编号对应的人员不一致"})
        if department is not None and person.department_id != department.pk:
            raise DjangoValidationError({"assigned_person_department": "部门与员工编号对应的人员不一致"})
    else:
        if not name:
            raise DjangoValidationError({"assigned_person_name": "没有员工编号时必须填写使用人姓名"})
        query = Person.objects.filter(name__iexact=name).select_related("department")
        if department is not None:
            query = query.filter(department_id=department.pk)
        matches = list(query)
        if not matches:
            raise DjangoValidationError({"assigned_person_name": "未找到对应的人员"})
        if len(matches) > 1:
            raise DjangoValidationError({"assigned_person_name": "姓名匹配到多名人员，请填写员工编号或所属部门"})
        person = matches[0]

    if not person.is_active:
        raise DjangoValidationError({"assigned_person_name": "停用人员不能作为使用人导入"})
    return {"action": "assign", "target_person": person.pk, "reason": reason}


def _prepare_payload(row, headers):
    asset_no = row.get("asset_no", "").strip()
    asset_name = row.get("name", "").strip()
    if not asset_no:
        raise DjangoValidationError({"asset_no": "资产编号不能为空"})
    if not asset_name:
        raise DjangoValidationError({"name": "资产名称不能为空"})

    manufacturer_name = row.get("manufacturer", "").strip()
    manufacturer = _named_active(Manufacturer.objects, manufacturer_name, "manufacturer", "厂商", allow_code=True) if manufacturer_name else None
    device_type_name = row.get("device_type", "").strip()
    if not device_type_name:
        raise DjangoValidationError({"device_type": "设备类型不能为空"})
    device_type = _named_active(DeviceType.objects, device_type_name, "device_type", "设备类型")

    asset_data_center_name = row.get("asset_data_center", "").strip()
    asset_data_center = _named_active(DataCenter.objects, asset_data_center_name, "asset_data_center", "数据中心") if asset_data_center_name else None

    status = row.get("status", "").strip() or get_system_settings().default_asset_status
    status = ASSET_STATUS_IMPORT_ALIASES.get(status, status)
    validate_asset_status_transition(None, status)

    configuration = dict(row)
    location_values = [row.get(key, "").strip() for key in ("data_center", "server_room", "rack_code", "rack_start_u", "rack_end_u")]
    if any(location_values):
        if not all(location_values):
            missing = [
                label for key, label in (
                    ("data_center", "数据中心"),
                    ("server_room", "机房"),
                    ("rack_code", "机柜编号"),
                    ("rack_start_u", "起始 U"),
                    ("rack_end_u", "结束 U"),
                ) if not row.get(key, "").strip()
            ]
            raise DjangoValidationError({"configuration": f"机柜位置未填写完整，请补充：{'、'.join(missing)}"})
        data_center = _named_location(DataCenter.objects, row["data_center"], "data_center", "数据中心")
        room = _named_location(ServerRoom.objects.filter(data_center=data_center), row["server_room"], "server_room", "机房")
        rack = _named_location(Rack.objects.filter(room=room), row["rack_code"], "rack_code", "机柜", lookup_field="code")
        configuration.update({
            "data_center": str(data_center.pk),
            "server_room_id": str(room.pk),
            "rack_id": str(rack.pk),
            "rack_start_u": _positive_integer(row["rack_start_u"], "rack_start_u", "起始 U"),
            "rack_end_u": _positive_integer(row["rack_end_u"], "rack_end_u", "结束 U"),
        })

    configuration["purchase_date"] = _iso_date(row.get("purchase_date"), "purchase_date", "采购日期")
    configuration["maintenance_start_date"] = _iso_date(row.get("maintenance_start_date"), "maintenance_start_date", "维保开始日")
    configuration["maintenance_expiry_date"] = _iso_date(row.get("maintenance_expiry_date"), "maintenance_expiry_date", "维保到期日")
    configuration["purchase_amount"] = _decimal_text(
        row.get("purchase_amount", ""),
        "purchase_amount",
        "采购金额",
        non_negative=True,
        max_decimal_places=2,
        max_digits=14,
    )

    depreciation_values = {
        "depreciation_start_date": row.get("depreciation_start_date", "").strip(),
        "depreciation_years": row.get("depreciation_years", "").strip(),
        "residual_rate": row.get("residual_rate", "").strip(),
    }
    if any(depreciation_values.values()):
        incomplete = {
            field: "折旧配置不完整"
            for field, value in depreciation_values.items()
            if not value
        }
        if incomplete:
            raise DjangoValidationError(incomplete)
        depreciation_start_date = _iso_date(
            depreciation_values["depreciation_start_date"],
            "depreciation_start_date",
            "折旧起算日",
        )
        depreciation_years = _positive_integer(
            depreciation_values["depreciation_years"],
            "depreciation_years",
            "折旧年限",
        )
        residual_rate = _percentage_rate(
            depreciation_values["residual_rate"],
            "residual_rate",
            "残值率",
        )
        depreciation_method = "straight_line"
    else:
        depreciation_start_date = None
        depreciation_years = None
        residual_rate = None
        depreciation_method = None

    custom_values = {}
    custom_headers = [header for header in headers if header.startswith("custom__")]
    fields_by_key = {
        field.key: field
        for field in CustomField.objects.filter(key__in=[header[8:] for header in custom_headers]).select_related("device_type").prefetch_related("options")
    }
    for header in custom_headers:
        field_key = header[8:]
        field = fields_by_key.get(field_key)
        if not field:
            raise DjangoValidationError({header: f"未知自定义字段编码“{field_key}”"})
        if not field.is_active:
            raise DjangoValidationError({header: "该自定义字段已停用，不能导入"})
        if field.device_type_id and field.device_type_id != device_type.pk:
            raise DjangoValidationError({header: f"该字段不适用于设备类型“{device_type.name}”"})
        raw = row.get(header, "").strip()
        if not raw:
            continue
        if field.field_type == "boolean":
            if raw.lower() in {"true", "1", "yes", "是"}:
                custom_values[field_key] = True
            elif raw.lower() in {"false", "0", "no", "否"}:
                custom_values[field_key] = False
            else:
                raise DjangoValidationError({header: "布尔值只能填写 true/false、是/否"})
        elif field.field_type == "multiselect":
            custom_values[field_key] = [
                _normalize_custom_option_value(field, value, header)
                for value in _split_values(raw)
            ]
        elif field.field_type == "select":
            custom_values[field_key] = _normalize_custom_option_value(field, raw, header)
        else:
            custom_values[field_key] = raw
        try:
            custom_values[field_key] = validate_custom_field_value(field, custom_values[field_key])
        except ValueError as exc:
            raise DjangoValidationError({header: str(exc)}) from exc

    tag_values = []
    for tag_name in _split_values(row.get("tags", "")):
        matches = list(Tag.objects.filter(name__iexact=tag_name))
        if not matches:
            raise DjangoValidationError({"tags": f"未找到标签“{tag_name}”"})
        if len(matches) > 1:
            raise DjangoValidationError({"tags": f"标签“{tag_name}”匹配到多个结果"})
        tag = matches[0]
        if not tag.is_active:
            raise DjangoValidationError({"tags": f"停用的标签“{tag_name}”不能用于新资产"})
        tag_values.append(tag.pk)

    return {
        "asset_no": asset_no,
        "name": asset_name,
        "manufacturer_id": manufacturer.pk if manufacturer else None,
        "device_type": device_type.pk,
        "asset_data_center": asset_data_center.pk if asset_data_center else None,
        "model": row.get("model", "").strip(),
        "manufacturer_model": row.get("manufacturer_model", "").strip(),
        "serial_number": row.get("serial_number", "").strip() or None,
        "purpose": row.get("purpose", "").strip(),
        "status": status,
        "notes": row.get("notes", "").strip(),
        "depreciation_start_date": depreciation_start_date,
        "depreciation_years": depreciation_years,
        "residual_rate": residual_rate,
        "depreciation_method": depreciation_method,
        "configuration": configuration,
        "tags": tag_values,
        "custom_values": custom_values,
        "assignment": _resolve_assigned_person(row),
    }


def _preview_row(line, row, headers, duplicate_counts, *, lock=False):
    asset_no = row.get("asset_no", "").strip()
    base = {
        "line": line,
        "asset_no": asset_no,
        "name": row.get("name", "").strip(),
        "device_type": row.get("device_type", "").strip(),
        "location": _location_text(row),
        "depreciation": _depreciation_preview_text(row),
        "valid": False,
        "action": "error",
        "changes": [],
        "errors": [],
    }
    if asset_no and duplicate_counts[asset_no] > 1:
        base["errors"] = [{"field": "asset_no", "label": "资产编号", "message": "文件内重复的资产编号"}]
        return base
    if asset_no:
        existing_query = Asset.objects.select_for_update() if lock else Asset.objects
        if existing_query.filter(asset_no=asset_no).exists():
            base["errors"] = [{"field": "asset_no", "label": "资产编号", "message": "资产编号已存在，资产导入仅支持新增"}]
            return base
    try:
        payload = _prepare_payload(row, headers)
        serializer = AssetWriteSerializer(data=payload)
        serializer.is_valid(raise_exception=True)
        with transaction.atomic():
            serializer.save()
            transaction.set_rollback(True)
    except Exception as exc:
        detail = getattr(exc, "detail", None) or getattr(exc, "message_dict", None) or str(exc)
        base["errors"] = _error_items(detail)
        return base
    base["valid"] = True
    base["action"] = "create"
    return base


def _preview_parsed(parsed, *, lock=False):
    duplicate_counts = Counter(row.get("asset_no", "").strip() for _line, row in parsed.rows if row.get("asset_no", "").strip())
    rows = [_preview_row(line, row, parsed.headers, duplicate_counts, lock=lock) for line, row in parsed.rows]
    valid = sum(1 for row in rows if row["valid"])
    invalid = len(rows) - valid
    return {
        "filename": parsed.filename,
        "total": len(rows),
        "valid": valid,
        "invalid": invalid,
        "summary": {"ready": valid, "conflicts": 0, "errors": invalid},
        "rows": rows,
    }


class AssetImportService:
    """Coordinate parse → normalize → preview → atomic commit for new assets."""

    @classmethod
    def parse(cls, upload):
        return _parse_upload(upload)

    @classmethod
    def preview(cls, upload):
        return _preview_parsed(cls.parse(upload))

    @classmethod
    def commit(cls, upload, request):
        parsed = cls.parse(upload)
        try:
            with transaction.atomic():
                preview = _preview_parsed(parsed, lock=True)
                if preview["invalid"]:
                    raise ImportValidationError(preview)
                created = 0
                for line, row in parsed.rows:
                    payload = _prepare_payload(row, parsed.headers)
                    serializer = AssetWriteSerializer(data=payload)
                    serializer.is_valid(raise_exception=True)
                    asset = serializer.save()
                    write_audit_log(
                        request,
                        action="create",
                        resource_type="asset",
                        resource_id=asset.pk,
                        after=asset_audit_snapshot(asset.pk),
                        extra={"source": "asset_import", "filename": parsed.filename, "line": line},
                    )
                    created += 1
                return {"created": created, "total": len(parsed.rows), "errors": []}
        except ImportValidationError:
            raise
        except (DjangoValidationError, DRFValidationError, IntegrityError):
            latest = _preview_parsed(parsed)
            raise ImportValidationError(latest, concurrent=True)


def build_import_template():
    """Build the current standard template from the same field definitions used by the importer."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "资产导入"
    active_fields = list(CustomField.objects.filter(is_active=True).select_related("device_type").order_by("device_type__name", "sort_order", "id"))
    headers = [key for key, _label, _required, _description in IMPORT_COLUMNS]
    headers.extend(f"custom__{field.key}" for field in active_fields)
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(len(header) + 4, 14), 28)

    guide = workbook.create_sheet("填写说明")
    guide.append(["模板格式", IMPORT_TEMPLATE_FORMAT])
    guide.append(["说明", "标准字段名位于资产导入 sheet 第一行；不要修改字段名。空单元格按当前资产表单语义处理。"])
    guide.append([])
    guide.append(["字段", "显示名称", "必填", "填写说明"])
    for key, label, required, description in IMPORT_COLUMNS:
        guide.append([key, label, "是" if required else "否", description])
    for field in active_fields:
        scope = field.device_type.name if field.device_type_id else "全部设备类型"
        custom_description = f"{scope} · 类型：{field.field_type}。{field.help_text or '按当前自定义字段选项填写。'}"
        if field.field_type in {"select", "multiselect"}:
            custom_description += " 单选/多选可填写启用选项的 value 或唯一显示名称；多选用英文分号、中文分号或逗号分隔，重复或停用选项不接受。"
        guide.append([f"custom__{field.key}", field.name, "是" if field.required else "否", custom_description])
    guide.append([])
    guide.append(["状态合法值", "可填 in_stock、in_use、idle、retired，或对应显示值：在库、在用、闲置、已报废；维修中仍只能由故障流程设置。"])
    guide.append(["日期格式", "Excel 日期单元格或 YYYY-MM-DD；不接受模糊日期。"])
    guide.append(["机柜位置", "数据中心 + 机房 + 机柜编号 + 起始 U + 结束 U 必须同时填写；位置按真实层级匹配并复用现有 U 位冲突校验。"])
    guide.append(["标签", "多个标签用英文分号、中文分号或逗号分隔；不存在或停用标签会阻止整批导入。"])
    guide.append(["资产台账导出", "导出文件是阅读型台账，不能直接上传；请使用本模板的标准字段名。引用导出值时，已上架资产的数据中心填入 data_center，未上架资产填入 asset_data_center；资产编号、序列号和网络地址仍需使用新的唯一值。"])
    guide.append(["导入策略", "只新增资产；已存在资产编号、文件内重复编号或任意校验错误都会阻止确认。"])
    guide.freeze_panes = "A5"
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 24
    guide.column_dimensions["C"].width = 12
    guide.column_dimensions["D"].width = 100
    for cell in guide[4]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    return workbook
