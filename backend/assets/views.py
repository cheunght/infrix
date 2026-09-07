from django.conf import settings
import logging
from django.db.models import BooleanField, Count, Exists, F, OuterRef, Q, Prefetch, Sum
from django.db.models.functions import Coalesce
from django.db import DatabaseError, IntegrityError, connection, transaction
from django.db.models.expressions import RawSQL
from django.db.models.deletion import ProtectedError
from uuid import uuid4
import re
from decimal import Decimal, InvalidOperation
from types import SimpleNamespace
from django.contrib.auth import authenticate, login, logout, update_session_auth_hash
from django.core.exceptions import ValidationError as DjangoValidationError
from django.middleware.csrf import get_token
from django.contrib.auth.models import Group, User
from django_filters.rest_framework import DjangoFilterBackend
from rest_framework import viewsets
from rest_framework.filters import OrderingFilter, SearchFilter
from rest_framework.decorators import action, api_view, permission_classes, parser_classes
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.permissions import AllowAny, IsAdminUser, IsAuthenticated
from rest_framework.response import Response
from rest_framework.exceptions import MethodNotAllowed, PermissionDenied, ValidationError as DRFValidationError
from drf_spectacular.types import OpenApiTypes
from drf_spectacular.utils import OpenApiParameter, extend_schema
from django.http import HttpResponse, QueryDict
from django.shortcuts import get_object_or_404
from django.utils import timezone
from datetime import date, datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from urllib.parse import quote
from .models import AuthThrottleState, AuditLog, Asset, AssetCustomValue, AssetNetworkAddress, AssetResponsibilityEvent, AssetTag, CustomField, CustomFieldOption, DataCenter, Department, DeviceType, DirectoryIdentity, FaultEvent, InventoryItem, InventoryTask, MaintenanceContract, Manufacturer, ProcurementRecord, Rack, RackUnitAllocation, RepairPartUsage, RepairRecord, ServerRoom, SoftwareLicense, SparePart, SparePartCategory, SpareStock, SpareStockTransaction, SystemSetting, Tag, UserSecurityProfile
from .enum_contracts import (
    ASSET_STATUS_LABELS,
    ASSET_STATUS_VALUES,
    INVENTORY_ITEM_STATUS_LABELS,
    INVENTORY_ITEM_STATUS_VALUES,
    INVENTORY_RESOLUTION_ACTION_LABELS,
    INVENTORY_RESOLUTION_STATUS_LABELS,
    INVENTORY_RESOLUTION_STATUS_VALUES,
    RACK_STATUS_VALUES,
    SPARE_UNIT_LABELS,
    STOCK_OPERATION_TYPE_LABELS,
    STOCK_OPERATION_TYPE_VALUES,
)
from .serializers import AdminPasswordResetSerializer, AssetBatchDeleteResponseSerializer, AssetBatchDeleteSerializer, AssetDetailSerializer, AssetListSerializer, AssetResponsibilityEventSerializer, AssetResponsibilityReturnSerializer, AssetResponsibilityTargetSerializer, AssetResponsibilityUserSerializer, AssetSerializer, AssetWriteSerializer, AuditLogSerializer, CurrentUserProfileSerializer, CustomFieldOptionSerializer, CustomFieldRuntimeSchemaSerializer, CustomFieldSerializer, DataCenterSerializer, DepartmentSerializer, DeviceTypeSerializer, FaultEventSerializer, GroupSerializer, InventoryBulkNormalResponseSerializer, InventoryBulkNormalSerializer, InventoryBulkResolutionResponseSerializer, InventoryBulkResolutionSerializer, InventoryInspectorSerializer, InventoryItemPageSerializer, InventoryItemSerializer, InventoryResolutionSerializer, InventoryScopePreviewQuerySerializer, InventoryScopePreviewSerializer, InventoryTaskSerializer, LdapConfigurationUpdateSerializer, ManufacturerSerializer, RackSerializer, RepairPartUsageCreateSerializer, RepairPartUsageSerializer, RepairRecordSerializer, ServerRoomSerializer, SoftwareLicenseSerializer, SparePartCategorySerializer, SparePartDetailSerializer, SparePartSerializer, SpareStockSerializer, SpareStockTransactionSerializer, SmtpTestEmailSerializer, SystemResetSerializer, SystemSettingsSerializer, TagSerializer, UserBatchStatusResponseSerializer, UserBatchStatusSerializer, UserSerializer, _default_references_option, _responsibility_user_name
from .services import (
    apply_spare_stock_transaction,
    confirm_inventory_item_normal,
    create_repair_part_usage,
    inventory_task_delete_block_reason,
    reopen_repair,
    reset_inventory_resolution,
    resolve_inventory_item,
    sync_asset_fault_status,
    sync_fault_completion,
    sync_repair_completion,
    update_asset_placement,
    assign_asset,
    return_asset,
    transfer_asset,
)
from .inventory import get_inventory_scope_assets
from .depreciation import calculate_asset_depreciation
from .license_status import LICENSE_STATUS_KEYS, LICENSE_STATUS_LABELS, filter_licenses_by_status, license_status_counts, license_status_value
from .audit import asset_audit_snapshot, asset_custom_value_changes, model_snapshot, software_license_audit_snapshot, spare_part_audit_snapshot, spare_stock_transaction_audit_snapshot, write_audit_log
from .imports import AssetImportService, ImportFileError, ImportValidationError, build_import_template
from .ldap_auth import (
    AUTH_SOURCE_LDAP,
    AUTH_SOURCE_LOCAL,
    LDAP_DIAGNOSTIC_CHECKS,
    LDAP_MODEL_BACKEND,
    LDAP_DIAGNOSTIC_MESSAGES,
    LDAPDirectoryClient,
    AuthenticationFailure,
    authenticate_with_source_routing,
    is_directory_managed,
    ldap_status_snapshot,
    ldap_is_enabled,
)
from .ldap_configuration import (
    ConfigurationIdentityError,
    ConfigurationSecretError,
    configuration_errors,
    get_effective_ldap_configuration,
    merge_configuration,
    public_configuration,
    save_configuration,
)
from .configuration_secrets import encrypt_secret
from .smtp import SmtpConfigurationError, send_smtp_test_email
from .permissions import BusinessRolePermission, CanExportAssets, CanExportFaults, CanExportInventory, CanExportLicenses, CanExportRacks, CanExportSpares, CanImportAssets, CanManageInventory, CanManageSystemSettings, CanResetSystem, CanViewAssetCustomFieldSchema, CanViewAssetTagsRuntime, CanViewAuditLog, CanViewDashboard, CanViewDepartmentRuntime, CanViewInventory, CanViewLicenses, CanViewManufacturerRuntime, CanViewSparePartCategoryRuntime, IsSystemAdministrator
from .roles import ROLE_DEFINITIONS, ROLE_NAME_TO_CODE, user_capabilities, user_has_capability, user_role_code, user_role_codes
from .reporting import (
    DashboardScopeError,
    build_dashboard_payload,
    build_alerts_payload,
    build_rack_capacity_rows,
    rack_effective_used_u,
    resolve_dashboard_scope,
)
from .reporting.constants import RACK_LAYOUT_EXPORT_MAX_RACKS, RACK_LAYOUT_EXPORT_MAX_U_POSITIONS
from .pagination import StandardPagination
from .system_reset import reset_system
from .system_settings import (
    get_local_account_security_policy,
    get_system_settings,
    local_password_expired,
    system_localdate,
    system_localtime,
    system_settings_snapshot,
    validate_local_password,
)


EXPORT_MAX_ROWS = 10_000
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
logger = logging.getLogger(__name__)


def _export_timestamp():
    return system_localtime().strftime("%Y%m%d_%H%M%S")


def _excel_value(value):
    """Keep user text as text so editable exports cannot inject formulas."""
    if value is None:
        return ""
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return f"'{value}"
    return value


def _append_excel_row(sheet, values):
    sheet.append([_excel_value(value) for value in values])


def _style_export_sheet(sheet, max_width=40):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
    for column in sheet.columns:
        values = [len(str(cell.value or "")) for cell in column]
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max(max(values) + 2, 12), max_width)
    sheet.freeze_panes = "A2"


def _xlsx_response(book, filename):
    response = HttpResponse(content_type=XLSX_CONTENT_TYPE)
    ascii_filename = re.sub(r"[^A-Za-z0-9._-]+", "_", filename).strip("_") or "export.xlsx"
    response["Content-Disposition"] = (
        f'attachment; filename="{ascii_filename}"; filename*=UTF-8\'\'{quote(filename, safe="")}'
    )
    book.save(response)
    return response


def _server_room_location_snapshot(room_id):
    room = ServerRoom.objects.select_related("data_center").get(pk=room_id)
    return {
        "data_center_id": room.data_center_id,
        "data_center": room.data_center.name,
        "server_room_id": room.pk,
        "server_room": room.name,
    }


def _data_center_location_snapshot(data_center_id):
    data_center = DataCenter.objects.get(pk=data_center_id)
    return {
        "data_center_id": data_center.pk,
        "data_center": data_center.name,
    }


def _export_limit_response(queryset, label):
    count = queryset.count()
    if count > EXPORT_MAX_ROWS:
        return Response({"detail": f"{label}超过 {EXPORT_MAX_ROWS} 条，请缩小筛选范围后重试"}, status=400)
    return None


def _export_decimal(value):
    if value is None or value == "":
        return ""
    return Decimal(str(value))


def _export_residual_rate(value):
    if value is None or value == "":
        return ""
    rate = Decimal(str(value)) * Decimal("100")
    return f"{format(rate, 'f').rstrip('0').rstrip('.') or '0'}%"


def _filtered_view_queryset(view_class, request, *, action="export", filter_backends=None):
    """Run the same DRF queryset/filter contract without pagination."""
    view = view_class()
    view.request = request
    view.args = ()
    view.kwargs = {}
    view.action = action
    if filter_backends is not None:
        view.filter_backends = filter_backends
    return view.filter_queryset(view.get_queryset())


class AuditedModelViewSetMixin:
    audit_resource = None

    def audit_snapshot(self, instance):
        return model_snapshot(instance)

    def audit_extra(self, before, after, *, action):
        return None

    @transaction.atomic
    def perform_create(self, serializer):
        instance = serializer.save()
        after = self.audit_snapshot(instance)
        write_audit_log(
            self.request,
            action="create",
            resource_type=self.audit_resource,
            resource_id=instance.pk,
            after=after,
            extra=self.audit_extra(None, after, action="create"),
        )

    @transaction.atomic
    def perform_update(self, serializer):
        before = self.audit_snapshot(serializer.instance)
        instance = serializer.save()
        after = self.audit_snapshot(instance)
        write_audit_log(
            self.request,
            action="update",
            resource_type=self.audit_resource,
            resource_id=instance.pk,
            before=before,
            after=after,
            extra=self.audit_extra(before, after, action="update"),
        )

    @transaction.atomic
    def perform_destroy(self, instance):
        before = self.audit_snapshot(instance)
        resource_id = instance.pk
        instance.delete()
        write_audit_log(
            self.request,
            action="delete",
            resource_type=self.audit_resource,
            resource_id=resource_id,
            before=before,
        )


def _date_filter_errors(request, fields):
    """Return a field error for malformed ISO dates before they reach SQL."""
    for field in fields:
        value = request.query_params.get(field, "").strip()
        if value:
            try:
                date.fromisoformat(value)
            except ValueError:
                return {field: "日期必须使用 YYYY-MM-DD 格式"}
    return None


CUSTOM_FILTER_MAX_CONDITIONS = 8
CUSTOM_FILTER_OPERATOR_SUFFIXES = ("contains", "gte", "lte", "eq")
CUSTOM_FILTER_OPERATORS_BY_TYPE = {
    "text": {"contains", "eq"},
    "textarea": {"contains", "eq"},
    "number": {"eq", "gte", "lte"},
    "date": {"eq", "gte", "lte"},
    "boolean": {"eq"},
    "select": {"eq"},
    "multiselect": {"contains"},
}


def _custom_filter_key_and_operator(query_key):
    field_part = query_key[len("custom__"):]
    for operator in CUSTOM_FILTER_OPERATOR_SUFFIXES:
        suffix = f"__{operator}"
        if field_part.endswith(suffix):
            return field_part[:-len(suffix)], operator
    return field_part, None


def _custom_filter_boolean(value):
    normalized = value.lower()
    if normalized in {"true", "1", "yes", "是"}:
        return True
    if normalized in {"false", "0", "no", "否"}:
        return False
    raise ValueError("必须是是或否")


def _custom_filter_multiselect_membership(value):
    if connection.vendor == "sqlite":
        return RawSQL(
            "EXISTS (SELECT 1 FROM json_each(json_value) WHERE json_type(json_value) = 'array' AND json_each.value = %s)",
            [value],
            output_field=BooleanField(),
        )
    if connection.vendor == "mysql":
        return RawSQL(
            "JSON_TYPE(json_value) = 'ARRAY' AND JSON_CONTAINS(json_value, JSON_QUOTE(%s), '$')",
            [value],
            output_field=BooleanField(),
        )
    return Q(json_value__contains=[value])


def _batch_error_message(error, fallback):
    detail = getattr(error, "detail", error)

    def flatten(value):
        if isinstance(value, dict):
            return "；".join(flatten(child) for child in value.values())
        if isinstance(value, (list, tuple)):
            return "；".join(flatten(child) for child in value)
        return str(value or "")

    return flatten(detail) or fallback


def _delete_asset_with_audit(instance, request, *, batch_operation_id=None):
    before = asset_audit_snapshot(instance.pk)
    resource_id = instance.pk
    try:
        instance.delete()
    except ProtectedError as exc:
        protected = list(exc.protected_objects)
        if any(isinstance(item, InventoryItem) for item in protected):
            raise DRFValidationError("资产存在历史盘点记录，不能删除") from exc
        if any(isinstance(item, FaultEvent) for item in protected):
            raise DRFValidationError("资产存在关联故障记录，不能删除") from exc
        if any(isinstance(item, AssetResponsibilityEvent) for item in protected):
            raise DRFValidationError("资产存在责任变化历史，不能删除") from exc
        raise DRFValidationError("资产存在关联数据，不能删除") from exc
    extra = {"batch_operation_id": str(batch_operation_id)} if batch_operation_id else None
    write_audit_log(
        request,
        action="delete",
        resource_type="asset",
        resource_id=resource_id,
        before=before,
        extra=extra,
    )


class AssetViewSet(AuditedModelViewSetMixin, viewsets.ModelViewSet):
    max_custom_columns = 12
    queryset = Asset.objects.select_related("department", "responsible_user", "manufacturer", "device_type", "asset_data_center", "rack_allocation__rack__room__data_center").prefetch_related("network_addresses", "procurement_records", "maintenance_contracts", "asset_tags__tag", "custom_values__field__options").order_by("asset_no", "id")
    serializer_class = AssetSerializer
    permission_classes = [BusinessRolePermission]
    permission_resource = "assets"
    audit_resource = "asset"
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ["status", "department", "manufacturer", "device_type", "model"]
    ordering_fields = ["asset_no", "name", "manufacturer_model", "serial_number"]
    ordering = ["asset_no", "id"]
    search_fields = [
        "asset_no", "name", "manufacturer_model", "serial_number", "purpose", "notes", "status",
        "manufacturer__name", "device_type__name", "device_type__color", "model", "department__name", "department__code",
        "responsible_user__username", "responsible_user__first_name", "responsible_user__last_name", "responsible_user__email",
        "network_addresses__address", "network_addresses__role", "network_addresses__status", "network_addresses__notes",
        "rack_allocation__rack__code", "rack_allocation__rack__room__name", "rack_allocation__rack__room__data_center__name",
        "rack_allocation__start_u", "rack_allocation__end_u",
        "procurement_records__purchase_date", "procurement_records__supplier", "procurement_records__order_no", "procurement_records__amount", "procurement_records__notes",
        "maintenance_contracts__provider", "maintenance_contracts__contract_no", "maintenance_contracts__notes",
        "maintenance_contracts__start_date", "maintenance_contracts__expiry_date",
        "asset_tags__tag__name", "custom_values__text_value", "custom_values__json_value",
    ]

    @extend_schema(
        parameters=[
            OpenApiParameter(
                name="custom_columns",
                type=OpenApiTypes.STR,
                required=False,
                description="compact 台账列表需要返回的动态字段 key，支持逗号分隔或重复参数，最多 12 个。",
            ),
            OpenApiParameter(
                name="compact",
                type=OpenApiTypes.BOOL,
                required=False,
                description="使用资产台账的轻量列表响应。",
            ),
            OpenApiParameter(
                name="tags",
                type=OpenApiTypes.STR,
                required=False,
                description="按标签 ID 筛选资产，支持逗号分隔或重复参数；多个标签为 OR（命中任一标签）。",
            ),
            OpenApiParameter(
                name="data_center",
                type=OpenApiTypes.INT,
                required=False,
                description="按数据中心 ID 筛选资产；已上架资产按机柜所属数据中心，未上架资产按所属数据中心。",
            ),
            OpenApiParameter(
                name="warranty",
                type=OpenApiTypes.STR,
                required=False,
                enum=["within_30_days", "expired"],
                description="按维保到期状态筛选：within_30_days 表示今天至未来 30 天内到期，expired 表示已过期。",
            ),
            OpenApiParameter(
                name="custom__{field_key}__{operator}",
                type=OpenApiTypes.STR,
                required=False,
                description="动态字段筛选；operator 按字段类型使用 eq、contains、gte 或 lte。多个条件为 AND。",
            ),
            OpenApiParameter(
                name="ordering",
                type=OpenApiTypes.STR,
                required=False,
                enum=[
                    "asset_no", "-asset_no", "name", "-name",
                    "manufacturer_model", "-manufacturer_model",
                    "serial_number", "-serial_number",
                ],
                description="资产列表排序字段；支持资产编号、名称、厂商/型号和序列号，前缀 - 表示降序。",
            ),
        ]
    )
    def list(self, request, *args, **kwargs):
        return super().list(request, *args, **kwargs)

    def get_queryset(self):
        queryset = super().get_queryset()
        if self.request.query_params.get("search", "").strip():
            queryset = queryset.distinct()
        if self.action == "list" and self.request.query_params.get("compact", "").lower() in {"1", "true", "yes"}:
            requested_custom_columns = self._requested_custom_columns()
            queryset = queryset.select_related(
                "responsible_user", "manufacturer", "device_type", "asset_data_center", "rack_allocation__rack__room__data_center"
            ).prefetch_related(None).prefetch_related(
                Prefetch(
                    "network_addresses",
                    queryset=AssetNetworkAddress.objects.only("id", "asset_id", "role", "address"),
                ),
                Prefetch(
                    "procurement_records",
                    queryset=ProcurementRecord.objects.only("id", "asset_id", "purchase_date", "supplier", "order_no", "amount").order_by("-purchase_date", "-id"),
                ),
                Prefetch(
                    "maintenance_contracts",
                    queryset=MaintenanceContract.objects.only("id", "asset_id", "provider", "expiry_date").order_by("-updated_at", "-id"),
                ),
                Prefetch("asset_tags", queryset=AssetTag.objects.select_related("tag")),
            )
            if requested_custom_columns:
                queryset = queryset.prefetch_related(
                    Prefetch(
                        "custom_values",
                        queryset=AssetCustomValue.objects.filter(
                            field__key__in=requested_custom_columns,
                            field__is_active=True,
                            field__list_visible=True,
                        ).select_related("field"),
                        to_attr="list_custom_values",
                    )
                )
        tag_ids = self._tag_filter_ids()
        if tag_ids:
            queryset = queryset.filter(asset_tags__tag_id__in=tag_ids)
        data_center = self.request.query_params.get("data_center", "").strip()
        try:
            data_center_id = int(data_center) if data_center else None
        except (TypeError, ValueError):
            data_center_id = None
        if data_center_id and data_center_id > 0:
            queryset = queryset.filter(
                Q(rack_allocation__rack__room__data_center_id=data_center_id)
                | Q(rack_allocation__isnull=True, asset_data_center_id=data_center_id)
            ).distinct()
        warranty = self.request.query_params.get("warranty", "").strip()
        if warranty == "within_30_days":
            today = system_localdate()
            queryset = queryset.filter(
                maintenance_contracts__expiry_date__gte=today,
                maintenance_contracts__expiry_date__lte=today + timedelta(days=30),
            ).distinct()
        elif warranty == "expired":
            today = system_localdate()
            queryset = queryset.filter(maintenance_contracts__expiry_date__lt=today).distinct()
        for field, operator, value in self._validated_custom_filters():
            values = AssetCustomValue.objects.filter(
                asset_id=OuterRef("pk"),
                field_id=field.id,
                field__is_active=True,
            ).filter(
                Q(field__device_type_id__isnull=True)
                | Q(field__device_type_id=OuterRef("device_type_id")),
            )
            if field.field_type in {"text", "textarea"}:
                lookup = "text_value__icontains" if operator == "contains" else "text_value"
                values = values.filter(**{lookup: value})
            elif field.field_type == "number":
                lookup = {"eq": "exact", "gte": "gte", "lte": "lte"}[operator]
                values = values.filter(**{f"number_value__{lookup}": value})
            elif field.field_type == "date":
                lookup = {"eq": "exact", "gte": "gte", "lte": "lte"}[operator]
                values = values.filter(**{f"date_value__{lookup}": value})
            elif field.field_type == "boolean":
                values = values.filter(boolean_value=value)
            elif field.field_type == "select":
                values = values.filter(text_value=value)
            else:
                values = values.filter(_custom_filter_multiselect_membership(value))
            queryset = queryset.filter(Exists(values))
        if tag_ids:
            queryset = queryset.distinct()
        return queryset

    def filter_queryset(self, queryset):
        """Keep list pagination deterministic after applying the ordering whitelist."""
        queryset = super().filter_queryset(queryset)
        ordering = list(queryset.query.order_by or ())
        if not ordering:
            ordering = list(self.ordering)
        if not any(str(term).lstrip("-") == "id" for term in ordering):
            last_term = str(ordering[-1]) if ordering else "asset_no"
            ordering.append("-id" if last_term.startswith("-") else "id")
        return queryset.order_by(*ordering)

    def _tag_filter_ids(self):
        """Parse the multi-tag ID filter."""
        raw_values = self.request.query_params.getlist("tags")
        if not raw_values:
            return []
        values = []
        invalid = []
        for raw_value in raw_values:
            for part in str(raw_value).split(","):
                value = part.strip()
                if not value:
                    continue
                if not value.isdigit() or int(value) <= 0:
                    invalid.append(value)
                    continue
                values.append(int(value))
        if invalid:
            raise DRFValidationError({"tags": "标签筛选参数必须是正整数 ID"})
        return list(dict.fromkeys(values))

    def _validated_custom_filters(self):
        raw_conditions = [
            (query_key, value)
            for query_key, values in self.request.query_params.lists()
            if query_key.startswith("custom__")
            for value in values
        ]
        if not raw_conditions:
            return []
        if len(raw_conditions) > CUSTOM_FILTER_MAX_CONDITIONS:
            raise DRFValidationError({"custom_filters": f"动态筛选条件最多 {CUSTOM_FILTER_MAX_CONDITIONS} 条"})

        parsed = []
        field_keys = set()
        for query_key, raw_value in raw_conditions:
            field_key, explicit_operator = _custom_filter_key_and_operator(query_key)
            parsed.append((query_key, field_key, explicit_operator, raw_value))
            if field_key:
                field_keys.add(field_key)
        fields = {
            field.key: field
            for field in CustomField.objects.filter(key__in=field_keys).select_related("device_type").prefetch_related("options")
        }
        errors = []
        validated = []
        for query_key, field_key, explicit_operator, raw_value in parsed:
            if explicit_operator is None:
                errors.append(f"{query_key}：必须指定操作符")
                continue
            field = fields.get(field_key)
            if not field:
                errors.append(f"{query_key}：字段不存在")
                continue
            if not field.is_active:
                errors.append(f"{field.name}：字段已停用")
                continue
            if not field.filterable:
                errors.append(f"{field.name}：字段未开启筛选")
                continue
            operator = explicit_operator
            if operator not in CUSTOM_FILTER_OPERATORS_BY_TYPE.get(field.field_type, set()):
                errors.append(f"{field.name}：不支持“{operator}”操作")
                continue
            value = str(raw_value).strip()
            if not value:
                errors.append(f"{field.name}：筛选值不能为空")
                continue
            try:
                if field.field_type == "number":
                    value = Decimal(value)
                    if not value.is_finite():
                        raise ValueError("必须是有限数字")
                elif field.field_type == "date":
                    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
                        raise ValueError("必须使用 YYYY-MM-DD 格式")
                    value = date.fromisoformat(value)
                elif field.field_type == "boolean":
                    value = _custom_filter_boolean(value)
                elif field.field_type in {"select", "multiselect"}:
                    active_options = {option.value for option in field.options.all() if option.is_active}
                    if value not in active_options:
                        raise ValueError("选项值不存在或已停用")
            except (InvalidOperation, ValueError) as exc:
                errors.append(f"{field.name}：{exc}")
                continue
            validated.append((field, operator, value))
        if errors:
            raise DRFValidationError({"custom_filters": errors})
        return validated

    def _requested_custom_columns(self) -> "list[str]":
        cached = getattr(self, "_requested_custom_columns_cache", None)
        if cached is not None:
            return cached

        raw_values = self.request.query_params.getlist("custom_columns")
        keys = []
        for raw_value in raw_values:
            keys.extend(part.strip() for part in raw_value.split(",") if part.strip())
        keys = list(dict.fromkeys(keys))
        if len(keys) > self.max_custom_columns:
            raise DRFValidationError({
                "custom_columns": f"动态列最多同时请求 {self.max_custom_columns} 个字段",
            })
        if not keys:
            self._requested_custom_columns_cache = []
            return []

        available = set(
            CustomField.objects.filter(
                key__in=keys,
                is_active=True,
                list_visible=True,
            ).values_list("key", flat=True)
        )
        missing = [key for key in keys if key not in available]
        if missing:
            raise DRFValidationError({
                "custom_columns": f"动态列不存在、已停用或未开启列表展示：{', '.join(missing)}",
            })
        self._requested_custom_columns_cache = keys
        return keys

    def get_serializer_context(self):
        context = super().get_serializer_context()
        if self.action == "list" and self.request.query_params.get("compact", "").lower() in {"1", "true", "yes"}:
            context["requested_custom_columns"] = self._requested_custom_columns()
        else:
            context["requested_custom_columns"] = []
        return context

    def get_serializer_class(self):
        if self.action in {"create", "update", "partial_update"}:
            return AssetWriteSerializer
        if self.action == "retrieve":
            return AssetDetailSerializer
        if self.request.query_params.get("compact", "").lower() in {"1", "true", "yes"}:
            return AssetListSerializer
        return AssetSerializer

    @extend_schema(
        responses=AssetResponsibilityUserSerializer(many=True),
        parameters=[
            OpenApiParameter(
                name="search",
                type=OpenApiTypes.STR,
                required=False,
                description="按用户名、姓名或邮箱搜索可作为资产责任人的启用用户。",
            ),
        ],
        description="返回资产责任动作可选择的启用用户，使用标准分页。",
    )
    @action(detail=False, methods=["get"], url_path="responsibility-users")
    def responsibility_users(self, request):
        queryset = User.objects.filter(is_active=True)
        search = request.query_params.get("search", "").strip()
        if search:
            queryset = queryset.filter(
                Q(username__icontains=search)
                | Q(first_name__icontains=search)
                | Q(last_name__icontains=search)
                | Q(email__icontains=search)
            )
        queryset = queryset.order_by("username", "id")
        page = self.paginate_queryset(queryset)
        serializer = AssetResponsibilityUserSerializer(page if page is not None else queryset, many=True)
        if page is not None:
            return self.get_paginated_response(serializer.data)
        return Response(serializer.data)

    @extend_schema(
        responses=AssetResponsibilityEventSerializer(many=True),
        description="按时间倒序分页返回资产责任人变化历史。历史记录只读。",
    )
    @action(detail=True, methods=["get"], url_path="responsibility-history")
    def responsibility_history(self, request, pk=None):
        asset = self.get_object()
        queryset = AssetResponsibilityEvent.objects.filter(asset_id=asset.pk).select_related(
            "from_user", "to_user", "operator"
        ).order_by("-created_at", "-id")
        page = self.paginate_queryset(queryset)
        serializer = AssetResponsibilityEventSerializer(page if page is not None else queryset, many=True)
        if page is not None:
            return self.get_paginated_response(serializer.data)
        return Response(serializer.data)

    @extend_schema(
        request=AssetResponsibilityTargetSerializer,
        responses=AssetDetailSerializer,
        description="将当前未分配责任人的资产领用给一个启用用户。",
    )
    @action(detail=True, methods=["post"], url_path="assign")
    def assign(self, request, pk=None):
        serializer = AssetResponsibilityTargetSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        asset, _event = assign_asset(
            asset_id=pk,
            target_user_id=serializer.validated_data["target_user"].pk,
            actor=request.user,
            request=request,
            reason=serializer.validated_data.get("reason", ""),
        )
        return Response(AssetDetailSerializer(asset, context=self.get_serializer_context()).data)

    @extend_schema(
        request=AssetResponsibilityReturnSerializer,
        responses=AssetDetailSerializer,
        description="归还当前有责任人的资产并清空当前责任人。",
    )
    @action(detail=True, methods=["post"], url_path="return")
    def return_asset(self, request, pk=None):
        serializer = AssetResponsibilityReturnSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        asset, _event = return_asset(
            asset_id=pk,
            actor=request.user,
            request=request,
            reason=serializer.validated_data.get("reason", ""),
        )
        return Response(AssetDetailSerializer(asset, context=self.get_serializer_context()).data)

    @extend_schema(
        request=AssetResponsibilityTargetSerializer,
        responses=AssetDetailSerializer,
        description="将资产责任人从当前用户调拨给另一个启用用户。",
    )
    @action(detail=True, methods=["post"], url_path="transfer")
    def transfer(self, request, pk=None):
        serializer = AssetResponsibilityTargetSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        asset, _event = transfer_asset(
            asset_id=pk,
            target_user_id=serializer.validated_data["target_user"].pk,
            actor=request.user,
            request=request,
            reason=serializer.validated_data.get("reason", ""),
        )
        return Response(AssetDetailSerializer(asset, context=self.get_serializer_context()).data)

    def audit_snapshot(self, instance):
        return asset_audit_snapshot(instance.pk)

    def audit_extra(self, before, after, *, action):
        if action not in {"create", "update"}:
            return None
        changes = asset_custom_value_changes(
            (before or {}).get("custom_value_snapshot"),
            (after or {}).get("custom_value_snapshot"),
        )
        return {"custom_changes": changes} if changes else None

    def perform_destroy(self, instance):
        try:
            # Serialize the single-delete path with history writers that lock
            # the asset before appending an immutable business record.  Keep
            # the atomic block inside the try so a deferred FK violation at
            # transaction exit is translated into the same 4xx contract as a
            # Django ProtectedError.
            with transaction.atomic():
                locked_instance = get_object_or_404(
                    Asset.objects.select_for_update(),
                    pk=instance.pk,
                )
                _delete_asset_with_audit(locked_instance, self.request)
        except IntegrityError as exc:
            raise DRFValidationError("资产存在关联数据，不能删除") from exc

    @extend_schema(
        request=AssetBatchDeleteSerializer,
        responses=AssetBatchDeleteResponseSerializer,
        description="批量删除资产；每条记录会重新执行单条删除保护并返回逐条结果。",
    )
    @action(detail=False, methods=["post"], url_path="batch-delete")
    def batch_delete(self, request):
        request_serializer = AssetBatchDeleteSerializer(data=request.data)
        request_serializer.is_valid(raise_exception=True)
        asset_ids = request_serializer.validated_data["ids"]
        batch_operation_id = uuid4()
        results = []
        succeeded = 0
        for asset_id in asset_ids:
            asset_no = f"ID {asset_id}"
            try:
                with transaction.atomic():
                    instance = Asset.objects.select_for_update().get(pk=asset_id)
                    asset_no = instance.asset_no
                    _delete_asset_with_audit(
                        instance,
                        request,
                        batch_operation_id=batch_operation_id,
                    )
            except Asset.DoesNotExist:
                results.append({
                    "id": asset_id,
                    "asset_no": f"ID {asset_id}",
                    "success": False,
                    "code": "NOT_FOUND",
                    "reason": "资产不存在或已被删除",
                })
            except DRFValidationError as exc:
                results.append({
                    "id": asset_id,
                    "asset_no": asset_no,
                    "success": False,
                    "code": "PROTECTED",
                    "reason": _batch_error_message(exc, "资产当前不能删除"),
                })
            except DatabaseError:
                results.append({
                    "id": asset_id,
                    "asset_no": asset_no,
                    "success": False,
                    "code": "CONFLICT",
                    "reason": "资产当前存在关联数据，未删除",
                })
            else:
                succeeded += 1
                results.append({
                    "id": asset_id,
                    "asset_no": asset_no,
                    "success": True,
                    "code": "",
                    "reason": "",
                })
        response_data = {
            "requested": len(asset_ids),
            "succeeded": succeeded,
            "failed": len(asset_ids) - succeeded,
            "results": results,
        }
        return Response(AssetBatchDeleteResponseSerializer(response_data).data)


class RackViewSet(AuditedModelViewSetMixin, viewsets.ModelViewSet):
    queryset = Rack.objects.select_related("room", "room__data_center").prefetch_related(
        Prefetch("allocations", queryset=RackUnitAllocation.objects.select_related("asset", "asset__manufacturer", "asset__device_type", "rack__room__data_center"))
    )
    serializer_class = RackSerializer
    permission_classes = [BusinessRolePermission]
    permission_resource = "racks"
    audit_resource = "rack"
    filter_backends = [DjangoFilterBackend, SearchFilter]
    filterset_fields = ["room", "room__data_center", "code", "status"]
    search_fields = ["code", "name", "rack_type", "owner_name", "room__name", "room__data_center__name"]

    @extend_schema(
        parameters=[
            OpenApiParameter(
                name="device_type",
                type=OpenApiTypes.INT,
                required=False,
                description="按已分配资产的设备类型 ID 筛选机柜。",
            ),
        ]
    )
    def list(self, request, *args, **kwargs):
        return super().list(request, *args, **kwargs)

    def get_queryset(self):
        queryset = super().get_queryset()
        device_type = self.request.query_params.get("device_type", "").strip()
        if device_type:
            try:
                device_type_id = int(device_type)
            except ValueError:
                return queryset.none()
            if device_type_id <= 0:
                return queryset.none()
            queryset = queryset.filter(allocations__asset__device_type_id=device_type_id).distinct()
        active = self.request.query_params.get("is_active", "true").strip().lower()
        if self.action in {"retrieve", "update", "partial_update", "destroy"}:
            return queryset.order_by("room__data_center__name", "room__name", "code")
        if active in {"true", "false"}:
            queryset = queryset.filter(is_active=active == "true")
        status = self.request.query_params.get("status", "").strip()
        if status in RACK_STATUS_VALUES:
            queryset = queryset.filter(status=status)
        return queryset.order_by("room__data_center__name", "room__name", "code")

    def audit_extra(self, before, after, *, action):
        if action != "update" or before.get("room") == after.get("room"):
            return None
        return {
            "source": "rack_location_correction",
            "old_location": _server_room_location_snapshot(before["room"]),
            "new_location": _server_room_location_snapshot(after["room"]),
            "affected_assets": RackUnitAllocation.objects.filter(rack_id=after["id"]).count(),
        }

    @transaction.atomic
    def perform_update(self, serializer):
        # Match the rack -> allocations lock order used by asset placement and
        # repeat validation after acquiring the locks. Otherwise a concurrent
        # placement can make an earlier capacity check stale before save.
        locked_rack = Rack.objects.select_for_update().get(pk=serializer.instance.pk)
        target_room = serializer.validated_data.get("room")
        target_room_id = target_room.pk if target_room is not None else locked_rack.room_id
        locked_target_room = ServerRoom.objects.select_for_update().get(pk=target_room_id)
        DataCenter.objects.select_for_update().get(pk=locked_target_room.data_center_id)
        list(
            RackUnitAllocation.objects.select_for_update()
            .filter(rack_id=locked_rack.pk)
            .order_by("pk")
        )
        locked_serializer = self.get_serializer(
            locked_rack,
            data=self.request.data,
            partial=self.request.method == "PATCH",
        )
        locked_serializer.is_valid(raise_exception=True)
        super().perform_update(locked_serializer)
        serializer.instance = locked_serializer.instance

    @transaction.atomic
    def perform_destroy(self, instance):
        from django.db.models.deletion import ProtectedError

        instance = Rack.objects.select_for_update().get(pk=instance.pk)
        if instance.allocations.exists():
            raise DRFValidationError("机柜仍有资产占用，不能删除，请先迁移资产或停用机柜")
        try:
            super().perform_destroy(instance)
        except ProtectedError as exc:
            raise DRFValidationError("机柜仍有资产占用，不能删除，请先迁移资产或停用机柜") from exc


class ServerRoomViewSet(AuditedModelViewSetMixin, viewsets.ModelViewSet):
    queryset = ServerRoom.objects.select_related("data_center").annotate(
        racks_count=Count("racks", distinct=True),
        assets_count=Count("racks__allocations__asset", distinct=True),
    ).order_by("data_center__name", "name")
    serializer_class = ServerRoomSerializer
    permission_classes = [BusinessRolePermission]
    permission_resource = "racks"
    audit_resource = "server_room"
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ["data_center", "is_active"]
    search_fields = ["name", "data_center__name"]
    ordering_fields = ["name", "created_at", "updated_at"]

    def get_queryset(self):
        queryset = super().get_queryset()
        if self.action in {"retrieve", "update", "partial_update", "destroy"}:
            return queryset
        active = self.request.query_params.get("is_active", "true").strip().lower()
        if active in {"true", "false"}:
            queryset = queryset.filter(is_active=active == "true")
        return queryset

    def audit_extra(self, before, after, *, action):
        if action != "update" or before.get("data_center") == after.get("data_center"):
            return None
        room_id = after["id"]
        return {
            "source": "server_room_location_correction",
            "old_location": {
                **_data_center_location_snapshot(before["data_center"]),
                "server_room_id": room_id,
                "server_room": before.get("name", ""),
            },
            "new_location": {
                **_data_center_location_snapshot(after["data_center"]),
                "server_room_id": room_id,
                "server_room": after.get("name", ""),
            },
            "affected_racks": Rack.objects.filter(room_id=room_id).count(),
            "affected_assets": RackUnitAllocation.objects.filter(rack__room_id=room_id).count(),
        }

    @transaction.atomic
    def perform_update(self, serializer):
        needs_location_lock = "data_center" in self.request.data
        locked_room = None
        if needs_location_lock:
            # A room reparenting changes the effective data center of every
            # rack below it. Lock the same rack -> allocation hierarchy used
            # by asset placement, then validate again against the locked rows.
            locked_room = ServerRoom.objects.select_for_update().get(pk=serializer.instance.pk)
            target_data_center = serializer.validated_data.get("data_center")
            target_data_center_id = (
                target_data_center.pk
                if target_data_center is not None
                else locked_room.data_center_id
            )
            DataCenter.objects.select_for_update().get(pk=target_data_center_id)
            list(
                Rack.objects.select_for_update()
                .filter(room_id=locked_room.pk)
                .order_by("pk")
            )
            list(
                RackUnitAllocation.objects.select_for_update()
                .filter(rack__room_id=locked_room.pk)
                .order_by("pk")
            )
            locked_serializer = self.get_serializer(
                locked_room,
                data=self.request.data,
                partial=self.request.method == "PATCH",
            )
            locked_serializer.is_valid(raise_exception=True)
            if locked_serializer.validated_data.get("is_active") is False:
                if locked_room.is_active and locked_room.spare_stocks.filter(quantity__gt=0).exists():
                    from rest_framework.exceptions import ValidationError as DRFValidationError
                    raise DRFValidationError("机房仍有备件库存，无法停用，请先调出、出库或报废库存")
            super().perform_update(locked_serializer)
            serializer.instance = locked_serializer.instance
            return
        elif serializer.validated_data.get("is_active") is False:
            locked_room = ServerRoom.objects.select_for_update().get(pk=serializer.instance.pk)
        if serializer.validated_data.get("is_active") is False:
            if locked_room is None:
                locked_room = ServerRoom.objects.select_for_update().get(pk=serializer.instance.pk)
            if locked_room.is_active and locked_room.spare_stocks.filter(quantity__gt=0).exists():
                from rest_framework.exceptions import ValidationError as DRFValidationError
                raise DRFValidationError("机房仍有备件库存，无法停用，请先调出、出库或报废库存")
        super().perform_update(serializer)

    def perform_destroy(self, instance):
        if instance.racks.exists():
            racks = list(instance.racks.values_list("code", flat=True)[:5])
            raise DRFValidationError(
                f"机房仍包含机柜（{', '.join(racks)}），不能删除，请先迁移或删除机柜后再操作"
            )
        if instance.inventory_tasks.exists():
            raise DRFValidationError("机房存在盘点任务记录，不能删除")
        try:
            super().perform_destroy(instance)
        except ProtectedError as exc:
            if instance.inventory_tasks.exists():
                raise DRFValidationError("机房存在盘点任务记录，不能删除") from exc
            if instance.spare_stocks.exists():
                raise DRFValidationError("机房仍有备件库存，不能删除，请先调整库存地点") from exc
            if instance.spare_source_transactions.exists() or instance.spare_target_transactions.exists():
                raise DRFValidationError("机房存在备件库存流水，不能删除") from exc
            raise DRFValidationError("机房仍被其他业务数据引用，不能删除") from exc


class DataCenterViewSet(AuditedModelViewSetMixin, viewsets.ModelViewSet):
    queryset = DataCenter.objects.annotate(
        assets_count=(
            Count("rooms__racks__allocations__asset", distinct=True)
            + Count(
                "unmounted_assets",
                filter=Q(unmounted_assets__rack_allocation__isnull=True),
                distinct=True,
            )
        ),
        rooms_count=Count("rooms", distinct=True),
    ).order_by("name")
    serializer_class = DataCenterSerializer
    permission_classes = [BusinessRolePermission]
    permission_resource = "racks"
    audit_resource = "data_center"

    def get_queryset(self):
        queryset = super().get_queryset()
        # Keep inactive data centers addressable for administrator actions.
        # The list defaults to active entries, but update/delete must still
        # resolve a previously disabled entry instead of returning 404.
        if self.action in {"retrieve", "update", "partial_update", "destroy"}:
            return queryset
        active = self.request.query_params.get("is_active", "true").strip().lower()
        if active in {"true", "false"}:
            queryset = queryset.filter(is_active=active == "true")
        search = self.request.query_params.get("search", "").strip()
        if search:
            queryset = queryset.filter(Q(name__icontains=search) | Q(address__icontains=search))
        return queryset

    @transaction.atomic
    def perform_update(self, serializer):
        if serializer.validated_data.get("is_active") is False:
            locked_data_center = DataCenter.objects.select_for_update().get(pk=serializer.instance.pk)
            if locked_data_center.is_active and locked_data_center.spare_stocks.filter(quantity__gt=0).exists():
                from rest_framework.exceptions import ValidationError as DRFValidationError
                raise DRFValidationError("数据中心仍有备件库存，无法停用，请先调出、出库或报废库存")
        super().perform_update(serializer)

    def perform_destroy(self, instance):
        if instance.rooms.exists():
            from rest_framework.exceptions import ValidationError as DRFValidationError
            rooms = list(instance.rooms.values_list("name", flat=True)[:5])
            raise DRFValidationError(f"数据中心仍包含机房（{', '.join(rooms)}），不能删除，请先删除机房或停用")
        try:
            super().perform_destroy(instance)
        except Exception as exc:
            from django.db.models.deletion import ProtectedError
            from rest_framework.exceptions import ValidationError as DRFValidationError
            if isinstance(exc, ProtectedError):
                if instance.unmounted_assets.exists():
                    raise DRFValidationError("数据中心仍有未上架资产归属，不能删除，请先调整资产所属数据中心") from exc
                if instance.inventory_tasks.exists():
                    raise DRFValidationError("数据中心仍有历史盘点任务，不能删除") from exc
                if instance.spare_stocks.exists():
                    raise DRFValidationError("数据中心仍有备件库存，不能删除，请先调整库存地点") from exc
                if instance.spare_source_transactions.exists() or instance.spare_target_transactions.exists():
                    raise DRFValidationError("数据中心存在备件库存流水，不能删除") from exc
            raise

    def destroy(self, request, *args, **kwargs):
        # Resolve by primary key directly so an inactive data center is not
        # hidden by the list filter during DELETE and turned into a 404.
        from rest_framework.exceptions import NotFound
        try:
            instance = DataCenter.objects.get(pk=kwargs.get(self.lookup_field, kwargs.get("pk")))
        except DataCenter.DoesNotExist as exc:
            raise NotFound("数据中心不存在") from exc
        self.perform_destroy(instance)
        from rest_framework.response import Response
        return Response(status=204)


class DictionaryViewSet(AuditedModelViewSetMixin, viewsets.ModelViewSet):
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    search_fields = ["name"]
    ordering_fields = ["name", "created_at", "updated_at"]
    ordering = ["name", "id"]

    permission_classes = [BusinessRolePermission]
    permission_resource = "settings"

    def can_view_inactive(self):
        return "settings.manage" in user_capabilities(self.request.user)

    def get_queryset(self):
        queryset = self.queryset.annotate(assets_count=Count("assets"))
        active = self.request.query_params.get("is_active", "true").strip().lower()
        if not self.can_view_inactive():
            active = "true"
        if active in {"true", "false"}:
            queryset = queryset.filter(is_active=active == "true")
        return queryset

    def perform_create(self, serializer):
        super().perform_create(serializer)
        instance = serializer.instance
        instance.assets_count = 0

    def perform_update(self, serializer):
        super().perform_update(serializer)
        instance = serializer.instance
        instance.assets_count = instance.assets.count()

    def perform_destroy(self, instance):
        if instance.assets.exists():
            from rest_framework.exceptions import ValidationError as DRFValidationError
            raise DRFValidationError("字典项正在被资产使用，不能删除，请先停用")
        super().perform_destroy(instance)


class DepartmentViewSet(AuditedModelViewSetMixin, viewsets.ModelViewSet):
    queryset = Department.objects.select_related("parent").annotate(
        assets_count=Count("assets", distinct=True),
    ).order_by("name", "id")
    serializer_class = DepartmentSerializer
    permission_classes = [BusinessRolePermission]
    permission_resource = "settings"
    audit_resource = "department"
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ["parent"]
    search_fields = ["name", "code", "parent__name"]
    ordering_fields = ["name", "code", "created_at", "updated_at"]
    ordering = ["name", "id"]

    def get_permissions(self):
        if self.action in {"list", "retrieve"}:
            return [CanViewDepartmentRuntime()]
        return super().get_permissions()

    def perform_destroy(self, instance):
        if instance.assets.exists():
            raise DRFValidationError("部门正在被资产使用，不能删除，请先调整资产归属")
        try:
            super().perform_destroy(instance)
        except ProtectedError as exc:
            raise DRFValidationError("部门仍有子部门使用，不能删除，请先调整部门层级") from exc


class ManufacturerViewSet(DictionaryViewSet):
    queryset = Manufacturer.objects.all()
    serializer_class = ManufacturerSerializer
    audit_resource = "manufacturer"

    def get_permissions(self):
        if self.action in {"list", "retrieve"}:
            return [CanViewManufacturerRuntime()]
        return super().get_permissions()

    def can_view_inactive(self):
        if super().can_view_inactive():
            return True
        return self.action in {"list", "retrieve"} and any(
            user_has_capability(self.request.user, capability)
            for capability in (
                "assets.view",
                "assets.manage",
                "licenses.view",
                "licenses.manage",
                "spares.view",
                "spares.manage",
            )
        )

    def get_queryset(self):
        return super().get_queryset().annotate(
            assets_count=Count("assets", distinct=True),
            licenses_count=Count("software_licenses", distinct=True),
            spare_parts_count=Count("spare_parts", distinct=True),
        )

    def perform_create(self, serializer):
        super().perform_create(serializer)
        instance = serializer.instance
        instance.assets_count = 0
        instance.licenses_count = 0
        instance.spare_parts_count = 0

    def perform_update(self, serializer):
        super().perform_update(serializer)
        instance = serializer.instance
        instance.assets_count = instance.assets.count()
        instance.licenses_count = instance.software_licenses.count()
        instance.spare_parts_count = instance.spare_parts.count()

    def perform_destroy(self, instance):
        if instance.assets.exists() or instance.software_licenses.exists() or instance.spare_parts.exists():
            raise DRFValidationError("厂商正在被资产、软件许可或备件使用，不能删除，请先停用")
        super().perform_destroy(instance)


class DeviceTypeViewSet(DictionaryViewSet):
    queryset = DeviceType.objects.all()
    serializer_class = DeviceTypeSerializer
    audit_resource = "device_type"

    def get_queryset(self):
        return super().get_queryset().annotate(
            custom_fields_count=Count("custom_fields", distinct=True),
        )

    def perform_create(self, serializer):
        super().perform_create(serializer)
        serializer.instance.custom_fields_count = 0

    def perform_update(self, serializer):
        super().perform_update(serializer)
        serializer.instance.custom_fields_count = serializer.instance.custom_fields.count()

    def perform_destroy(self, instance):
        if instance.assets.exists():
            raise DRFValidationError("设备类型正在被资产使用，不能删除，请先停用")
        if instance.custom_fields.exists():
            raise DRFValidationError("设备类型仍被自定义字段使用，不能删除，请先停用或解除字段绑定")
        try:
            super().perform_destroy(instance)
        except ProtectedError as exc:
            raise DRFValidationError("设备类型仍被业务数据使用，不能删除，请先停用") from exc


class SparePartCategoryViewSet(AuditedModelViewSetMixin, viewsets.ModelViewSet):
    queryset = SparePartCategory.objects.annotate(spare_parts_count=Count("spare_parts", distinct=True)).order_by("name", "id")
    serializer_class = SparePartCategorySerializer
    permission_classes = [BusinessRolePermission]
    permission_resource = "settings"
    audit_resource = "spare_part_category"
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ["is_active"]
    search_fields = ["name", "code"]
    ordering_fields = ["name", "code", "created_at", "updated_at"]

    def get_permissions(self):
        if self.request.method in {"GET", "HEAD", "OPTIONS"} or self.action in {"list", "retrieve"}:
            return [CanViewSparePartCategoryRuntime()]
        return super().get_permissions()

    def get_queryset(self):
        queryset = super().get_queryset()
        if self.action in {"retrieve", "update", "partial_update", "destroy"}:
            return queryset
        active = self.request.query_params.get("is_active", "true").strip().lower()
        if "settings.manage" not in user_capabilities(self.request.user):
            active = "true"
        if active in {"true", "false"}:
            queryset = queryset.filter(is_active=active == "true")
        return queryset

    def perform_destroy(self, instance):
        if instance.spare_parts.exists():
            raise DRFValidationError("备件类型正在被备件使用，不能删除，请先停用")
        super().perform_destroy(instance)


class CustomFieldViewSet(AuditedModelViewSetMixin, viewsets.ModelViewSet):
    queryset = CustomField.objects.select_related("device_type").prefetch_related("options").annotate(
        assets_count=Count("asset_values__asset", distinct=True)
    ).order_by("device_type__name", "sort_order", "id")
    serializer_class = CustomFieldSerializer
    permission_classes = [BusinessRolePermission]
    permission_resource = "custom_fields"
    audit_resource = "custom_field"
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ["device_type", "field_type", "is_active"]
    search_fields = ["key", "name", "device_type__name"]
    ordering_fields = ["name", "sort_order", "created_at", "updated_at"]

    def get_permissions(self):
        if self.action == "schema_fields":
            return [CanViewAssetCustomFieldSchema()]
        return super().get_permissions()

    def get_queryset(self):
        queryset = super().get_queryset()
        if self.action not in {"retrieve", "update", "partial_update", "destroy"}:
            active = self.request.query_params.get("is_active", "true").strip().lower()
            if "custom_fields.manage" not in user_capabilities(self.request.user):
                active = "true"
            if active in {"true", "false"}:
                queryset = queryset.filter(is_active=active == "true")
        return queryset

    @extend_schema(
        parameters=[
            OpenApiParameter(
                name="device_type",
                type=OpenApiTypes.INT,
                required=False,
                description="仅返回全局字段与指定设备类型字段。",
            ),
            OpenApiParameter(
                name="list_visible",
                type=OpenApiTypes.BOOL,
                required=False,
                description="仅返回启用且允许台账列展示的字段。",
            ),
            OpenApiParameter(
                name="filterable",
                type=OpenApiTypes.BOOL,
                required=False,
                description="仅返回启用且允许台账筛选的字段；返回选项仅包含启用选项。",
            ),
        ],
        responses=CustomFieldRuntimeSchemaSerializer(many=True),
    )
    @action(detail=False, methods=["get"], url_path="schema", pagination_class=None, filter_backends=[])
    def schema_fields(self, request):
        list_visible = request.query_params.get("list_visible", "").strip().lower() in {"1", "true", "yes"}
        filterable = request.query_params.get("filterable", "").strip().lower() in {"1", "true", "yes"}
        if list_visible or filterable:
            filters = {"is_active": True}
            if list_visible:
                filters["list_visible"] = True
            if filterable:
                filters["filterable"] = True
            options = CustomFieldOption.objects.all().order_by("sort_order", "id")
            if filterable:
                options = options.filter(is_active=True)
            queryset = CustomField.objects.filter(**filters).select_related("device_type").prefetch_related(
                Prefetch("options", queryset=options)
            ).order_by("device_type__name", "sort_order", "id")
            return Response(CustomFieldRuntimeSchemaSerializer(queryset, many=True).data)

        device_type = request.query_params.get("device_type", "").strip()
        if device_type:
            scope = Q(device_type__isnull=True) | Q(device_type_id=device_type)
        else:
            scope = Q(device_type__isnull=True)
        queryset = CustomField.objects.filter(is_active=True).filter(scope).select_related("device_type").prefetch_related(
            Prefetch("options", queryset=CustomFieldOption.objects.filter(is_active=True))
        ).order_by("sort_order", "id")
        return Response(CustomFieldRuntimeSchemaSerializer(queryset, many=True).data)

    def perform_destroy(self, instance):
        from rest_framework.exceptions import ValidationError as DRFValidationError
        if instance.asset_values.exists():
            raise DRFValidationError("字段已有资产值，不能删除，请先停用")
        if instance.options.exists():
            raise DRFValidationError("字段仍有选项，不能删除，请先删除选项")
        super().perform_destroy(instance)


class CustomFieldOptionViewSet(AuditedModelViewSetMixin, viewsets.ModelViewSet):
    queryset = CustomFieldOption.objects.select_related("field", "field__device_type").order_by("field_id", "sort_order", "id")
    serializer_class = CustomFieldOptionSerializer
    permission_classes = [BusinessRolePermission]
    permission_resource = "custom_fields"
    audit_resource = "custom_field_option"
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ["field", "is_active"]
    search_fields = ["value", "label", "field__name"]
    ordering_fields = ["sort_order", "created_at", "updated_at"]

    def perform_destroy(self, instance):
        from rest_framework.exceptions import ValidationError as DRFValidationError
        if _default_references_option(instance.field, instance.value):
            raise DRFValidationError("选项被字段默认值引用，不能删除")
        used = False
        for text_value, json_value in AssetCustomValue.objects.filter(
            field_id=instance.field_id
        ).values_list("text_value", "json_value"):
            if text_value == instance.value or (
                isinstance(json_value, list) and instance.value in json_value
            ):
                used = True
                break
        if used:
            raise DRFValidationError("选项已有资产使用，不能删除，请先停用")
        super().perform_destroy(instance)


class TagViewSet(AuditedModelViewSetMixin, viewsets.ModelViewSet):
    queryset = Tag.objects.annotate(assets_count=Count("asset_tags__asset", distinct=True)).order_by("name", "id")
    serializer_class = TagSerializer
    permission_classes = [BusinessRolePermission]
    permission_resource = "tags"
    audit_resource = "tag"
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ["is_active"]
    search_fields = ["name"]
    ordering_fields = ["name", "created_at", "updated_at"]

    def get_permissions(self):
        if self.action in {"list", "retrieve"} or self.request.method in {"GET", "HEAD", "OPTIONS"}:
            return [CanViewAssetTagsRuntime()]
        return super().get_permissions()

    def get_queryset(self):
        queryset = super().get_queryset()
        if self.action not in {"retrieve", "update", "partial_update", "destroy"}:
            active = self.request.query_params.get("is_active", "true").strip().lower()
            if "tags.manage" not in user_capabilities(self.request.user):
                active = "true"
            if active in {"true", "false"}:
                queryset = queryset.filter(is_active=active == "true")
        return queryset

    def perform_destroy(self, instance):
        from rest_framework.exceptions import ValidationError as DRFValidationError
        if instance.asset_tags.exists():
            raise DRFValidationError("标签正在被资产使用，不能删除，请先停用")
        super().perform_destroy(instance)


class SparePartViewSet(AuditedModelViewSetMixin, viewsets.ModelViewSet):
    queryset = SparePart.objects.select_related("manufacturer", "category").annotate(
        total_quantity=Coalesce(Sum("stocks__quantity"), 0),
        location_count=Count("stocks", distinct=True),
        stock_movement_exists=Exists(SpareStockTransaction.objects.filter(part=OuterRef("pk"))),
    ).order_by("name", "category__name", "id")
    serializer_class = SparePartSerializer
    permission_classes = [BusinessRolePermission]
    permission_resource = "spares"
    audit_resource = "spare_part"
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ["category", "manufacturer"]
    search_fields = [
        "code", "name", "category__name", "category__code", "manufacturer__name", "model", "specification",
        "storage_location", "notes",
    ]
    ordering_fields = ["code", "name", "category", "category__name", "safety_stock", "created_at", "updated_at"]

    def audit_snapshot(self, instance):
        return spare_part_audit_snapshot(instance)

    def get_serializer_class(self):
        if self.action == "retrieve":
            return SparePartDetailSerializer
        return super().get_serializer_class()

    def get_queryset(self):
        queryset = super().get_queryset()
        # A part can be stored at more than one location. Filter through its
        # stock balances and de-duplicate the result when it matches several
        # locations.
        data_center = self.request.query_params.get("data_center", "").strip()
        server_room = self.request.query_params.get("server_room", "").strip()
        if data_center:
            queryset = queryset.filter(stocks__data_center_id=data_center)
        if server_room:
            queryset = queryset.filter(stocks__server_room_id=server_room)
        if data_center or server_room:
            queryset = queryset.distinct()
        return queryset

    @transaction.atomic
    def perform_create(self, serializer):
        initial_quantity = serializer.validated_data.pop("initial_quantity", 0)
        initial_data_center = serializer.validated_data.pop("initial_data_center", None)
        initial_server_room = serializer.validated_data.pop("initial_server_room", None)
        instance = serializer.save()
        after = self.audit_snapshot(instance)
        write_audit_log(
            self.request,
            action="create",
            resource_type=self.audit_resource,
            resource_id=instance.pk,
            after=after,
            extra=self.audit_extra(None, after, action="create"),
        )
        if initial_quantity <= 0:
            return
        try:
            transaction_row = apply_spare_stock_transaction(
                {
                    "part": instance,
                    "operation_type": "initial",
                    "quantity": initial_quantity,
                    "target_data_center": initial_data_center,
                    "target_server_room": initial_server_room,
                    "reference": "初始库存",
                    "notes": "新增备件时登记的初始库存",
                },
                self.request.user,
            )
        except DjangoValidationError as exc:
            detail = getattr(exc, "message_dict", None) or {"detail": "; ".join(exc.messages)}
            raise DRFValidationError(detail) from exc
        write_audit_log(
            self.request,
            action="create",
            resource_type="spare_stock_transaction",
            resource_id=transaction_row.pk,
            after=SpareStockTransactionSerializer(transaction_row).data,
        )

    @transaction.atomic
    def perform_destroy(self, instance):
        from rest_framework.exceptions import ValidationError as DRFValidationError
        # Keep the delete lock order identical to apply_spare_stock_transaction:
        # part first, then related stock/ledger rows.
        try:
            locked_part = SparePart.objects.select_for_update().get(pk=instance.pk)
        except SparePart.DoesNotExist as exc:
            raise DRFValidationError("备件不存在") from exc
        if locked_part.transactions.exists():
            raise DRFValidationError("该备件已有库存流水，无法删除。")
        if locked_part.repair_part_usages.exists():
            raise DRFValidationError("该备件已有维修用件记录，无法删除。")
        # Stock balances without a movement are not historical records. They
        # must be removed before the part because their FK is also PROTECT;
        # deletion eligibility is determined only by transaction history.
        locked_part.stocks.all().delete()
        super().perform_destroy(locked_part)


class SpareStockViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = SpareStock.objects.select_related("part", "part__category", "part__manufacturer", "data_center", "server_room").order_by(
        "part__name", "data_center__name", "server_room__name", "id"
    )
    serializer_class = SpareStockSerializer
    permission_classes = [BusinessRolePermission]
    permission_resource = "spares"
    filter_backends = [DjangoFilterBackend, OrderingFilter]
    filterset_fields = ["part", "data_center", "server_room"]
    ordering_fields = ["quantity", "updated_at"]


class SpareStockTransactionViewSet(viewsets.ModelViewSet):
    queryset = SpareStockTransaction.objects.select_related(
        "part", "part__category", "part__manufacturer", "operator", "source_data_center", "source_server_room",
        "target_data_center", "target_server_room",
    ).order_by("-created_at", "-id")
    serializer_class = SpareStockTransactionSerializer
    permission_classes = [BusinessRolePermission]
    permission_resource = "spares"
    http_method_names = ["get", "post", "head", "options"]
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ["part", "operation_type", "source_data_center", "target_data_center", "operator"]
    search_fields = ["part__name", "part__model", "reference", "notes", "operator__username"]
    ordering_fields = ["created_at", "quantity_delta", "before_quantity", "after_quantity", "operation_type"]

    def get_queryset(self):
        queryset = super().get_queryset()
        start = self.request.query_params.get("start", "").strip()
        end = self.request.query_params.get("end", "").strip()
        if start:
            try:
                queryset = queryset.filter(created_at__date__gte=date.fromisoformat(start))
            except ValueError:
                return queryset.none()
        if end:
            try:
                queryset = queryset.filter(created_at__date__lte=date.fromisoformat(end))
            except ValueError:
                return queryset.none()
        return queryset

    def list(self, request, *args, **kwargs):
        errors = _date_filter_errors(request, ("start", "end"))
        if errors:
            return Response(errors, status=400)
        return super().list(request, *args, **kwargs)

    @transaction.atomic
    def perform_create(self, serializer):
        from rest_framework.exceptions import ValidationError as DRFValidationError
        audit_context = {}
        try:
            transaction_row = apply_spare_stock_transaction(
                serializer.validated_data,
                self.request.user,
                audit_context=audit_context,
            )
        except DjangoValidationError as exc:
            detail = getattr(exc, "message_dict", None) or {"detail": "; ".join(exc.messages)}
            raise DRFValidationError(detail) from exc
        serializer.instance = transaction_row
        write_audit_log(
            self.request,
            action="create",
            resource_type="spare_stock_transaction",
            resource_id=transaction_row.pk,
            after=spare_stock_transaction_audit_snapshot(
                transaction_row,
                transfer_context=audit_context,
            ),
        )


class SoftwareLicenseViewSet(AuditedModelViewSetMixin, viewsets.ModelViewSet):
    queryset = SoftwareLicense.objects.select_related("manufacturer")
    serializer_class = SoftwareLicenseSerializer
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ["manufacturer"]
    search_fields = ["name", "manufacturer__name", "license_type", "notes"]
    ordering_fields = ["name", "manufacturer__name", "expiry_date", "authorized_count", "used_count", "created_at"]
    ordering = ["expiry_date", "name", "id"]
    permission_classes = [BusinessRolePermission]
    permission_resource = "licenses"
    audit_resource = "software_license"

    def audit_snapshot(self, instance):
        return software_license_audit_snapshot(instance)

    def get_queryset(self):
        queryset = super().get_queryset()
        status = self.request.query_params.get("status", "").strip()
        return filter_licenses_by_status(queryset, status)


class GroupViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = Group.objects.filter(name__in=ROLE_NAME_TO_CODE).annotate(user_count=Count("user")).order_by("id")
    serializer_class = GroupSerializer
    permission_classes = [IsSystemAdministrator]
    filter_backends = [SearchFilter, OrderingFilter]
    search_fields = ["name"]
    ordering_fields = ["name", "user_count"]


def _user_audit_snapshot(user):
    role_code = user_role_code(user)
    return {
        "username": user.username,
        "display_name": user.get_full_name() or user.username,
        "first_name": user.first_name,
        "last_name": user.last_name,
        "email": user.email,
        "role_code": role_code,
        "role_name": ROLE_DEFINITIONS.get(role_code, {}).get("name", ""),
        "is_active": user.is_active,
    }


class UserViewSet(viewsets.ModelViewSet):
    queryset = User.objects.select_related("directory_identity").prefetch_related("groups").order_by("username")
    serializer_class = UserSerializer
    permission_classes = [IsSystemAdministrator]
    filter_backends = [SearchFilter, OrderingFilter]
    search_fields = ["username", "first_name", "last_name", "email"]
    ordering_fields = ["username", "date_joined", "last_login"]

    @transaction.atomic
    def perform_create(self, serializer):
        instance = serializer.save()
        write_audit_log(
            self.request,
            action="create",
            resource_type="user",
            resource_id=instance.pk,
            after=_user_audit_snapshot(instance),
        )

    def _validate_protected_update(self, serializer):
        instance = serializer.instance
        requested_role = serializer.validated_data.get("role_code")
        requested_active = serializer.validated_data.get("is_active")
        current_role = user_role_code(instance)
        is_current_user = instance.pk == self.request.user.pk
        if requested_role is not None and requested_role != current_role:
            if instance.is_superuser:
                raise DRFValidationError({"role_code": "不能修改超级管理员角色"})
            if is_current_user:
                raise DRFValidationError({"role_code": "不能修改当前登录账号的角色"})
        if requested_active is False and (instance.is_superuser or is_current_user):
            raise DRFValidationError({"is_active": "不能停用当前登录账号或超级管理员"})

    @transaction.atomic
    def perform_update(self, serializer):
        self._validate_protected_update(serializer)
        before = _user_audit_snapshot(serializer.instance)
        instance = serializer.save()
        batch_operation_id = getattr(self, "_batch_operation_id", None)
        write_audit_log(
            self.request,
            action="update",
            resource_type="user",
            resource_id=instance.pk,
            before=before,
            after=_user_audit_snapshot(instance),
            extra={"batch_operation_id": batch_operation_id} if batch_operation_id else None,
        )

    @extend_schema(
        request=UserBatchStatusSerializer,
        responses=UserBatchStatusResponseSerializer,
        description="批量启用或停用用户；每条记录会重新执行单条用户保护校验并返回逐条结果。",
    )
    @action(detail=False, methods=["post"], url_path="batch-status")
    def batch_status(self, request):
        request_serializer = UserBatchStatusSerializer(data=request.data)
        request_serializer.is_valid(raise_exception=True)
        user_ids = request_serializer.validated_data["ids"]
        is_active = request_serializer.validated_data["is_active"]
        batch_operation_id = str(uuid4())
        results = []
        succeeded = 0
        self._batch_operation_id = batch_operation_id
        try:
            for user_id in user_ids:
                username = f"ID {user_id}"
                try:
                    with transaction.atomic():
                        instance = User.objects.select_for_update().prefetch_related("groups").get(pk=user_id)
                        username = instance.username
                        serializer = UserSerializer(
                            instance=instance,
                            data={"is_active": is_active},
                            partial=True,
                            context={"request": request},
                        )
                        serializer.is_valid(raise_exception=True)
                        self.perform_update(serializer)
                except User.DoesNotExist:
                    results.append({
                        "id": user_id,
                        "username": username,
                        "success": False,
                        "code": "NOT_FOUND",
                        "reason": "用户不存在或已被删除",
                    })
                except DRFValidationError as exc:
                    detail = getattr(exc, "detail", {})
                    code = "PROTECTED" if isinstance(detail, dict) and "is_active" in detail else "INVALID_STATE"
                    results.append({
                        "id": user_id,
                        "username": username,
                        "success": False,
                        "code": code,
                        "reason": _batch_error_message(exc, "用户当前不能更新状态"),
                    })
                except DatabaseError:
                    results.append({
                        "id": user_id,
                        "username": username,
                        "success": False,
                        "code": "CONFLICT",
                        "reason": "用户状态更新发生并发冲突，请重试",
                    })
                else:
                    succeeded += 1
                    results.append({
                        "id": user_id,
                        "username": username,
                        "success": True,
                        "code": "",
                        "reason": "",
                    })
        finally:
            delattr(self, "_batch_operation_id")

        response_data = {
            "requested": len(user_ids),
            "succeeded": succeeded,
            "failed": len(user_ids) - succeeded,
            "results": results,
        }
        return Response(UserBatchStatusResponseSerializer(response_data).data)

    @action(detail=True, methods=["post"], url_path="reset-password")
    @transaction.atomic
    def reset_password(self, request, pk=None):
        instance = self.get_object()
        if is_directory_managed(instance):
            raise DRFValidationError(
                {
                    "detail": "Directory-managed accounts must change passwords through the corporate directory.",
                    "code": "directory_password_managed",
                }
            )
        serializer = AdminPasswordResetSerializer(
            data=request.data,
            context={"user": instance},
        )
        serializer.is_valid(raise_exception=True)
        instance.set_password(serializer.validated_data["new_password"])
        instance.save(update_fields=["password"])
        profile = _security_profile(instance)
        profile.must_change_password = True
        profile.password_changed_at = None
        profile.save(update_fields=["must_change_password", "password_changed_at", "updated_at"])
        write_audit_log(
            request,
            action="update",
            resource_type="user",
            resource_id=instance.pk,
            extra={"password_reset": True, "target_username": instance.username},
        )
        return Response({"ok": True})

    @transaction.atomic
    def perform_destroy(self, instance):
        if instance.pk == self.request.user.pk:
            raise DRFValidationError("不能删除当前登录账号")
        if instance.is_superuser:
            raise DRFValidationError("不能通过业务接口删除超级管理员")
        if is_directory_managed(instance):
            raise DRFValidationError(
                {
                    "detail": "Directory-managed accounts must be removed from the corporate directory first.",
                    "code": "directory_user_protected",
                }
            )
        if Asset.objects.filter(responsible_user_id=instance.pk).exists():
            raise DRFValidationError("用户仍是资产责任人，归还或调拨资产后才能删除")
        if InventoryTask.objects.filter(inspector_id=instance.pk).exists():
            raise DRFValidationError("用户仍被盘点任务引用，不能删除")
        before = _user_audit_snapshot(instance)
        resource_id = instance.pk
        try:
            instance.delete()
        except ProtectedError as exc:
            raise DRFValidationError("用户仍被其他业务数据引用，不能删除") from exc
        write_audit_log(
            self.request,
            action="delete",
            resource_type="user",
            resource_id=resource_id,
            before=before,
        )


class FaultEventViewSet(AuditedModelViewSetMixin, viewsets.ModelViewSet):
    queryset = FaultEvent.objects.select_related("asset", "repair")
    serializer_class = FaultEventSerializer
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ["asset", "is_closed"]
    search_fields = ["asset__asset_no", "asset__name", "reason", "description"]
    ordering_fields = ["occurred_at", "created_at"]
    ordering = ["-occurred_at"]
    permission_classes = [BusinessRolePermission]
    permission_resource = "faults"
    audit_resource = "fault_event"

    def get_queryset(self):
        queryset = super().get_queryset()
        start = self.request.query_params.get("start")
        end = self.request.query_params.get("end")
        if start:
            queryset = queryset.filter(occurred_at__date__gte=start)
        if end:
            queryset = queryset.filter(occurred_at__date__lte=end)
        return queryset

    def list(self, request, *args, **kwargs):
        errors = _date_filter_errors(request, ("start", "end"))
        if errors:
            return Response(errors, status=400)
        return super().list(request, *args, **kwargs)

    @extend_schema(
        request=RepairPartUsageCreateSerializer,
        responses=RepairPartUsageSerializer,
        description="读取或记录一条故障维修用件；用件记录只允许创建和读取。",
    )
    @action(detail=True, methods=["get", "post"], url_path="part-usages")
    def part_usages(self, request, pk=None):
        fault = self.get_object()
        if request.method == "GET":
            queryset = fault.part_usages.select_related(
                "spare_part",
                "spare_stock__data_center",
                "spare_stock__server_room",
                "operator",
            ).order_by("-created_at", "-id")
            page = self.paginate_queryset(queryset)
            if page is not None:
                return self.get_paginated_response(RepairPartUsageSerializer(page, many=True).data)
            return Response(RepairPartUsageSerializer(queryset, many=True).data)

        input_serializer = RepairPartUsageCreateSerializer(data=request.data)
        input_serializer.is_valid(raise_exception=True)
        usage = create_repair_part_usage(
            fault_id=fault.pk,
            validated_data=input_serializer.validated_data,
            operator=request.user,
            request=request,
        )
        return Response(RepairPartUsageSerializer(usage).data, status=201)

    @transaction.atomic
    def perform_create(self, serializer):
        super().perform_create(serializer)
        fault = serializer.instance
        sync_asset_fault_status(fault.asset_id, request=self.request, fault_id=fault.pk)

    @transaction.atomic
    def perform_update(self, serializer):
        # Validation initially happens before perform_update. Re-lock and
        # validate again so a concurrent completion cannot turn this into a
        # historical write after the first validation pass.
        locked_fault = FaultEvent.objects.select_for_update().get(pk=serializer.instance.pk)
        locked_serializer = self.get_serializer(
            locked_fault,
            data=self.request.data,
            partial=self.request.method == "PATCH",
        )
        locked_serializer.is_valid(raise_exception=True)
        previous_asset_id = locked_fault.asset_id
        super().perform_update(locked_serializer)
        serializer.instance = locked_serializer.instance
        fault = locked_serializer.instance
        sync_asset_fault_status(previous_asset_id, request=self.request)
        if fault.asset_id != previous_asset_id:
            sync_asset_fault_status(fault.asset_id, request=self.request, fault_id=fault.pk)

    @transaction.atomic
    def perform_destroy(self, instance):
        locked_fault = FaultEvent.objects.select_for_update().get(pk=instance.pk)
        if locked_fault.is_closed:
            raise DRFValidationError({
                "detail": "已关闭故障不可删除。",
                "code": "closed_fault_immutable",
            })
        if locked_fault.part_usages.exists():
            raise DRFValidationError("故障存在维修用件记录，不能删除。")
        asset_id = locked_fault.asset_id
        super().perform_destroy(locked_fault)
        sync_asset_fault_status(asset_id, request=self.request)


class RepairRecordViewSet(AuditedModelViewSetMixin, viewsets.ModelViewSet):
    queryset = RepairRecord.objects.select_related("fault", "fault__asset")
    serializer_class = RepairRecordSerializer
    filter_backends = [DjangoFilterBackend, OrderingFilter]
    filterset_fields = ["fault"]
    ordering_fields = ["created_at", "started_at", "finished_at"]
    ordering = ["-created_at"]
    permission_classes = [BusinessRolePermission]
    permission_resource = "faults"
    audit_resource = "repair_record"

    @transaction.atomic
    def perform_create(self, serializer):
        super().perform_create(serializer)
        repair = serializer.instance
        sync_repair_completion(repair, request=self.request)

    @transaction.atomic
    def perform_update(self, serializer):
        # Revalidate after locking so a concurrent finish cannot be followed
        # by an ordinary update to completed historical facts.
        locked_repair = RepairRecord.objects.select_for_update().get(pk=serializer.instance.pk)
        locked_serializer = self.get_serializer(
            locked_repair,
            data=self.request.data,
            partial=self.request.method == "PATCH",
        )
        locked_serializer.is_valid(raise_exception=True)
        super().perform_update(locked_serializer)
        serializer.instance = locked_serializer.instance
        repair = locked_serializer.instance
        sync_repair_completion(repair, request=self.request)

    @transaction.atomic
    def perform_destroy(self, instance):
        locked_repair = RepairRecord.objects.select_for_update().get(pk=instance.pk)
        if locked_repair.finished_at is not None:
            raise DRFValidationError({
                "detail": "已完成维修不可删除，请使用重新打开操作。",
                "code": "finished_repair_immutable",
            })
        fault_id = locked_repair.fault_id
        repair_id = locked_repair.pk
        super().perform_destroy(locked_repair)
        sync_fault_completion(
            fault_id=fault_id,
            finished_at=None,
            request=self.request,
            repair_id=repair_id,
        )

    @extend_schema(
        responses=RepairRecordSerializer,
        description="重新打开已完成维修；保留维修记录和用件历史，并按现有生命周期规则重开故障。",
    )
    @action(detail=True, methods=["post"], url_path="reopen")
    def reopen(self, request, pk=None):
        repair = reopen_repair(repair_id=self.get_object().pk, request=request)
        return Response(RepairRecordSerializer(repair).data)


def _inventory_snapshot(asset):
    allocation = getattr(asset, "rack_allocation", None)
    network = {
        item.role: item.address
        for item in getattr(asset, "network_addresses", []).all()
    }
    if allocation:
        location = {
            "data_center_id": allocation.rack.room.data_center_id,
            "data_center": allocation.rack.room.data_center.name,
            "server_room_id": allocation.rack.room_id,
            "server_room": allocation.rack.room.name,
            "rack_id": allocation.rack_id,
            "rack_code": allocation.rack.code,
            "start_u": allocation.start_u,
            "end_u": allocation.end_u,
        }
    else:
        location = {
            "data_center_id": asset.asset_data_center_id,
            "data_center": asset.asset_data_center.name if asset.asset_data_center_id else "",
            "server_room_id": None,
            "server_room": "",
            "rack_id": None,
            "rack_code": "",
            "start_u": None,
            "end_u": None,
        }
    return {
        "asset_no": asset.asset_no,
        "name": asset.name,
        "device_type_name": asset.device_type.name if asset.device_type_id else "",
        "serial_number": asset.serial_number,
        "status": asset.status,
        "business_ip": network.get("business", ""),
        "management_ip": network.get("management", ""),
        "oob_ip": network.get("oob", ""),
        **location,
    }


def _inventory_items_queryset(task, request):
    queryset = InventoryItem.objects.filter(task=task).select_related(
        "asset", "checked_by", "resolved_by", "actual_rack__room__data_center"
    ).order_by("asset__asset_no", "id")
    search = request.query_params.get("search", "").strip()
    if search:
        queryset = queryset.filter(
            Q(asset__asset_no__icontains=search)
            | Q(asset__name__icontains=search)
            | Q(asset__serial_number__icontains=search)
            | Q(asset__network_addresses__address__icontains=search)
        ).distinct()
    status = request.query_params.get("status", "").strip()
    if status:
        queryset = queryset.filter(status=status)
    resolution_status = request.query_params.get("resolution_status", "").strip()
    if resolution_status:
        queryset = queryset.filter(resolution_status=resolution_status)
    return queryset


class InventoryTaskViewSet(AuditedModelViewSetMixin, viewsets.ModelViewSet):
    queryset = InventoryTask.objects.select_related(
        "data_center", "server_room", "inspector"
    ).prefetch_related("items")
    serializer_class = InventoryTaskSerializer
    permission_classes = [BusinessRolePermission]
    permission_resource = "inventory"
    audit_resource = "inventory_task"
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ["status", "data_center", "server_room", "inspector"]
    search_fields = ["name", "data_center__name", "server_room__name", "inspector__username"]
    ordering_fields = ["created_at", "start_at", "end_at", "name"]
    ordering = ["-created_at", "-id"]

    def get_permissions(self):
        if self.action == "scope_preview":
            return [CanManageInventory()]
        if self.action == "export":
            return [CanExportInventory()]
        if self.action in {"complete", "reopen"}:
            return [CanManageInventory()]
        return super().get_permissions()

    @transaction.atomic
    def perform_create(self, serializer):
        task = serializer.save(inspector=serializer.validated_data.get("inspector") or self.request.user)
        assets = get_inventory_scope_assets(task.data_center, task.server_room)
        InventoryItem.objects.bulk_create([
            InventoryItem(task=task, asset=asset, system_snapshot=_inventory_snapshot(asset))
            for asset in assets
        ])
        write_audit_log(
            self.request,
            action="create",
            resource_type=self.audit_resource,
            resource_id=task.pk,
            after={"task": model_snapshot(task), "items_created": len(assets)},
        )

    @transaction.atomic
    def perform_destroy(self, instance):
        items = list(InventoryItem.objects.filter(task_id=instance.pk).select_for_update().order_by("pk").only(
            "status",
            "resolution_status",
            "resolution_action",
            "resolution_note",
            "resolved_by",
            "resolved_at",
        ))
        task = InventoryTask.objects.select_for_update().get(pk=instance.pk)
        block_reason = inventory_task_delete_block_reason(task, items)
        if block_reason:
            raise DRFValidationError({"detail": block_reason})
        super().perform_destroy(task)

    @extend_schema(
        parameters=[InventoryScopePreviewQuerySerializer],
        responses=InventoryScopePreviewSerializer,
        description="预览创建盘点任务时将生成的资产范围，不创建任务或盘点明细。",
    )
    @action(detail=False, methods=["get"], url_path="scope-preview")
    def scope_preview(self, request):
        query_params = {
            "data_center": request.query_params.get("data_center"),
            "server_room": request.query_params.get("server_room") or None,
        }
        serializer = InventoryScopePreviewQuerySerializer(data=query_params)
        serializer.is_valid(raise_exception=True)
        data_center = serializer.validated_data["data_center"]
        server_room = serializer.validated_data.get("server_room")
        assets = get_inventory_scope_assets(data_center, server_room)

        total = assets.count()
        racked = assets.filter(rack_allocation__isnull=False).count()
        unracked = assets.filter(rack_allocation__isnull=True).count()
        retired = assets.filter(status="retired").count()
        warnings = []
        if unracked:
            warnings.append(f"当前范围包含 {unracked} 台未上架资产")
        if retired:
            warnings.append(f"当前范围包含 {retired} 台已报废资产")

        payload = {
            "data_center": {"id": data_center.id, "name": data_center.name},
            "server_room": (
                {"id": server_room.id, "name": server_room.name}
                if server_room is not None
                else None
            ),
            "scope_label": (
                f"{data_center.name} / {server_room.name}"
                if server_room is not None
                else f"{data_center.name} / 整个数据中心"
            ),
            "total": total,
            "racked": racked,
            "unracked": unracked,
            "retired": retired,
            "includes_unracked": server_room is None,
            "warnings": warnings,
        }
        return Response(InventoryScopePreviewSerializer(payload).data)

    @action(detail=True, methods=["get"], url_path="items")
    def items(self, request, pk=None):
        # ``get_object`` applies the list search backends to query parameters.
        # The item search belongs to the nested item queryset, so resolving the
        # task from the base queryset prevents ``?search=SN...`` from filtering
        # the task itself to a false 404.
        task = get_object_or_404(InventoryTask, pk=pk)
        queryset = _inventory_items_queryset(task, request)
        page = self.paginate_queryset(queryset)
        serializer = InventoryItemSerializer(page or queryset, many=True)
        if page is not None:
            return self.get_paginated_response(serializer.data)
        return Response(serializer.data)

    @action(detail=True, methods=["post"], url_path="complete")
    @transaction.atomic
    def complete(self, request, pk=None):
        task = InventoryTask.objects.select_for_update().get(pk=self.get_object().pk)
        if task.status == "completed":
            return Response(InventoryTaskSerializer(task).data)
        items = list(
            InventoryItem.objects.select_for_update()
            .filter(task_id=task.pk)
            .order_by("pk")
            .only("status")
        )
        pending = sum(item.status == "pending" for item in items)
        if pending:
            return Response({"detail": f"还有 {pending} 台设备未盘点，不能完成任务"}, status=400)
        task.status = "completed"
        task.completed_at = timezone.now()
        task.save(update_fields=["status", "completed_at", "updated_at"])
        write_audit_log(
            request,
            action="complete",
            resource_type=self.audit_resource,
            resource_id=task.pk,
            after={"status": task.status, "completed_at": task.completed_at},
        )
        return Response(InventoryTaskSerializer(task).data)

    @action(detail=True, methods=["post"], url_path="reopen")
    @transaction.atomic
    def reopen(self, request, pk=None):
        task = InventoryTask.objects.select_for_update().get(pk=self.get_object().pk)
        task.status = "in_progress"
        task.completed_at = None
        task.save(update_fields=["status", "completed_at", "updated_at"])
        write_audit_log(
            request,
            action="reopen",
            resource_type=self.audit_resource,
            resource_id=task.pk,
            after={"status": task.status},
        )
        return Response(InventoryTaskSerializer(task).data)

    @extend_schema(
        parameters=[
            OpenApiParameter(name="search", type=OpenApiTypes.STR, required=False),
            OpenApiParameter(name="status", type=OpenApiTypes.STR, required=False, enum=INVENTORY_ITEM_STATUS_VALUES),
            OpenApiParameter(name="resolution_status", type=OpenApiTypes.STR, required=False, enum=INVENTORY_RESOLUTION_STATUS_VALUES),
        ],
        responses=OpenApiTypes.BINARY,
        description="导出当前盘点任务在当前明细筛选条件下的完整结果集，不受分页参数影响。",
    )
    @action(detail=True, methods=["get"], url_path="export")
    def export(self, request, pk=None):
        task = get_object_or_404(
            InventoryTask.objects.select_related("data_center", "server_room", "inspector"),
            pk=pk,
        )
        items = _inventory_items_queryset(task, request)
        limit_response = _export_limit_response(items, "盘点结果导出")
        if limit_response:
            return limit_response
        book = Workbook()
        sheet = book.active
        sheet.title = "盘点结果"
        headers = [
            "盘点名称", "数据中心", "机房", "任务盘点人", "资产编号", "资产名称", "序列号", "设备类型", "业务 IP", "管理 IP", "带外 IP",
            "系统机柜", "系统 U 位", "盘点结果", "实际数据中心", "实际机房", "实际机柜", "实际 U 位",
            "盘点时间", "实际盘点人", "备注", "处理状态", "处理方式", "处理人", "处理时间", "处理备注",
        ]
        _append_excel_row(sheet, headers)
        for item in items:
            snapshot = item.system_snapshot or {}
            system_u = ""
            if snapshot.get("start_u") is not None:
                system_u = f"U{snapshot.get('start_u')}–U{snapshot.get('end_u')}"
            actual_u = ""
            if item.actual_start_u is not None:
                actual_u = f"U{item.actual_start_u}–U{item.actual_end_u}"
            _append_excel_row(sheet, [
                task.name, task.data_center.name, task.server_room.name if task.server_room_id else "整个数据中心",
                task.inspector.get_full_name() or task.inspector.username, item.asset.asset_no, item.asset.name,
                item.asset.serial_number or "", snapshot.get("device_type_name", ""),
                snapshot.get("business_ip", ""), snapshot.get("management_ip", ""), snapshot.get("oob_ip", ""),
                snapshot.get("rack_code", ""), system_u,
                INVENTORY_ITEM_STATUS_LABELS.get(item.status, item.status),
                item.actual_rack.room.data_center.name if item.actual_rack_id else "",
                item.actual_rack.room.name if item.actual_rack_id else "",
                item.actual_rack.code if item.actual_rack_id else "", actual_u,
                system_localtime(item.checked_at).replace(tzinfo=None) if item.checked_at else "",
                item.checked_by.get_full_name() or item.checked_by.username if item.checked_by_id else "",
                item.notes,
                INVENTORY_RESOLUTION_STATUS_LABELS.get(item.resolution_status, item.resolution_status),
                INVENTORY_RESOLUTION_ACTION_LABELS.get(item.resolution_action, "") if item.resolution_action else "",
                item.resolved_by.get_full_name() or item.resolved_by.username if item.resolved_by_id else "",
                system_localtime(item.resolved_at).replace(tzinfo=None) if item.resolved_at else "",
                item.resolution_note,
            ])
        _style_export_sheet(sheet, max_width=32)
        return _xlsx_response(book, f"盘点结果_{_export_timestamp()}.xlsx")


def _inventory_error_message(error, fallback):
    detail = getattr(error, "detail", error)

    def flatten(value):
        if isinstance(value, dict):
            return "；".join(flatten(child) for child in value.values())
        if isinstance(value, (list, tuple)):
            return "；".join(flatten(child) for child in value)
        return str(value or "")

    return flatten(detail) or fallback


def _inventory_resolution_error_message(error):
    return _inventory_error_message(error, "盘点异常处理失败")


class InventoryItemViewSet(viewsets.ModelViewSet):
    queryset = InventoryItem.objects.select_related(
        "task", "asset", "checked_by", "resolved_by", "actual_rack__room__data_center"
    )
    serializer_class = InventoryItemSerializer
    permission_classes = [BusinessRolePermission]
    permission_resource = "inventory"
    http_method_names = ["get", "post", "patch", "head", "options"]
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ["task", "status", "resolution_status"]

    def create(self, request, *args, **kwargs):
        raise MethodNotAllowed("POST")

    @transaction.atomic
    def perform_update(self, serializer):
        # Validation initially happens before perform_update. Lock in the same
        # task -> item order as completion, then validate the submitted result
        # again so completion cannot race with this write.
        task = InventoryTask.objects.select_for_update().get(pk=serializer.instance.task_id)
        item = self.get_queryset().select_for_update().get(pk=serializer.instance.pk)
        locked_serializer = self.get_serializer(
            item,
            data=self.request.data,
            partial=True,
        )
        locked_serializer.is_valid(raise_exception=True)
        if task.status == "completed":
            raise DRFValidationError({"detail": "已完成的盘点任务已锁定，不能修改"})
        if locked_serializer.validated_data.get("status", item.status) == "normal":
            updated = confirm_inventory_item_normal(
                item_id=item.pk,
                actor=self.request.user,
                request=self.request,
                notes=locked_serializer.validated_data.get("notes"),
                source="inventory_result",
            )
            serializer.instance = updated
            return
        before_result = (
            item.status,
            item.actual_rack_id,
            item.actual_start_u,
            item.actual_end_u,
            item.notes,
        )
        before = InventoryItemSerializer(item).data
        updated = locked_serializer.save()
        if updated.status == "pending":
            updated.checked_by = None
            updated.checked_at = None
        else:
            updated.checked_by = self.request.user
            updated.checked_at = timezone.now()
        after_result = (
            updated.status,
            updated.actual_rack_id,
            updated.actual_start_u,
            updated.actual_end_u,
            updated.notes,
        )
        resolution_reset = updated.status in {"pending", "normal"} or before_result != after_result
        if resolution_reset:
            reset_inventory_resolution(updated)
        update_fields = ["checked_by", "checked_at", "updated_at"]
        if resolution_reset:
            update_fields.extend([
                "resolution_status",
                "resolution_action",
                "resolution_note",
                "resolved_by",
                "resolved_at",
            ])
        updated.save(update_fields=update_fields)
        write_audit_log(
            self.request,
            action="update",
            resource_type="inventory_item",
            resource_id=updated.pk,
            before=before,
            after=InventoryItemSerializer(updated).data,
        )
        serializer.instance = updated

    @extend_schema(
        request=InventoryResolutionSerializer,
        responses=InventoryItemSerializer,
        description="处理盘点异常；盘点结果本身在任务完成后仍保持锁定。",
    )
    @action(detail=True, methods=["post"], url_path="resolve")
    def resolve(self, request, pk=None):
        request_serializer = InventoryResolutionSerializer(data=request.data)
        request_serializer.is_valid(raise_exception=True)
        action_name = request_serializer.validated_data["action"]
        note = request_serializer.validated_data.get("note", "")
        if action_name == "update_asset" and not user_has_capability(request.user, "assets.manage"):
            raise PermissionDenied("更新资产台账需要 assets.manage 权限")
        item = self.get_object()
        resolved_item = resolve_inventory_item(
            item_id=item.pk,
            action=action_name,
            note=note,
            actor=request.user,
            request=request,
        )
        return Response(InventoryItemSerializer(resolved_item).data)

    @extend_schema(
        request=InventoryBulkResolutionSerializer,
        responses=InventoryBulkResolutionResponseSerializer,
        description="批量处理当前盘点任务中的盘点异常；不支持批量更新资产台账。",
    )
    @action(detail=False, methods=["post"], url_path="bulk-resolve")
    def bulk_resolve(self, request):
        request_serializer = InventoryBulkResolutionSerializer(data=request.data)
        request_serializer.is_valid(raise_exception=True)
        item_ids = request_serializer.validated_data["item_ids"]
        action_name = request_serializer.validated_data["action"]
        note = request_serializer.validated_data.get("note", "")

        items = list(
            InventoryItem.objects.filter(pk__in=item_ids).select_related("task", "asset")
        )
        items_by_id = {item.pk: item for item in items}
        missing_ids = [item_id for item_id in item_ids if item_id not in items_by_id]
        if missing_ids:
            raise DRFValidationError({"item_ids": f"盘点项不存在：{', '.join(map(str, missing_ids))}"})
        task_ids = {item.task_id for item in items}
        if len(task_ids) != 1:
            raise DRFValidationError({"item_ids": "所有盘点项必须属于同一个盘点任务"})

        results = []
        succeeded = 0
        for item_id in item_ids:
            item = items_by_id[item_id]
            try:
                resolved_item = resolve_inventory_item(
                    item_id=item_id,
                    action=action_name,
                    note=note,
                    actor=request.user,
                    request=request,
                )
            except (DRFValidationError, PermissionDenied) as exc:
                results.append({
                    "item_id": item_id,
                    "asset_no": item.asset.asset_no,
                    "success": False,
                    "reason": _inventory_resolution_error_message(exc),
                })
            except InventoryItem.DoesNotExist:
                results.append({
                    "item_id": item_id,
                    "asset_no": item.asset.asset_no,
                    "success": False,
                    "reason": "盘点项不存在",
                })
            else:
                succeeded += 1
                results.append({
                    "item_id": item_id,
                    "asset_no": resolved_item.asset.asset_no,
                    "success": True,
                    "reason": "",
                })

        response_data = {
            "requested": len(item_ids),
            "succeeded": succeeded,
            "failed": len(item_ids) - succeeded,
            "results": results,
        }
        return Response(InventoryBulkResolutionResponseSerializer(response_data).data)

    @extend_schema(
        request=InventoryBulkNormalSerializer,
        responses=InventoryBulkNormalResponseSerializer,
        description="批量确认当前进行中盘点任务中的未盘点项为正常；位置由任务创建时的系统快照确定。",
    )
    @action(detail=False, methods=["post"], url_path="bulk-confirm-normal")
    def bulk_confirm_normal(self, request):
        request_serializer = InventoryBulkNormalSerializer(data=request.data)
        request_serializer.is_valid(raise_exception=True)
        item_ids = request_serializer.validated_data["item_ids"]

        items = list(
            InventoryItem.objects.filter(pk__in=item_ids).select_related("task", "asset")
        )
        items_by_id = {item.pk: item for item in items}
        missing_ids = [item_id for item_id in item_ids if item_id not in items_by_id]
        if missing_ids:
            raise DRFValidationError({"item_ids": f"盘点项不存在：{', '.join(map(str, missing_ids))}"})
        task_ids = {item.task_id for item in items}
        if len(task_ids) != 1:
            raise DRFValidationError({"item_ids": "所有盘点项必须属于同一个盘点任务"})

        results = []
        succeeded = 0
        for item_id in item_ids:
            item = items_by_id[item_id]
            try:
                confirmed_item = confirm_inventory_item_normal(
                    item_id=item_id,
                    actor=request.user,
                    request=request,
                    require_pending=True,
                    source="inventory_bulk_confirm_normal",
                )
            except (DRFValidationError, PermissionDenied) as exc:
                results.append({
                    "item_id": item_id,
                    "asset_no": item.asset.asset_no,
                    "success": False,
                    "reason": _inventory_error_message(exc, "批量标记正常失败"),
                })
            except InventoryItem.DoesNotExist:
                results.append({
                    "item_id": item_id,
                    "asset_no": item.asset.asset_no,
                    "success": False,
                    "reason": "盘点项不存在",
                })
            else:
                succeeded += 1
                results.append({
                    "item_id": item_id,
                    "asset_no": confirmed_item.asset.asset_no,
                    "success": True,
                    "reason": "",
                })

        response_data = {
            "requested": len(item_ids),
            "succeeded": succeeded,
            "failed": len(item_ids) - succeeded,
            "results": results,
        }
        return Response(InventoryBulkNormalResponseSerializer(response_data).data)


@extend_schema(
    parameters=[
        OpenApiParameter(name="page", type=OpenApiTypes.INT, required=False),
        OpenApiParameter(name="page_size", type=OpenApiTypes.INT, required=False, description="每页数量，最大 100。"),
    ],
    responses=InventoryItemPageSerializer,
)
@api_view(["GET"])
@permission_classes([CanViewInventory])
def asset_inventory_records(request, pk):
    if not Asset.objects.filter(pk=pk).exists():
        from rest_framework.exceptions import NotFound
        raise NotFound("资产不存在")
    records = InventoryItem.objects.filter(asset_id=pk).select_related(
        "asset",
        "asset__device_type",
        "task",
        "task__data_center",
        "task__server_room",
        "checked_by",
        "resolved_by",
        "actual_rack__room__data_center",
    ).order_by("asset__asset_no", "id")
    paginator = StandardPagination()
    page = paginator.paginate_queryset(records, request)
    return paginator.get_paginated_response(InventoryItemSerializer(page, many=True).data)


@extend_schema(responses=InventoryInspectorSerializer(many=True))
@api_view(["GET"])
@permission_classes([CanManageInventory])
def inventory_inspectors(request):
    users = User.objects.filter(is_active=True).order_by("username")
    return Response([
        {"id": user.id, "username": user.username, "display_name": user.get_full_name() or user.username}
        for user in users
    ])


class AuditLogViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = AuditLog.objects.select_related("actor").order_by("-created_at", "-id")
    serializer_class = AuditLogSerializer
    permission_classes = [CanViewAuditLog]
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ["actor", "resource_type", "action", "resource_id"]
    search_fields = ["actor__username", "actor__first_name", "actor__last_name", "resource_type", "resource_id"]
    ordering_fields = ["created_at"]

    def get_queryset(self):
        queryset = super().get_queryset()
        start = self.request.query_params.get("start")
        end = self.request.query_params.get("end")
        if start:
            try:
                start = date.fromisoformat(start)
            except ValueError:
                return queryset.none()
            queryset = queryset.filter(created_at__date__gte=start)
        if end:
            try:
                end = date.fromisoformat(end)
            except ValueError:
                return queryset.none()
            queryset = queryset.filter(created_at__date__lte=end)
        return queryset

    def list(self, request, *args, **kwargs):
        errors = _date_filter_errors(request, ("start", "end"))
        if errors:
            return Response(errors, status=400)
        return super().list(request, *args, **kwargs)


LOGIN_USERNAME_MAX_LENGTH = User._meta.get_field(User.USERNAME_FIELD).max_length or 150


def _security_profile(user):
    profile, _ = UserSecurityProfile.objects.get_or_create(
        user=user,
        defaults={
            "must_change_password": False,
            "locale": get_system_settings().default_locale,
        },
    )
    return profile


def _login_ip(request):
    # REMOTE_ADDR is the only trusted value unless a deployment explicitly
    # adds a trusted proxy middleware in front of Django.
    return (request.META.get("REMOTE_ADDR") or "unknown")[:255]


def _login_account_key(username):
    return (username or "").strip().casefold()


def _auth_audit_extra(request, username, **extra):
    payload = {
        "username": username,
        "ip": _login_ip(request),
        "user_agent": (request.META.get("HTTP_USER_AGENT") or "")[:500],
    }
    payload.update(extra)
    return payload


def _throttle_state(scope, key, now):
    state, _ = AuthThrottleState.objects.select_for_update().get_or_create(
        scope=scope,
        key=key,
    )
    if state.locked_until and state.locked_until <= now:
        state.failure_count = 0
        state.first_failed_at = None
        state.locked_until = None
        state.save(update_fields=["failure_count", "first_failed_at", "locked_until", "updated_at"])
    return state


LDAP_DIAGNOSTIC_MAX_ATTEMPTS = 3
LDAP_DIAGNOSTIC_WINDOW_SECONDS = 60
LDAP_DIAGNOSTIC_LOCK_SECONDS = 60


def _ldap_diagnostic_throttle_key(request) -> str:
    return f"ldap-diagnostic:{request.user.pk}:{_login_ip(request)}"[:255]


def _consume_ldap_diagnostic_slot(request) -> tuple[bool, int]:
    now = timezone.now()
    window = timedelta(seconds=LDAP_DIAGNOSTIC_WINDOW_SECONDS)
    lock_duration = timedelta(seconds=LDAP_DIAGNOSTIC_LOCK_SECONDS)
    with transaction.atomic():
        state = _throttle_state("ip", _ldap_diagnostic_throttle_key(request), now)
        if state.first_failed_at and now - state.first_failed_at > window:
            state.failure_count = 0
            state.first_failed_at = None
            state.locked_until = None
        if state.locked_until and state.locked_until > now:
            return False, max(1, int((state.locked_until - now).total_seconds()))
        if state.failure_count >= LDAP_DIAGNOSTIC_MAX_ATTEMPTS:
            state.locked_until = now + lock_duration
            state.save(update_fields=["locked_until", "updated_at"])
            return False, LDAP_DIAGNOSTIC_LOCK_SECONDS
        if state.first_failed_at is None:
            state.first_failed_at = now
        state.failure_count += 1
        state.save(update_fields=["failure_count", "first_failed_at", "locked_until", "updated_at"])
    return True, 0


def _login_lock_status(username, ip):
    now = timezone.now()
    with transaction.atomic():
        states = [_throttle_state("account", _login_account_key(username), now), _throttle_state("ip", ip, now)]
        locked_until = max(
            (state.locked_until for state in states if state.locked_until and state.locked_until > now),
            default=None,
        )
    if locked_until:
        return True, max(1, int((locked_until - now).total_seconds()))
    return False, 0


def _register_login_failure(
    request,
    username,
    actor=None,
    *,
    reason="invalid_credentials",
    auth_source=None,
):
    now = timezone.now()
    ip = _login_ip(request)
    policy = get_local_account_security_policy()
    window = timedelta(seconds=max(1, policy["login_window_seconds"]))
    lock_duration = timedelta(seconds=max(1, policy["login_lock_seconds"]))
    max_attempts = max(1, policy["login_max_attempts"])
    with transaction.atomic():
        states = [_throttle_state("account", _login_account_key(username), now), _throttle_state("ip", ip, now)]
        locked_until = None
        for state in states:
            if not state.first_failed_at or now - state.first_failed_at > window:
                state.failure_count = 0
                state.first_failed_at = now
            state.failure_count += 1
            if state.failure_count >= max_attempts:
                state.locked_until = now + lock_duration
                locked_until = max(locked_until or now, state.locked_until)
            state.save(update_fields=["failure_count", "first_failed_at", "locked_until", "updated_at"])
        is_locked = bool(locked_until)
        retry_after = max(1, int((locked_until - now).total_seconds())) if locked_until else 0
        write_audit_log(
            request,
            action="login_locked" if is_locked else "login_failure",
            resource_type="auth_login",
            resource_id=username,
            actor=actor,
            extra=_auth_audit_extra(
                request,
                username,
                reason=reason,
                retry_after=retry_after,
                **({"auth_source": auth_source} if auth_source else {}),
            ),
        )
    return is_locked, retry_after


def _clear_login_throttle(username, ip):
    with transaction.atomic():
        AuthThrottleState.objects.select_for_update().filter(
            scope="account", key=_login_account_key(username),
        ).update(failure_count=0, first_failed_at=None, locked_until=None)
        AuthThrottleState.objects.select_for_update().filter(
            scope="ip", key=ip,
        ).update(failure_count=0, first_failed_at=None, locked_until=None)


def _auth_response(user):
    role_codes = user_role_codes(user)
    role_code = role_codes[0] if role_codes else None
    security_profile = _security_profile(user)
    directory_identity = DirectoryIdentity.objects.filter(user_id=user.pk).first()
    return {
        "username": user.username,
        "display_name": user.get_full_name() or user.username,
        "first_name": user.first_name,
        "last_name": user.last_name,
        "email": user.email,
        "is_active": user.is_active,
        "is_staff": user.is_staff,
        "is_superuser": user.is_superuser,
        "is_admin": user_has_capability(user, "organization.manage"),
        "role_code": role_code,
        "role_name": ROLE_DEFINITIONS.get(role_code, {}).get("name", ""),
        "roles": [
            {"code": code, "name": ROLE_DEFINITIONS[code]["name"]}
            for code in role_codes
        ],
        "permissions": user_capabilities(user),
        "auth_source": AUTH_SOURCE_LDAP if directory_identity is not None else AUTH_SOURCE_LOCAL,
        "directory_provider": directory_identity.provider if directory_identity is not None else None,
        "directory_login_identifier": (
            directory_identity.current_login_identifier if directory_identity is not None else None
        ),
        "directory_last_seen_at": (
            directory_identity.last_seen_at if directory_identity is not None else None
        ),
        "password_change_required": security_profile.must_change_password,
        "locale": security_profile.locale,
        "last_login": user.last_login,
    }


@extend_schema(request=OpenApiTypes.OBJECT, responses=OpenApiTypes.OBJECT)
@api_view(["POST"])
@permission_classes([AllowAny])
def auth_login(request):
    username = str(request.data.get("username", "")).strip()
    password = request.data.get("password", "")
    if not isinstance(password, str):
        password = ""
    if len(username) > LOGIN_USERNAME_MAX_LENGTH:
        return Response(
            {
                "detail": f"用户名不能超过 {LOGIN_USERNAME_MAX_LENGTH} 个字符",
                "code": "invalid_username",
            },
            status=400,
        )
    ip = _login_ip(request)
    locked, retry_after = _login_lock_status(username, ip)
    matched_user = User.objects.filter(username__iexact=username).first() if username else None
    if locked:
        with transaction.atomic():
            write_audit_log(
                request,
                action="login_locked",
                resource_type="auth_login",
                resource_id=username or "unknown",
                actor=matched_user,
                extra=_auth_audit_extra(request, username, reason="throttle_locked", retry_after=retry_after),
            )
        return Response(
            {"detail": "登录失败次数过多，请稍后再试", "code": "login_locked", "retry_after": retry_after},
            status=429,
            headers={"Retry-After": str(retry_after)},
        )
    auth_source = (
        AUTH_SOURCE_LDAP
        if matched_user and DirectoryIdentity.objects.filter(user_id=matched_user.pk).exists()
        else AUTH_SOURCE_LOCAL
        if matched_user
        else AUTH_SOURCE_LDAP
        if ldap_is_enabled()
        else None
    )
    try:
        authentication = authenticate_with_source_routing(request, username, password)
    except AuthenticationFailure as exc:
        actor = exc.actor or matched_user
        source = exc.auth_source or auth_source
        if exc.infrastructure:
            write_audit_log(
                request,
                action="login_failure",
                resource_type="auth_login",
                resource_id=username or "unknown",
                actor=actor,
                extra=_auth_audit_extra(
                    request,
                    username,
                    reason=exc.reason,
                    **({"auth_source": source} if source else {}),
                ),
            )
            return Response(
                {
                    "detail": "认证服务暂时不可用",
                    "code": "authentication_service_unavailable",
                },
                status=503,
            )
        if not exc.throttle:
            write_audit_log(
                request,
                action="login_failure",
                resource_type="auth_login",
                resource_id=username or "unknown",
                actor=actor,
                extra=_auth_audit_extra(
                    request,
                    username,
                    reason=exc.reason,
                    **({"auth_source": source} if source else {}),
                ),
            )
            return Response({"detail": "用户名或密码错误"}, status=400)
        is_locked, retry_after = _register_login_failure(
            request,
            username,
            actor,
            reason=exc.reason,
            auth_source=source,
        )
        if is_locked:
            return Response(
                {"detail": "登录失败次数过多，请稍后再试", "code": "login_locked", "retry_after": retry_after},
                status=429,
                headers={"Retry-After": str(retry_after)},
            )
        return Response({"detail": "用户名或密码错误"}, status=400)
    user = authentication.user
    _clear_login_throttle(username, ip)
    if authentication.auth_source == AUTH_SOURCE_LOCAL:
        profile = _security_profile(user)
        if not profile.must_change_password and local_password_expired(user, profile=profile):
            profile.must_change_password = True
            profile.save(update_fields=["must_change_password", "updated_at"])
    if authentication.auth_source == AUTH_SOURCE_LDAP:
        login(request, user, backend=LDAP_MODEL_BACKEND)
    else:
        login(request, user)
    write_audit_log(
        request,
        action="login_success",
        resource_type="auth_login",
        resource_id=user.username,
        actor=user,
        extra=_auth_audit_extra(
            request,
            user.username,
            auth_source=authentication.auth_source,
        ),
    )
    return Response(_auth_response(user))


@extend_schema(responses=OpenApiTypes.OBJECT)
@api_view(["GET"])
@permission_classes([AllowAny])
def auth_csrf(request):
    return Response({"csrfToken": get_token(request)})


@extend_schema(responses=OpenApiTypes.OBJECT)
@api_view(["GET"])
@permission_classes([IsAuthenticated, IsSystemAdministrator])
def auth_ldap_status(request):
    return Response(_ldap_status_payload())


def _ldap_last_diagnostic_payload() -> dict[str, object | None]:
    audit = (
        AuditLog.objects.filter(resource_type="ldap", action="ldap_diagnostic")
        .order_by("-created_at", "-id")
        .first()
    )
    if audit is None:
        return {
            "last_diagnostic_at": None,
            "last_diagnostic_success": None,
            "last_diagnostic_code": None,
        }
    extra = (audit.payload or {}).get("extra")
    extra = extra if isinstance(extra, dict) else {}
    return {
        "last_diagnostic_at": audit.created_at,
        "last_diagnostic_success": bool(extra.get("success")),
        "last_diagnostic_code": extra.get("code") if isinstance(extra.get("code"), str) else None,
    }


def _ldap_status_payload() -> dict:
    return {**ldap_status_snapshot(), **_ldap_last_diagnostic_payload()}


@extend_schema(responses=OpenApiTypes.OBJECT)
@api_view(["GET", "PATCH"])
@permission_classes([IsAuthenticated, IsSystemAdministrator])
def auth_ldap_config(request):
    if request.method == "GET":
        return Response({
            **public_configuration(get_effective_ldap_configuration()),
            **_ldap_last_diagnostic_payload(),
        })

    serializer = LdapConfigurationUpdateSerializer(data=request.data, partial=True)
    if not serializer.is_valid():
        return Response(serializer.errors, status=400)

    current = get_effective_ldap_configuration()
    password_value = serializer.validated_data.get("bind_password", "")
    password_submitted = bool(password_value)
    candidate = merge_configuration(
        current,
        serializer.validated_data,
        password_submitted=password_submitted,
        password_value=password_value,
    )
    errors = configuration_errors(candidate, require_password=candidate.enabled)
    if candidate.enabled and errors:
        return Response(errors, status=400)

    try:
        save_configuration(
            candidate,
            password_submitted=password_submitted,
            password_value=password_value,
        )
    except ConfigurationIdentityError as exc:
        return Response(exc.errors, status=400)
    except ConfigurationSecretError as exc:
        return Response({"bind_password": str(exc)}, status=400)

    updated = get_effective_ldap_configuration()
    before_public = public_configuration(current)
    after_public = public_configuration(updated)
    changed_fields = [
        field
        for field in serializer.validated_data
        if field != "bind_password" and before_public.get(field) != after_public.get(field)
    ]
    if password_submitted:
        changed_fields.append("bind_password")
    write_audit_log(
        request,
        action="ldap_configuration_updated",
        resource_type="ldap",
        resource_id="configuration",
        extra={"changed_fields": sorted(set(changed_fields))},
    )
    if current.enabled != updated.enabled:
        write_audit_log(
            request,
            action="ldap_enabled" if updated.enabled else "ldap_disabled",
            resource_type="ldap",
            resource_id="configuration",
            extra={"enabled": updated.enabled},
        )
    if password_submitted:
        write_audit_log(
            request,
            action="ldap_bind_password_updated",
            resource_type="ldap",
            resource_id="configuration",
            extra={"configured": True},
        )
    return Response({
        **public_configuration(updated),
        **_ldap_last_diagnostic_payload(),
    })


def _safe_ldap_diagnostic_payload(result) -> dict:
    if hasattr(result, "as_dict"):
        candidate = result.as_dict()
    elif isinstance(result, dict):
        candidate = result
    else:
        candidate = {}

    allowed_stages = {"configuration", "connection", "tls", "service_bind", "search"}
    stage = candidate.get("stage") if candidate.get("stage") in allowed_stages else "configuration"
    success = bool(candidate.get("success"))
    checks = []
    for check in candidate.get("checks", []) if isinstance(candidate.get("checks"), list) else []:
        if not isinstance(check, dict):
            continue
        name = check.get("name")
        status = check.get("status")
        if name in LDAP_DIAGNOSTIC_CHECKS and status in {"success", "error", "disabled"}:
            checks.append({"name": name, "status": status})
    payload = {"success": success, "stage": stage, "checks": checks}
    if success:
        return payload
    code = candidate.get("code")
    if code not in LDAP_DIAGNOSTIC_MESSAGES:
        code = "unexpected_error"
    payload["code"] = code
    payload["message"] = LDAP_DIAGNOSTIC_MESSAGES[code]
    return payload


def _write_ldap_diagnostic_audit(request, payload: dict) -> None:
    write_audit_log(
        request,
        action="ldap_diagnostic",
        resource_type="ldap",
        resource_id="configuration",
        extra={
            "success": payload["success"],
            "stage": payload["stage"],
            **({"code": payload["code"]} if "code" in payload else {}),
        },
    )


@extend_schema(request=OpenApiTypes.OBJECT, responses=OpenApiTypes.OBJECT)
@api_view(["POST"])
@permission_classes([IsAuthenticated, IsSystemAdministrator])
def auth_ldap_diagnostics(request):
    allowed, retry_after = _consume_ldap_diagnostic_slot(request)
    if not allowed:
        write_audit_log(
            request,
            action="ldap_diagnostic",
            resource_type="ldap",
            resource_id="configuration",
            extra={"success": False, "stage": "throttled", "code": "throttled"},
        )
        return Response(
            {
                "detail": "LDAP 诊断请求过于频繁，请稍后再试",
                "code": "diagnostic_throttled",
                "retry_after": retry_after,
            },
            status=429,
            headers={"Retry-After": str(retry_after)},
        )

    try:
        staged_configuration = None
        if request.data:
            serializer = LdapConfigurationUpdateSerializer(data=request.data, partial=True)
            if not serializer.is_valid():
                _write_ldap_diagnostic_audit(
                    request,
                    {"success": False, "stage": "configuration", "code": "configuration_error"},
                )
                return Response(serializer.errors, status=400)
            current = get_effective_ldap_configuration()
            password_value = serializer.validated_data.get("bind_password", "")
            staged_configuration = merge_configuration(
                current,
                serializer.validated_data,
                password_submitted=bool(password_value),
                password_value=password_value,
            )
            staged_errors = configuration_errors(staged_configuration, require_password=True)
            if staged_errors:
                _write_ldap_diagnostic_audit(
                    request,
                    {"success": False, "stage": "configuration", "code": "configuration_error"},
                )
                return Response(staged_errors, status=400)
        result = (
            LDAPDirectoryClient(configuration=staged_configuration).diagnose(
                staged_configuration,
                allow_disabled=staged_configuration is not None,
            )
            if staged_configuration is not None
            else LDAPDirectoryClient().diagnose()
        )
        payload = _safe_ldap_diagnostic_payload(result)
    except Exception as exc:
        logger.error("LDAP diagnostic endpoint failed (%s)", type(exc).__name__)
        payload = _safe_ldap_diagnostic_payload({
            "success": False,
            "stage": "configuration",
            "code": "unexpected_error",
        })

    _write_ldap_diagnostic_audit(request, payload)
    if payload["success"]:
        return Response(payload)
    status = 200 if payload.get("code") == "disabled" else 400 if payload.get("code") == "configuration_error" else 503
    return Response(payload, status=status)


@extend_schema(request=CurrentUserProfileSerializer, responses=OpenApiTypes.OBJECT)
@api_view(["GET", "PATCH"])
@permission_classes([IsAuthenticated])
def auth_me(request):
    if request.method == "PATCH":
        serializer = CurrentUserProfileSerializer(
            data=request.data,
            partial=True,
        )
        serializer.is_valid(raise_exception=True)
        before = _user_audit_snapshot(request.user)
        with transaction.atomic():
            validated = serializer.validated_data
            user = request.user
            user_fields = [field for field in ("first_name", "last_name", "email") if field in validated]
            for field in user_fields:
                setattr(user, field, validated[field])
            if user_fields:
                user.save(update_fields=user_fields)
            if "locale" in validated:
                profile = _security_profile(user)
                profile.locale = validated["locale"]
                profile.save(update_fields=["locale", "updated_at"])
            write_audit_log(
                request,
                action="update",
                resource_type="user",
                resource_id=user.pk,
                before=before,
                after=_user_audit_snapshot(user),
                extra={"profile_update": True},
            )
        return Response(_auth_response(request.user))
    return Response(_auth_response(request.user))


@extend_schema(request=OpenApiTypes.OBJECT, responses=OpenApiTypes.OBJECT)
@api_view(["POST"])
@permission_classes([IsAuthenticated])
def auth_logout(request):
    logout(request)
    return Response({"ok": True})


@extend_schema(request=OpenApiTypes.OBJECT, responses=OpenApiTypes.OBJECT)
@api_view(["POST"])
@permission_classes([IsAuthenticated])
@transaction.atomic
def auth_change_password(request):
    if is_directory_managed(request.user):
        return Response(
            {
                "detail": "Directory-managed accounts must change passwords through the corporate directory.",
                "code": "directory_password_managed",
            },
            status=400,
        )
    old_password = request.data.get("old_password", "")
    new_password = request.data.get("new_password", "")
    confirm_password = request.data.get("confirm_password")
    if not request.user.check_password(old_password):
        return Response({"detail": "原密码错误"}, status=400)
    if not new_password:
        return Response({"new_password": ["新密码不能为空"]}, status=400)
    if confirm_password is not None and new_password != confirm_password:
        return Response({"confirm_password": ["两次输入的新密码不一致"]}, status=400)
    if request.user.check_password(new_password):
        return Response({"new_password": ["新密码不能与当前密码相同"]}, status=400)
    try:
        validate_local_password(new_password, user=request.user)
    except DjangoValidationError as exc:
        return Response({"new_password": list(exc.messages)}, status=400)
    request.user.set_password(new_password)
    request.user.save(update_fields=["password"])
    profile = _security_profile(request.user)
    profile.must_change_password = False
    profile.password_changed_at = timezone.now()
    profile.save(update_fields=["must_change_password", "password_changed_at", "updated_at"])
    write_audit_log(
        request,
        action="update",
        resource_type="user",
        resource_id=request.user.pk,
        extra={"password_change": True},
    )
    update_session_auth_hash(request, request.user)
    return Response({"ok": True, "password_change_required": False})


@extend_schema(request=SystemResetSerializer, responses=OpenApiTypes.OBJECT)
@api_view(["POST"])
@permission_classes([CanResetSystem])
def system_reset(request):
    serializer = SystemResetSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    return Response(reset_system(actor=request.user, request=request))


@extend_schema(request=SystemSettingsSerializer, responses=SystemSettingsSerializer)
@api_view(["GET", "PUT"])
@permission_classes([CanManageSystemSettings])
def system_settings(request):
    if request.method == "GET":
        return Response(SystemSettingsSerializer(get_system_settings()).data)

    try:
        with transaction.atomic():
            setting = get_system_settings()
            setting = SystemSetting.objects.select_for_update().get(pk=setting.pk)
            before = system_settings_snapshot(setting)
            serializer = SystemSettingsSerializer(setting, data=request.data, partial=True)
            serializer.is_valid(raise_exception=True)
            smtp_password = serializer.validated_data.pop("smtp_password", "")
            encrypted_smtp_password = (
                encrypt_secret(smtp_password, field_name="SMTP 密码")
                if smtp_password
                else None
            )
            setting = serializer.save()
            if encrypted_smtp_password is not None:
                setting.smtp_password_encrypted = encrypted_smtp_password
                setting.save(update_fields=["smtp_password_encrypted", "updated_at"])
            after = system_settings_snapshot(setting)
            changed_fields = [
                key for key in after
                if before.get(key) != after.get(key)
            ]
            if smtp_password:
                changed_fields.append("smtp_password")
            if changed_fields:
                write_audit_log(
                    request,
                    action="update",
                    resource_type="system_settings",
                    resource_id="system",
                    before=before,
                    after=after,
                    extra={"changed_fields": sorted(set(changed_fields))},
                )
    except ConfigurationSecretError as exc:
        return Response({"smtp_password": [str(exc)]}, status=400)
    return Response(SystemSettingsSerializer(setting).data)


@extend_schema(request=SmtpTestEmailSerializer, responses=OpenApiTypes.OBJECT)
@api_view(["POST"])
@permission_classes([CanManageSystemSettings])
def smtp_test_email(request):
    serializer = SmtpTestEmailSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    setting = get_system_settings()
    try:
        send_smtp_test_email(setting, serializer.validated_data["recipient"])
    except SmtpConfigurationError as exc:
        write_audit_log(
            request,
            action="smtp_test",
            resource_type="system_settings",
            resource_id="smtp",
            extra={"success": False, "code": exc.code},
        )
        status = 400 if exc.code in {
            "smtp_disabled",
            "smtp_incomplete",
            "password_missing",
            "secret_unavailable",
        } else 502
        return Response({"detail": exc.detail, "code": exc.code}, status=status)
    except Exception:
        write_audit_log(
            request,
            action="smtp_test",
            resource_type="system_settings",
            resource_id="smtp",
            extra={"success": False, "code": "delivery_failed"},
        )
        return Response(
            {
                "detail": "无法发送测试邮件，请检查 SMTP 配置",
                "code": "delivery_failed",
            },
            status=502,
        )
    write_audit_log(
        request,
        action="smtp_test",
        resource_type="system_settings",
        resource_id="smtp",
        extra={"success": True},
    )
    return Response({"ok": True, "detail": "测试邮件已发送"})


@extend_schema(responses=OpenApiTypes.BINARY)
@api_view(["GET"])
@permission_classes([CanImportAssets])
def asset_import_template(request):
    book = build_import_template()
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="asset-import-template.xlsx"'
    book.save(response)
    return response


@extend_schema(request=OpenApiTypes.OBJECT, responses=OpenApiTypes.OBJECT)
@api_view(["POST"])
@permission_classes([CanImportAssets])
@parser_classes([MultiPartParser, FormParser])
def asset_import_preview(request):
    try:
        return Response(AssetImportService.preview(request.FILES.get("file")))
    except ImportFileError as exc:
        return Response({"detail": str(exc)}, status=400)


@extend_schema(request=OpenApiTypes.OBJECT, responses=OpenApiTypes.OBJECT)
@api_view(["POST"])
@permission_classes([CanImportAssets])
@parser_classes([MultiPartParser, FormParser])
def asset_import(request):
    try:
        return Response(AssetImportService.commit(request.FILES.get("file"), request))
    except ImportFileError as exc:
        return Response({"detail": str(exc)}, status=400)
    except ImportValidationError as exc:
        status = 409 if exc.concurrent else 400
        message = "确认导入前数据已发生变化，请查看最新校验结果" if exc.concurrent else "导入文件存在异常，请先修正后再确认"
        return Response({"detail": message, "preview": exc.preview}, status=status)


@extend_schema(
    parameters=[
        OpenApiParameter(name="search", type=OpenApiTypes.STR, required=False),
        OpenApiParameter(name="status", type=OpenApiTypes.STR, required=False, enum=ASSET_STATUS_VALUES),
        OpenApiParameter(name="device_type", type=OpenApiTypes.INT, required=False),
        OpenApiParameter(name="manufacturer", type=OpenApiTypes.INT, required=False),
        OpenApiParameter(name="model", type=OpenApiTypes.STR, required=False),
        OpenApiParameter(name="data_center", type=OpenApiTypes.INT, required=False),
        OpenApiParameter(name="tags", type=OpenApiTypes.STR, required=False, description="标签 ID，支持逗号分隔。"),
        OpenApiParameter(name="warranty", type=OpenApiTypes.STR, required=False, enum=["within_30_days", "expired"]),
        OpenApiParameter(name="custom__{field_key}__{operator}", type=OpenApiTypes.STR, required=False),
    ],
    responses=OpenApiTypes.BINARY,
    description="导出当前资产台账筛选结果，不受分页参数影响。",
)
@api_view(["GET"])
@permission_classes([CanExportAssets])
def asset_export(request):
    as_of_date = system_localdate()
    queryset = _filtered_view_queryset(
        AssetViewSet,
        request,
        filter_backends=[DjangoFilterBackend, SearchFilter],
    ).order_by("asset_no")
    limit_response = _export_limit_response(queryset, "资产导出结果")
    if limit_response:
        return limit_response
    assets = list(queryset)
    custom_fields = list(
        CustomField.objects.filter(
            is_active=True,
        ).filter(
            Q(device_type__isnull=True) | Q(device_type__assets__in=assets)
        ).distinct().prefetch_related("options").order_by("device_type__name", "sort_order", "id")
    ) if assets else []
    headers = [
        "资产编号", "资产名称", "设备类型", "厂商", "型号", "厂商/型号", "序列号", "用途", "状态", "当前责任人", "部门",
        "数据中心", "机房", "机柜", "起始 U", "结束 U", "业务 IP", "管理 IP", "带外 IP", "采购日期",
        "供应商", "采购单号", "采购金额", "折旧方法", "折旧起算日", "折旧年限", "残值率", "资产原值", "预计残值", "月折旧额", "累计折旧", "当前净值", "折旧状态",
        "维保厂商", "维保合同号", "维保开始日", "维保到期日", "维保备注", "备注", "标签",
    ]
    headers.extend([field.name or field.key for field in custom_fields])
    status_labels = ASSET_STATUS_LABELS
    book = Workbook()
    sheet = book.active
    sheet.title = "资产台账"
    _append_excel_row(sheet, headers)
    for asset in assets:
        networks = {item.role: item.address for item in asset.network_addresses.all()}
        rack = getattr(asset, "rack_allocation", None)
        procurement = next(iter(asset.procurement_records.all()), None)
        maintenance = next(iter(asset.maintenance_contracts.all()), None)
        depreciation = calculate_asset_depreciation(asset, as_of_date=as_of_date)
        depreciation_method_labels = {"straight_line": "直线法"}
        custom_by_key = {}
        for item in asset.custom_values.all():
            field = item.field
            option_labels = {option.value: option.label for option in field.options.all()}
            if field.field_type in {"text", "textarea"}:
                value = item.text_value
            elif field.field_type == "number":
                value = item.number_value
            elif field.field_type == "date":
                value = item.date_value
            elif field.field_type == "boolean":
                value = "" if item.boolean_value is None else ("是" if item.boolean_value else "否")
            elif field.field_type == "select":
                value = option_labels.get(item.text_value, item.text_value)
            else:
                values = item.json_value if isinstance(item.json_value, list) else []
                value = "; ".join(option_labels.get(str(raw), str(raw)) for raw in values)
            custom_by_key[field.key] = value
        tag_text = ", ".join(item.tag.name for item in asset.asset_tags.all())
        row_values = [
            asset.asset_no, asset.name, asset.device_type.name if asset.device_type_id else "",
            asset.manufacturer.name if asset.manufacturer_id else "", asset.model or "", asset.manufacturer_model, asset.serial_number or "", asset.purpose, status_labels.get(asset.status, asset.status), _responsibility_user_name(asset.responsible_user),
            asset.department.name if asset.department_id else "",
            rack.rack.room.data_center.name if rack else (asset.asset_data_center.name if asset.asset_data_center_id else ""), rack.rack.room.name if rack else "", rack.rack.code if rack else "",
            rack.start_u if rack else "", rack.end_u if rack else "", networks.get("business", ""), networks.get("management", ""), networks.get("oob", ""),
            procurement.purchase_date if procurement else "", procurement.supplier if procurement else "", procurement.order_no if procurement else "", procurement.amount if procurement else "",
            depreciation_method_labels.get(depreciation["method"], depreciation["method"] or ""),
            asset.depreciation_start_date or "", asset.depreciation_years or "", _export_residual_rate(asset.residual_rate),
            _export_decimal(depreciation["original_value"]), _export_decimal(depreciation["residual_value"]),
            _export_decimal(depreciation["monthly_depreciation"]), _export_decimal(depreciation["accumulated_depreciation"]),
            _export_decimal(depreciation["net_book_value"]),
            {"unconfigured": "未配置", "not_started": "尚未开始", "depreciating": "折旧中", "fully_depreciated": "已折旧完"}.get(depreciation["status"], depreciation["status"]),
            maintenance.provider if maintenance else "", maintenance.contract_no if maintenance else "", maintenance.start_date if maintenance else "", maintenance.expiry_date if maintenance else "", maintenance.notes if maintenance else "", asset.notes,
            tag_text,
        ]
        row_values.extend(custom_by_key.get(field.key, "") for field in custom_fields)
        _append_excel_row(sheet, row_values)
    _style_export_sheet(sheet, max_width=36)
    return _xlsx_response(book, f"资产台账_{_export_timestamp()}.xlsx")


def _repair_queryset(request):
    return _filtered_view_queryset(
        FaultEventViewSet,
        request,
        filter_backends=[DjangoFilterBackend, SearchFilter],
    ).order_by("-occurred_at")


@extend_schema(
    parameters=[
        OpenApiParameter(name="search", type=OpenApiTypes.STR, required=False),
        OpenApiParameter(name="is_closed", type=OpenApiTypes.BOOL, required=False),
        OpenApiParameter(name="start", type=OpenApiTypes.DATE, required=False),
        OpenApiParameter(name="end", type=OpenApiTypes.DATE, required=False),
    ],
    responses=OpenApiTypes.BINARY,
)
@api_view(["GET"])
@permission_classes([CanExportFaults])
def repair_record_export(request):
    errors = _date_filter_errors(request, ("start", "end"))
    if errors:
        return Response(errors, status=400)
    book = Workbook()
    sheet = book.active
    sheet.title = "维修记录"
    queryset = _repair_queryset(request)
    limit_response = _export_limit_response(queryset, "故障维修导出结果")
    if limit_response:
        return limit_response
    _append_excel_row(sheet, [
        "资产编号", "资产名称", "故障发生时间", "故障原因", "故障描述", "是否关闭",
        "维修厂商", "维修开始时间", "维修完成时间", "维修费用", "维修备注",
    ])
    for fault in queryset:
        repair = getattr(fault, "repair", None)
        _append_excel_row(sheet, [
            fault.asset.asset_no, fault.asset.name, system_localtime(fault.occurred_at).replace(tzinfo=None),
            fault.reason, fault.description, "是" if fault.is_closed else "否",
            repair.provider if repair else "",
            system_localtime(repair.started_at).replace(tzinfo=None) if repair and repair.started_at else "",
            system_localtime(repair.finished_at).replace(tzinfo=None) if repair and repair.finished_at else "",
            repair.cost if repair and repair.cost is not None else "",
            repair.notes if repair else "",
        ])
    _style_export_sheet(sheet, max_width=40)
    return _xlsx_response(book, f"故障维修_{_export_timestamp()}.xlsx")


@extend_schema(responses=OpenApiTypes.OBJECT)
@api_view(["GET"])
@permission_classes([CanViewLicenses])
def license_summary(request):
    counts = license_status_counts(today=system_localdate())
    return Response(counts)


@extend_schema(
    parameters=[
        OpenApiParameter(name="search", type=OpenApiTypes.STR, required=False),
        OpenApiParameter(name="manufacturer", type=OpenApiTypes.INT, required=False),
        OpenApiParameter(
            name="status",
            type=OpenApiTypes.STR,
            required=False,
            enum=LICENSE_STATUS_KEYS,
        ),
    ],
    responses=OpenApiTypes.BINARY,
    description="导出当前软件许可筛选结果，不受分页参数影响。",
)
@api_view(["GET"])
@permission_classes([CanExportLicenses])
def license_export(request):
    queryset = _filtered_view_queryset(
        SoftwareLicenseViewSet,
        request,
        filter_backends=[DjangoFilterBackend, SearchFilter],
    ).order_by("expiry_date", "name", "id")
    limit_response = _export_limit_response(queryset, "软件许可导出结果")
    if limit_response:
        return limit_response

    status_labels = LICENSE_STATUS_LABELS
    book = Workbook()
    sheet = book.active
    sheet.title = "软件许可"
    _append_excel_row(sheet, [
        "软件名称", "厂商", "许可类型", "授权数量", "已用数量", "剩余数量", "状态", "到期日", "备注",
    ])
    today = system_localdate()
    for license_row in queryset:
        status = license_status_value(license_row, today=today)
        _append_excel_row(sheet, [
            license_row.name, license_row.manufacturer.name if license_row.manufacturer_id else "", license_row.license_type,
            license_row.authorized_count, license_row.used_count,
            license_row.authorized_count - license_row.used_count,
            status_labels[status], license_row.expiry_date, license_row.notes,
        ])
    _style_export_sheet(sheet, max_width=36)
    return _xlsx_response(book, f"软件许可_{_export_timestamp()}.xlsx")


def _spare_part_export_queryset(request):
    return _filtered_view_queryset(
        SparePartViewSet,
        request,
        filter_backends=[DjangoFilterBackend, SearchFilter],
    ).order_by("name", "category__name", "id")


@extend_schema(
    parameters=[
        OpenApiParameter(name="search", type=OpenApiTypes.STR, required=False),
        OpenApiParameter(name="category", type=OpenApiTypes.INT, required=False),
        OpenApiParameter(name="manufacturer", type=OpenApiTypes.INT, required=False),
        OpenApiParameter(name="data_center", type=OpenApiTypes.INT, required=False),
        OpenApiParameter(name="server_room", type=OpenApiTypes.INT, required=False),
    ],
    responses=OpenApiTypes.BINARY,
    description="导出当前备件主数据筛选结果及真实库存聚合，不受分页参数影响。",
)
@api_view(["GET"])
@permission_classes([CanExportSpares])
def spare_part_export(request):
    queryset = _spare_part_export_queryset(request)
    limit_response = _export_limit_response(queryset, "备件导出结果")
    if limit_response:
        return limit_response

    book = Workbook()
    sheet = book.active
    sheet.title = "备件管理"
    _append_excel_row(sheet, [
        "备件编码", "备件名称", "备件类型", "厂商", "型号", "规格", "计量单位", "库存数量", "安全库存", "库存地点数", "存放位置", "备注",
    ])
    for part in queryset:
        _append_excel_row(sheet, [
            part.code, part.name, part.category.name, part.manufacturer.name if part.manufacturer_id else "", part.model,
            part.specification, SPARE_UNIT_LABELS.get(part.unit, part.unit), part.total_quantity, part.safety_stock, part.location_count,
            part.storage_location, part.notes,
        ])
    _style_export_sheet(sheet, max_width=36)
    return _xlsx_response(book, f"备件管理_{_export_timestamp()}.xlsx")


@extend_schema(
    parameters=[
        OpenApiParameter(name="part", type=OpenApiTypes.INT, required=False),
        OpenApiParameter(name="operation_type", type=OpenApiTypes.STR, required=False, enum=STOCK_OPERATION_TYPE_VALUES),
        OpenApiParameter(name="source_data_center", type=OpenApiTypes.INT, required=False),
        OpenApiParameter(name="target_data_center", type=OpenApiTypes.INT, required=False),
        OpenApiParameter(name="search", type=OpenApiTypes.STR, required=False),
        OpenApiParameter(name="start", type=OpenApiTypes.DATE, required=False),
        OpenApiParameter(name="end", type=OpenApiTypes.DATE, required=False),
    ],
    responses=OpenApiTypes.BINARY,
    description="导出当前备件库存流水筛选结果，不受分页参数影响。",
)
@api_view(["GET"])
@permission_classes([CanExportSpares])
def spare_transaction_export(request):
    errors = _date_filter_errors(request, ("start", "end"))
    if errors:
        return Response(errors, status=400)
    queryset = _filtered_view_queryset(
        SpareStockTransactionViewSet,
        request,
        filter_backends=[DjangoFilterBackend, SearchFilter],
    ).order_by("-created_at", "-id")
    limit_response = _export_limit_response(queryset, "备件流水导出结果")
    if limit_response:
        return limit_response

    operation_labels = STOCK_OPERATION_TYPE_LABELS
    book = Workbook()
    sheet = book.active
    sheet.title = "库存流水"
    _append_excel_row(sheet, [
        "备件编码", "备件名称", "备件类型", "操作", "变化数量", "计量单位", "来源数据中心", "来源机房",
        "目标数据中心", "目标机房", "操作前库存", "操作后库存", "操作人", "参考单号/用途", "备注", "发生时间",
    ])
    for row in queryset:
        movement_quantity = row.quantity_delta
        _append_excel_row(sheet, [
            row.part.code, row.part.name, row.part.category.name, operation_labels.get(row.operation_type, row.operation_type),
            movement_quantity, SPARE_UNIT_LABELS.get(row.part.unit, row.part.unit),
            row.source_data_center.name if row.source_data_center_id else "",
            row.source_server_room.name if row.source_server_room_id else "",
            row.target_data_center.name if row.target_data_center_id else "",
            row.target_server_room.name if row.target_server_room_id else "",
            row.before_quantity, row.after_quantity,
            row.operator.get_full_name() or row.operator.username if row.operator_id else "已删除账号",
            row.reference, row.notes, system_localtime(row.created_at).replace(tzinfo=None),
        ])
    _style_export_sheet(sheet, max_width=40)
    return _xlsx_response(book, f"备件流水_{_export_timestamp()}.xlsx")


@extend_schema(responses=OpenApiTypes.OBJECT)
@api_view(["GET"])
@permission_classes([CanViewDashboard])
def dashboard_overview(request):
    try:
        scope = resolve_dashboard_scope(request)
    except DashboardScopeError as exc:
        return Response({"detail": str(exc)}, status=400)
    return Response(build_dashboard_payload(
        scope,
        include_faults=user_has_capability(request.user, "faults.view"),
        include_licenses=user_has_capability(request.user, "licenses.view"),
    ))


@extend_schema(responses=OpenApiTypes.OBJECT)
@api_view(["GET"])
@permission_classes([CanViewDashboard])
def alerts_overview(request):
    """Return actionable reminders assembled from existing business data."""
    return Response(build_alerts_payload(
        include_assets=user_has_capability(request.user, "assets.view"),
        include_licenses=user_has_capability(request.user, "licenses.view"),
        include_faults=user_has_capability(request.user, "faults.view"),
        include_inventory=user_has_capability(request.user, "inventory.view"),
        include_spares=user_has_capability(request.user, "spares.view"),
    ))


@extend_schema(responses=OpenApiTypes.OBJECT)
@api_view(["GET"])
@permission_classes([CanViewDashboard])
def facilities_summary(request):
    """Compact aggregates for the room/rack management tabs."""
    data_center_id = request.query_params.get("data_center", "").strip()
    room_id = request.query_params.get("server_room", "").strip()
    try:
        data_center_pk = int(data_center_id) if data_center_id else None
        room_pk = int(room_id) if room_id else None
    except (TypeError, ValueError):
        return Response({"detail": "数据中心或机房筛选无效"}, status=400)

    if data_center_pk is not None and not DataCenter.objects.filter(
        pk=data_center_pk, is_active=True
    ).exists():
        return Response({"detail": "数据中心筛选无效"}, status=400)
    selected_room = None
    if room_pk is not None:
        selected_room = ServerRoom.objects.select_related("data_center").filter(
            pk=room_pk, is_active=True, data_center__is_active=True
        ).first()
        if selected_room is None:
            return Response({"detail": "机房筛选无效"}, status=400)
        if data_center_pk is not None and selected_room.data_center_id != data_center_pk:
            return Response({"detail": "所选机房不属于当前数据中心"}, status=400)

    # Keep all rooms/racks for status totals, but use only active locations for
    # capacity and device aggregates. This prevents disabled infrastructure
    # from inflating U totals while preserving maintenance-page counts.
    rooms = ServerRoom.objects.filter(data_center__is_active=True).select_related("data_center")
    if data_center_pk is not None:
        rooms = rooms.filter(data_center_id=data_center_pk)
    if room_pk is not None:
        rooms = rooms.filter(pk=room_pk)
    all_racks = Rack.objects.filter(room__in=rooms)
    racks = all_racks.filter(
        is_active=True,
        room__is_active=True,
        room__data_center__is_active=True,
    ).select_related("room__data_center").prefetch_related("allocations")
    room_ids = list(rooms.values_list("id", flat=True))
    rack_rows, totals = build_rack_capacity_rows(
        list(racks.order_by("room__data_center__name", "room__name", "code")),
        include_status=True,
    )
    room_rows = []
    for room in rooms.order_by("data_center__name", "name"):
        room_racks = [
            row for row in rack_rows if row["server_room_id"] == room.id
        ]
        room_rows.append({
            "id": room.id,
            "name": room.name,
            "data_center": room.data_center.name,
            "data_center_id": room.data_center_id,
            "racks_count": len(room_racks),
            "assets_count": sum(row["device_count"] for row in room_racks),
            "total_u": sum(row["total_u"] for row in room_racks),
            "used_u": sum(row["used_u"] for row in room_racks),
            "is_active": room.is_active,
        })

    data_centers = DataCenter.objects.filter(is_active=True).order_by("name")
    if data_center_pk is not None:
        data_centers = data_centers.filter(pk=data_center_pk)
    elif room_pk is not None:
        # A room filter implicitly narrows the data-center summary as well.
        data_centers = data_centers.filter(pk=selected_room.data_center_id)
    data_center_rows = []
    for data_center in data_centers:
        center_rooms = [row for row in room_rows if row["data_center_id"] == data_center.id]
        total_u = sum(row["total_u"] for row in center_rooms)
        used_u = sum(row["used_u"] for row in center_rooms)
        data_center_rows.append({
            "id": data_center.id,
            "name": data_center.name,
            "rooms_count": len(center_rooms),
            "assets_count": sum(row["assets_count"] for row in center_rooms),
            "racks_count": sum(row["racks_count"] for row in center_rooms),
            "total_u": total_u,
            "used_u": used_u,
            "free_u": max(total_u - used_u, 0),
            "utilization": round(used_u / total_u * 100, 1) if total_u else 0,
            "is_active": data_center.is_active,
        })
    return Response({
        "rooms_total": len(room_ids),
        "rooms_in_use": sum(1 for room in room_rows if room["is_active"]),
        "rooms_disabled": ServerRoom.objects.filter(pk__in=room_ids, is_active=False).count(),
        "racks_total": all_racks.count(),
        "racks_in_use": all_racks.filter(
            is_active=True,
            status="in_use",
            room__is_active=True,
            room__data_center__is_active=True,
        ).count(),
        "total_u": totals["total_u"],
        "used_u": totals["used_u"],
        "free_u": max(totals["total_u"] - totals["used_u"], 0),
        "data_centers": data_center_rows,
        "rooms": room_rows,
        "racks": rack_rows,
    })


def _rack_sheet_name(name, used):
    cleaned = re.sub(r"[\\/*?:\[\]]", "_", name or "未命名数据中心")[:31] or "未命名数据中心"
    candidate, suffix = cleaned, 1
    while candidate in used:
        suffix_text = f"_{suffix}"
        candidate = f"{cleaned[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1
    used.add(candidate)
    return candidate


def _rack_prefix(code):
    match = re.search(r"\d", code or "")
    return (code[:match.start()] if match else code or "未分组").rstrip("-_ ") or "未分组"


def _rack_sort_key(rack):
    match = re.search(r"(\d+)(?!.*\d)", rack.code or "")
    return (int(match.group(1)) if match else 0, rack.code or "")


def _rack_layout_export_limit_response(rack_count, total_u):
    if rack_count <= RACK_LAYOUT_EXPORT_MAX_RACKS and total_u <= RACK_LAYOUT_EXPORT_MAX_U_POSITIONS:
        return None

    return Response(
        {
            "code": "rack_layout_export_too_large",
            "detail": (
                f"机柜布局导出范围过大：包含 {rack_count} 个机柜、{total_u} 个 U 位；"
                f"当前同步导出最多支持 {RACK_LAYOUT_EXPORT_MAX_RACKS} 个机柜、"
                f"{RACK_LAYOUT_EXPORT_MAX_U_POSITIONS} 个 U 位，请缩小范围后重试。"
            ),
            "requested": {"racks": rack_count, "u_positions": total_u},
            "limits": {
                "racks": RACK_LAYOUT_EXPORT_MAX_RACKS,
                "u_positions": RACK_LAYOUT_EXPORT_MAX_U_POSITIONS,
            },
        },
        status=400,
    )


@extend_schema(
    parameters=[
        OpenApiParameter(
            name="room__data_center",
            type=OpenApiTypes.INT,
            required=False,
            description="按现有 Rack 列表 API 的数据中心 ID 筛选导出范围。",
        ),
        OpenApiParameter(
            name="room",
            type=OpenApiTypes.INT,
            required=False,
            description="按现有 Rack 列表 API 的机房 ID 筛选导出范围。",
        ),
    ],
    responses=OpenApiTypes.BINARY,
)
@api_view(["GET"])
@permission_classes([CanExportRacks])
def rack_layout_export(request):
    rack_queryset = Rack.objects.select_related("room__data_center").prefetch_related("allocations__asset__device_type").order_by("room__data_center__name", "code")
    spatial_params = QueryDict("", mutable=True)
    for parameter_name in ("room__data_center", "room"):
        values = request.query_params.getlist(parameter_name)
        if values:
            spatial_params.setlist(parameter_name, values)
    rack_queryset = DjangoFilterBackend().filter_queryset(
        SimpleNamespace(query_params=spatial_params),
        rack_queryset,
        RackViewSet(),
    )
    rack_count = rack_queryset.count()
    total_u = rack_queryset.aggregate(total_u=Coalesce(Sum("total_u"), 0))["total_u"] or 0
    limit_response = _rack_layout_export_limit_response(rack_count, total_u)
    if limit_response:
        return limit_response

    racks = list(rack_queryset)
    book = Workbook()
    book.remove(book.active)
    used_sheet_names = set()
    thin = Side(style="thin", color="B8C5D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    status_fills = {
        "in_use": PatternFill("solid", fgColor="C6EFCE"),
        "repair": PatternFill("solid", fgColor="FCE4D6"),
        "idle": PatternFill("solid", fgColor="E4DFEC"),
        "in_stock": PatternFill("solid", fgColor="D9EAF7"),
        "retired": PatternFill("solid", fgColor="E7E6E6"),
    }
    grouped = {}
    for rack in racks:
        grouped.setdefault(rack.room.data_center.name, []).append(rack)

    for data_center_name, data_center_racks in grouped.items():
        sheet = book.create_sheet(_rack_sheet_name(data_center_name, used_sheet_names))
        sheet.sheet_view.showGridLines = False
        sheet["A1"] = f"数据中心：{data_center_name}"
        sheet["A1"].font = Font(size=14, bold=True, color="17365D")
        sheet.freeze_panes = "A5"
        row_cursor = 3
        by_prefix = {}
        for rack in data_center_racks:
            by_prefix.setdefault(_rack_prefix(rack.code), []).append(rack)
        for prefix, prefix_racks in sorted(by_prefix.items(), key=lambda item: item[0]):
            prefix_racks.sort(key=_rack_sort_key)
            rack_width, gap = 5, 2
            group_start = 1
            group_end = group_start + len(prefix_racks) * (rack_width + gap) - gap - 1
            sheet.merge_cells(start_row=row_cursor, start_column=group_start, end_row=row_cursor, end_column=group_end)
            group_cell = sheet.cell(row_cursor, group_start, f"{prefix} 组")
            group_cell.font = Font(bold=True, color="FFFFFF")
            group_cell.fill = PatternFill("solid", fgColor="1F4E78")
            group_cell.alignment = Alignment(horizontal="center")
            row_cursor += 1
            rack_top = row_cursor
            max_u = max((rack.total_u for rack in prefix_racks), default=45)
            for index, rack in enumerate(prefix_racks):
                start_col = group_start + index * (rack_width + gap)
                end_col = start_col + rack_width - 1
                sheet.merge_cells(start_row=rack_top, start_column=start_col, end_row=rack_top, end_column=end_col)
                header = sheet.cell(rack_top, start_col, f"机柜 {rack.code}")
                header.font = Font(bold=True, color="FFFFFF")
                header.fill = PatternFill("solid", fgColor="4472C4")
                header.alignment = Alignment(horizontal="center")
                sheet.merge_cells(start_row=rack_top + 1, start_column=start_col, end_row=rack_top + 1, end_column=end_col)
                effective_used_u = rack_effective_used_u(list(rack.allocations.all()), rack.total_u)
                meta = sheet.cell(rack_top + 1, start_col, f"{rack.room.name} | {rack.total_u} U | 已用 {effective_used_u} U | 可用 {rack.total_u - effective_used_u} U")
                meta.font = Font(size=9, color="44546A")
                meta.alignment = Alignment(horizontal="center")
                for col in (start_col, end_col):
                    sheet.column_dimensions[get_column_letter(col)].width = 7
                for col in range(start_col + 1, end_col):
                    sheet.column_dimensions[get_column_letter(col)].width = 13
                for col in range(start_col, end_col + 1):
                    sheet.cell(rack_top + 2, col).border = border
                sheet.cell(rack_top + 2, start_col, "U").alignment = Alignment(horizontal="center")
                sheet.cell(rack_top + 2, start_col + 1, "设备信息").alignment = Alignment(horizontal="center")
                sheet.merge_cells(start_row=rack_top + 2, start_column=start_col + 1, end_row=rack_top + 2, end_column=end_col - 1)
                sheet.cell(rack_top + 2, end_col, "U").alignment = Alignment(horizontal="center")
                for offset in range(rack.total_u):
                    row = rack_top + 3 + offset
                    u = rack.total_u - offset
                    sheet.cell(row, start_col, u)
                    sheet.cell(row, end_col, u)
                    for col in range(start_col, end_col + 1):
                        cell = sheet.cell(row, col)
                        cell.border = border
                        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                ordered_allocations = sorted(rack.allocations.all(), key=lambda allocation: (allocation.start_u, allocation.end_u))
                for previous, current in zip(ordered_allocations, ordered_allocations[1:]):
                    if current.start_u - previous.end_u - 1 == 1:
                        gap_row = rack_top + 3 + (rack.total_u - (previous.end_u + 1))
                        for col in range(start_col + 1, end_col):
                            gap_cell = sheet.cell(gap_row, col)
                            gap_cell.fill = PatternFill("solid", fgColor="AEB8C8")
                            gap_cell.font = Font(size=9, color="69778C", italic=True)
                for allocation in rack.allocations.all():
                    top_row = rack_top + 3 + (rack.total_u - allocation.end_u)
                    bottom_row = rack_top + 3 + (rack.total_u - allocation.start_u)
                    asset = allocation.asset
                    text = "\n".join(filter(None, [asset.asset_no, asset.name, asset.device_type.name if asset.device_type_id else "", asset.manufacturer_model, f"SN: {asset.serial_number}" if asset.serial_number else ""]))
                    for row in range(top_row, bottom_row + 1):
                        for col in range(start_col + 1, end_col):
                            sheet.cell(row, col).fill = status_fills.get(asset.status, PatternFill("solid", fgColor="D9EAF7"))
                    if top_row != bottom_row:
                        sheet.merge_cells(start_row=top_row, start_column=start_col + 1, end_row=bottom_row, end_column=end_col - 1)
                    cell = sheet.cell(top_row, start_col + 1, text)
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.font = Font(size=9, bold=True)
            row_cursor = rack_top + 3 + max_u + 2
    if not racks:
        book.create_sheet("无机柜数据")["A1"] = "暂无机柜数据"
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="infrix-rack-layout-{system_localdate().isoformat()}.xlsx"'
    book.save(response)
    return response
