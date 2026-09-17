from io import BytesIO

from django.db import IntegrityError, transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from rest_framework.exceptions import ValidationError

from .models import AssetModel, CustomFieldSet, DeviceType, Manufacturer
from .serializers import AssetModelSerializer


ASSET_MODEL_IMPORT_COLUMNS = (
    ("name", "型号名称", True, "同一厂商内唯一。"),
    ("model_number", "型号编号", False, "填写时在同一厂商内唯一。"),
    ("manufacturer", "厂商", True, "必须提前维护且处于启用状态。"),
    ("device_type", "设备类型", True, "必须提前维护且处于启用状态。"),
    ("fieldset", "字段集", False, "留空时继承设备类型的默认字段集。"),
    ("default_warranty_months", "默认保修月数", False, "非负整数。"),
    ("expected_life_months", "预计寿命月数", False, "非负整数。"),
    ("notes", "备注", False, "型号说明。"),
    ("is_active", "状态", False, "可填 true/false、启用/停用，默认启用。"),
)


def build_asset_model_import_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "资产型号导入"
    sheet.append([item[0] for item in ASSET_MODEL_IMPORT_COLUMNS])
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    guide = workbook.create_sheet("填写说明")
    guide.append(["字段", "显示名称", "必填", "填写说明"])
    for key, label, required, description in ASSET_MODEL_IMPORT_COLUMNS:
        guide.append([key, label, "是" if required else "否", description])
    guide.column_dimensions["A"].width = 28
    guide.column_dimensions["B"].width = 24
    guide.column_dimensions["C"].width = 10
    guide.column_dimensions["D"].width = 72
    return workbook


def _named(queryset, value, label):
    value = str(value or "").strip()
    if not value:
        raise ValidationError({label: f"{label}不能为空"})
    matches = list(queryset.filter(name__iexact=value)[:2])
    if not matches:
        raise ValidationError({label: f"未找到{label}“{value}”"})
    if not matches[0].is_active:
        raise ValidationError({label: f"停用的{label}不能用于资产型号"})
    return matches[0]


def _optional_months(value, field):
    text = str(value or "").strip()
    if not text:
        return None
    if not text.isdigit():
        raise ValidationError({field: "必须是非负整数"})
    return int(text)


def _active(value):
    text = str(value or "").strip().lower()
    if not text:
        return True
    if text in {"true", "1", "yes", "启用"}:
        return True
    if text in {"false", "0", "no", "停用"}:
        return False
    raise ValidationError({"is_active": "状态只能填写 true/false 或启用/停用"})


def _payload(row):
    manufacturer = _named(Manufacturer.objects, row.get("manufacturer"), "厂商")
    device_type = _named(DeviceType.objects, row.get("device_type"), "设备类型")
    fieldset_text = str(row.get("fieldset") or "").strip()
    fieldset = _named(CustomFieldSet.objects, fieldset_text, "字段集") if fieldset_text else None
    return {
        "name": str(row.get("name") or "").strip(),
        "model_number": str(row.get("model_number") or "").strip(),
        "manufacturer": manufacturer.pk,
        "device_type": device_type.pk,
        "fieldset": fieldset.pk if fieldset else None,
        "default_warranty_months": _optional_months(row.get("default_warranty_months"), "default_warranty_months"),
        "expected_life_months": _optional_months(row.get("expected_life_months"), "expected_life_months"),
        "notes": str(row.get("notes") or "").strip(),
        "is_active": _active(row.get("is_active")),
    }


def _rows(upload):
    if upload is None:
        raise ValidationError({"file": "请选择导入文件"})
    if not upload.name.lower().endswith(".xlsx"):
        raise ValidationError({"file": "资产型号导入仅支持 .xlsx 文件"})
    workbook = load_workbook(BytesIO(upload.read()), read_only=True, data_only=True)
    sheet = workbook["资产型号导入"] if "资产型号导入" in workbook.sheetnames else workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows, ())]
    expected = {column[0] for column in ASSET_MODEL_IMPORT_COLUMNS}
    if not headers or not expected.issubset(set(headers)):
        raise ValidationError({"file": "模板字段不完整，请重新下载最新模板"})
    parsed = []
    for line, values in enumerate(rows, start=2):
        row = dict(zip(headers, values))
        if not any(value not in (None, "") for value in row.values()):
            continue
        parsed.append((line, row))
    return parsed


def preview_asset_model_import(upload):
    rows = _rows(upload)
    result = []
    seen = set()
    seen_numbers = set()
    for line, row in rows:
        item = {"line": line, "name": str(row.get("name") or "").strip(), "valid": False, "errors": []}
        try:
            payload = _payload(row)
            key = (payload["manufacturer"], payload["name"].casefold())
            if key in seen:
                raise ValidationError({"name": "文件内存在重复的厂商与型号名称"})
            seen.add(key)
            model_number = payload["model_number"].casefold()
            number_key = (payload["manufacturer"], model_number)
            if model_number and number_key in seen_numbers:
                raise ValidationError({"model_number": "文件内存在重复的厂商与型号编号"})
            if model_number:
                seen_numbers.add(number_key)
            serializer = AssetModelSerializer(data=payload)
            serializer.is_valid(raise_exception=True)
            item["payload"] = payload
            item["valid"] = True
        except ValidationError as exc:
            item["errors"] = exc.detail
        result.append(item)
    return {
        "total": len(result),
        "valid": sum(1 for item in result if item["valid"]),
        "invalid": sum(1 for item in result if not item["valid"]),
        "rows": result,
    }


@transaction.atomic
def commit_asset_model_import(upload):
    preview = preview_asset_model_import(upload)
    if preview["invalid"]:
        raise ValidationError({"detail": "导入文件存在异常，请先修正", "preview": preview})
    created = []
    try:
        for item in preview["rows"]:
            serializer = AssetModelSerializer(data=item["payload"])
            serializer.is_valid(raise_exception=True)
            created.append(serializer.save())
    except IntegrityError as exc:
        raise ValidationError({"detail": "导入期间型号目录已发生变化，请重新预览后再试"}) from exc
    return created
