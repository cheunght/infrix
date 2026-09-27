"""Bounded preview/commit support for department hierarchy imports.

Department ``code`` is the only import identity.  Names are validated against
the existing Department model, but are never used to resolve a parent or an
update target.  The hierarchy is checked as a complete proposed graph before
any row is saved.
"""

from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import IntegrityError, transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .audit import write_audit_log
from .imports import ImportFileError, ImportValidationError
from .models import Department
from .operation_limits import MAX_ASSET_MODEL_IMPORT_BYTES, MAX_ASSET_MODEL_IMPORT_ROWS


DEPARTMENT_IMPORT_COLUMNS = (
    ("code", "部门编码", True, "稳定部门标识；创建和更新都按部门编码匹配，请按文本填写以保留前导零。"),
    ("name", "部门名称", True, "创建时必填；更新时留空表示保留原名称。"),
    ("parent_code", "上级部门编码", False, "创建时留空表示根部门；更新时留空表示保留原上级部门。"),
)
DEPARTMENT_IMPORT_FIELD_LABELS = {
    key: label for key, label, _required, _description in DEPARTMENT_IMPORT_COLUMNS
}
DEPARTMENT_IMPORT_FIELDS = frozenset(DEPARTMENT_IMPORT_FIELD_LABELS)
DEPARTMENT_IMPORT_REQUIRED_HEADERS = frozenset({"code", "name"})
DEPARTMENT_IMPORT_MAX_LENGTHS = {
    "code": Department._meta.get_field("code").max_length,
    "name": Department._meta.get_field("name").max_length,
    "parent_code": Department._meta.get_field("code").max_length,
}

MAX_DEPARTMENT_IMPORT_BYTES = MAX_ASSET_MODEL_IMPORT_BYTES
MAX_DEPARTMENT_IMPORT_ROWS = MAX_ASSET_MODEL_IMPORT_ROWS


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, Decimal)):
        return str(value)
    if isinstance(value, float):
        try:
            decimal_value = Decimal(str(value))
        except InvalidOperation:
            return str(value)
        if not decimal_value.is_finite():
            return str(value)
        return format(decimal_value, "f").rstrip("0").rstrip(".") or "0"
    return str(value).strip()


def _is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _identity_key(value):
    # Matching follows the existing serializer's case-insensitive uniqueness
    # checks, while all stored/displayed values retain their original case.
    return str(value or "").strip().lower()


def _append_error(result, field, code, message):
    result["errors"].append({
        "field": field,
        "label": DEPARTMENT_IMPORT_FIELD_LABELS.get(field, field),
        "code": code,
        "message": message,
    })


def _normalize_headers(cells):
    headers = [_cell_text(cell.value).lstrip("\ufeff").strip() for cell in cells]
    while headers and not headers[-1]:
        headers.pop()
    if not headers or any(not header for header in headers):
        raise ImportFileError("部门导入文件第一行必须包含完整的字段名")
    if len(headers) != len(set(headers)):
        raise ImportFileError("部门导入文件包含重复字段名，请保留每个字段一列")
    missing = sorted(DEPARTMENT_IMPORT_REQUIRED_HEADERS.difference(headers))
    if missing:
        labels = "、".join(DEPARTMENT_IMPORT_FIELD_LABELS[field] for field in missing)
        raise ImportFileError(f"部门导入文件缺少必填列：{labels}（{', '.join(missing)}）")
    ignored = [header for header in headers if header not in DEPARTMENT_IMPORT_FIELDS]
    return tuple(headers), tuple(ignored)


def _read_upload(upload):
    if not upload:
        raise ImportFileError("请上传 .xlsx 文件")
    size = int(getattr(upload, "size", 0) or 0)
    if size > MAX_DEPARTMENT_IMPORT_BYTES:
        raise ImportFileError(
            f"部门导入文件不能超过 {MAX_DEPARTMENT_IMPORT_BYTES // (1024 * 1024)} MB"
        )
    filename = str(getattr(upload, "name", "") or "department-import.xlsx")
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix != "xlsx":
        raise ImportFileError("部门导入仅支持 .xlsx 文件")

    upload.seek(0)
    try:
        workbook = load_workbook(
            upload,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:  # openpyxl exposes several parser-specific exceptions
        raise ImportFileError("Excel 文件无法解析，请确认文件为有效的 .xlsx 工作簿") from exc

    try:
        if not workbook.worksheets:
            raise ImportFileError("Excel 文件没有可读取的工作表")
        sheet = next(
            (item for item in workbook.worksheets if item.title == "部门导入"),
            workbook.worksheets[0],
        )
        iterator = sheet.iter_rows()
        header_cells = next(iterator, None)
        if not header_cells:
            raise ImportFileError("部门导入文件第一行必须包含字段名")
        headers, ignored_columns = _normalize_headers(header_cells)
        rows = []
        for line, cells in enumerate(iterator, start=2):
            if not any(not _is_blank(cell.value) for cell in cells):
                continue
            if any(cell.data_type == "f" for cell in cells):
                raise ImportFileError(f"第 {line} 行包含公式，请将公式结果复制为值后再导入")
            if len(cells) > len(headers) and any(
                not _is_blank(cell.value) for cell in cells[len(headers):]
            ):
                raise ImportFileError(f"第 {line} 行包含未命名的多余数据列")
            values = [cell.value for cell in cells[:len(headers)]]
            values += [""] * max(0, len(headers) - len(values))
            rows.append((line, {header: _cell_text(value) for header, value in zip(headers, values)}))
            if len(rows) > MAX_DEPARTMENT_IMPORT_ROWS:
                raise ImportFileError(
                    f"单次部门导入最多支持 {MAX_DEPARTMENT_IMPORT_ROWS} 条业务记录，请拆分文件后重试"
                )
        if not rows:
            raise ImportFileError("部门导入文件中没有部门数据")
        return filename, headers, ignored_columns, tuple(rows)
    finally:
        workbook.close()


def _load_departments(*, lock=False):
    queryset = Department.objects.all().order_by("id")
    if lock:
        queryset = queryset.select_for_update()
    return tuple(queryset)


def _maps(departments):
    by_code = defaultdict(list)
    by_id = {}
    for department in departments:
        by_code[_identity_key(department.code)].append(department)
        by_id[department.pk] = department
    code_by_id = {
        department_id: _identity_key(department.code)
        for department_id, department in by_id.items()
    }
    parent_by_code = {
        _identity_key(department.code): code_by_id.get(department.parent_id)
        for department in departments
        if len(by_code[_identity_key(department.code)]) == 1
    }
    return by_code, by_id, parent_by_code


def _find_cycle_nodes(parent_by_code):
    cycle_nodes = set()
    visited = set()
    for start in parent_by_code:
        if start in visited:
            continue
        path = []
        positions = {}
        current = start
        while current and current in parent_by_code and current not in visited:
            if current in positions:
                cycle_nodes.update(path[positions[current]:])
                break
            positions[current] = len(path)
            path.append(current)
            current = parent_by_code[current]
        visited.update(path)
    return cycle_nodes


def _dependency_order(specs):
    """Return import rows with same-file parents before their children."""
    by_key = {spec["key"]: spec for spec in specs}
    input_order = {spec["key"]: index for index, spec in enumerate(specs)}
    pending = set(by_key)
    ordered = []
    while pending:
        ready = [
            key for key in pending
            if by_key[key]["parent_key"] not in pending
        ]
        if not ready:
            # The complete graph check runs before this function.  Keep this
            # guard so a future caller cannot accidentally save a partial DAG.
            raise IntegrityError("department dependency graph is cyclic")
        for key in sorted(ready, key=input_order.__getitem__):
            ordered.append(by_key[key])
            pending.remove(key)
    return ordered


def _base_row(line, row):
    return {
        "line": line,
        "code": row.get("code", "").strip(),
        "name": row.get("name", "").strip(),
        "parent_code": row.get("parent_code", "").strip(),
        "parent": "",
        "operation": "error",
        "valid": False,
        "changes": [],
        "errors": [],
    }


def _analyze_rows(parsed, departments):
    filename, headers, ignored_columns, parsed_rows = parsed
    by_code, by_id, current_parent_by_code = _maps(departments)
    code_counts = Counter(
        _identity_key(row.get("code"))
        for _line, row in parsed_rows
        if row.get("code", "").strip()
    )
    workbook_code_keys = set(code_counts)
    row_results = []
    specs = []
    result_by_key = {}

    for line, row in parsed_rows:
        result = _base_row(line, row)
        code = result["code"]
        name = result["name"]
        incoming_parent_code = result["parent_code"]
        code_key = _identity_key(code)
        matches = by_code.get(code_key, []) if code_key else []
        existing = matches[0] if len(matches) == 1 else None
        if not code:
            _append_error(result, "code", "required", "部门编码不能为空")
        elif len(code) > DEPARTMENT_IMPORT_MAX_LENGTHS["code"]:
            _append_error(
                result,
                "code",
                "max_length",
                f"部门编码不能超过 {DEPARTMENT_IMPORT_MAX_LENGTHS['code']} 个字符",
            )
        elif code_counts[code_key] > 1:
            _append_error(result, "code", "duplicate_code", "文件内重复的部门编码")
        elif len(matches) > 1:
            _append_error(result, "code", "ambiguous_code", "系统中存在多个相同部门编码，无法安全判断更新对象")

        if not name and existing is not None:
            name = existing.name
        elif not name:
            _append_error(result, "name", "required", "创建新部门时部门名称不能为空")
        elif len(name) > DEPARTMENT_IMPORT_MAX_LENGTHS["name"]:
            _append_error(
                result,
                "name",
                "max_length",
                f"部门名称不能超过 {DEPARTMENT_IMPORT_MAX_LENGTHS['name']} 个字符",
            )
        result["name"] = name

        if len(incoming_parent_code) > DEPARTMENT_IMPORT_MAX_LENGTHS["parent_code"]:
            _append_error(
                result,
                "parent_code",
                "max_length",
                f"上级部门编码不能超过 {DEPARTMENT_IMPORT_MAX_LENGTHS['parent_code']} 个字符",
            )

        if not code_key or code_counts[code_key] > 1 or len(matches) > 1:
            row_results.append(result)
            continue

        parent_key = _identity_key(incoming_parent_code) if incoming_parent_code else None
        if not incoming_parent_code and existing is not None:
            parent_key = current_parent_by_code.get(code_key)
        if parent_key == code_key:
            _append_error(result, "parent_code", "self_parent", "部门不能将自身作为上级部门")
        elif parent_key is not None:
            parent_matches = by_code.get(parent_key, [])
            parent_in_file = parent_key in workbook_code_keys
            if not parent_matches and not parent_in_file:
                _append_error(result, "parent_code", "missing_parent", f"未找到上级部门编码“{incoming_parent_code}”")
            elif len(parent_matches) > 1:
                _append_error(result, "parent_code", "ambiguous_parent", "上级部门编码匹配到多个部门")
            elif code_counts.get(parent_key, 0) > 1:
                _append_error(result, "parent_code", "parent_row_invalid", "上级部门编码在文件内重复，无法建立层级")

        parent = None
        if parent_key is not None:
            parent_matches = by_code.get(parent_key, [])
            if len(parent_matches) == 1:
                parent = parent_matches[0]
                result["parent"] = parent.name
            elif parent_key in result_by_key:
                result["parent"] = result_by_key[parent_key]["name"]
            else:
                result["parent"] = incoming_parent_code
        result["parent_code"] = (
            incoming_parent_code
            if incoming_parent_code
            else (parent.code if parent is not None else "")
        )
        spec = {
            "key": code_key,
            "code": code,
            "name": name,
            "parent_key": parent_key,
            "existing": existing,
            "result": result,
            "input_index": len(row_results),
        }
        specs.append(spec)
        result_by_key[code_key] = result
        row_results.append(result)

    # Resolve same-file parent display names after every row has been seen.
    row_by_key = {spec["key"]: spec for spec in specs}
    for spec in specs:
        parent_key = spec["parent_key"]
        if parent_key in row_by_key:
            spec["result"]["parent"] = row_by_key[parent_key]["name"]
            if not spec["result"]["parent_code"]:
                spec["result"]["parent_code"] = row_by_key[parent_key]["result"]["code"]

    # Enforce the model's global name constraint without ever using name as an
    # identity.  This also catches two new rows choosing the same name.
    existing_names = defaultdict(list)
    for department in departments:
        existing_names[_identity_key(department.name)].append(_identity_key(department.code))
    proposed_names = defaultdict(list)
    for spec in specs:
        name_key = _identity_key(spec["name"])
        proposed_names[name_key].append(spec["key"])
        conflicting_codes = {
            code_key for code_key in existing_names.get(name_key, []) if code_key != spec["key"]
        }
        if conflicting_codes:
            _append_error(spec["result"], "name", "duplicate_name", "部门名称已存在")
    for name_key, code_keys in proposed_names.items():
        if name_key and len(code_keys) > 1:
            for code_key in code_keys:
                _append_error(row_by_key[code_key]["result"], "name", "duplicate_name", "文件内重复的部门名称")

    # Build the complete proposed graph, including untouched database rows.
    proposed_parent_by_code = dict(current_parent_by_code)
    for spec in specs:
        if spec["parent_key"] is not None and spec["parent_key"] not in by_code and spec["parent_key"] not in row_by_key:
            continue
        proposed_parent_by_code[spec["key"]] = spec["parent_key"]
    cycle_nodes = _find_cycle_nodes(proposed_parent_by_code)
    for spec in specs:
        if spec["key"] in cycle_nodes:
            _append_error(spec["result"], "parent_code", "hierarchy_cycle", "导入后的部门层级会形成循环")
    if cycle_nodes and not any(spec["key"] in cycle_nodes for spec in specs):
        for spec in specs:
            _append_error(spec["result"], "parent_code", "hierarchy_cycle", "系统现有部门层级已包含循环，无法安全导入")

    for spec in specs:
        result = spec["result"]
        if result["errors"]:
            continue
        existing = spec["existing"]
        if existing is None:
            result["operation"] = "create"
            result["valid"] = True
        else:
            changes = []
            if existing.name != spec["name"]:
                changes.append("name")
            if current_parent_by_code.get(spec["key"]) != spec["parent_key"]:
                changes.append("parent_code")
            result["changes"] = changes
            result["operation"] = "update" if changes else "unchanged"
            result["valid"] = True

    error_count = sum(1 for result in row_results if result["errors"])
    preview = {
        "filename": filename,
        "total": len(row_results),
        "create": sum(1 for result in row_results if result["operation"] == "create"),
        "update": sum(1 for result in row_results if result["operation"] == "update"),
        "unchanged": sum(1 for result in row_results if result["operation"] == "unchanged"),
        "error": error_count,
        "valid": len(row_results) - error_count,
        "invalid": error_count,
        "ignored_columns": list(ignored_columns),
        "rows": row_results,
    }
    return preview, specs


def _parse(upload):
    return _read_upload(upload)


def preview_department_import(upload):
    parsed = _parse(upload)
    return _analyze_rows(parsed, _load_departments())[0]


def commit_department_import(upload, request):
    parsed = _parse(upload)
    try:
        with transaction.atomic():
            departments = _load_departments(lock=True)
            preview, specs = _analyze_rows(parsed, departments)
            if preview["error"]:
                raise ImportValidationError(preview)

            by_code = {
                _identity_key(department.code): department
                for department in departments
            }
            valid_specs = [spec for spec in specs if spec["result"]["valid"]]
            created = updated = unchanged = 0
            for spec in _dependency_order(valid_specs):
                department = by_code.get(spec["key"])
                parent = by_code.get(spec["parent_key"]) if spec["parent_key"] else None
                if department is None:
                    department = Department(code=spec["code"], name=spec["name"], parent=parent)
                    department.save()
                    created += 1
                elif spec["result"]["operation"] == "update":
                    department.name = spec["name"]
                    department.parent = parent
                    department.save(update_fields=["name", "parent", "updated_at"])
                    updated += 1
                else:
                    unchanged += 1
                by_code[spec["key"]] = department

            write_audit_log(
                request,
                action="import",
                resource_type="department",
                resource_id="bulk",
                extra={
                    "source": "department_import",
                    "filename": parsed[0][:160],
                    "total_count": len(valid_specs),
                    "created_count": created,
                    "updated_count": updated,
                    "unchanged_count": unchanged,
                    "ignored_columns": list(parsed[2])[:20],
                },
            )
            return {
                "created": created,
                "updated": updated,
                "unchanged": unchanged,
                "total": len(valid_specs),
                "errors": [],
            }
    except ImportValidationError:
        raise
    except (DjangoValidationError, IntegrityError) as exc:
        latest, _ = _analyze_rows(parsed, _load_departments())
        raise ImportValidationError(latest, concurrent=True) from exc


def build_department_import_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "部门导入"
    headers = [key for key, _label, _required, _description in DEPARTMENT_IMPORT_COLUMNS]
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    for row in range(2, MAX_DEPARTMENT_IMPORT_ROWS + 2):
        sheet.cell(row, 1).number_format = "@"
        sheet.cell(row, 3).number_format = "@"
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for index in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 24

    guide = workbook.create_sheet("填写说明")
    guide.append(["字段", "显示名称", "创建必填", "更新行为", "填写说明"])
    for key, label, required, description in DEPARTMENT_IMPORT_COLUMNS:
        update_behavior = "空白保留原值" if key != "code" else "只用于匹配，不修改部门编码"
        guide.append([key, label, "是" if required else "否", update_behavior, description])
    for cell in guide[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for column, width in {"A": 22, "B": 18, "C": 14, "D": 24, "E": 72}.items():
        guide.column_dimensions[column].width = width
    guide.freeze_panes = "A2"
    return workbook
