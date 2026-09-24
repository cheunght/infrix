"""Bounded preview/commit support for responsibility-subject imports.

People remain separate from login accounts.  This module only resolves the
stable employee number, validates the existing Person serializer contract, and
coordinates an atomic import transaction.
"""

from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import IntegrityError, transaction
from django.db.models.functions import Lower
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import ValidationError as DRFValidationError

from .audit import write_audit_log
from .imports import ImportFileError, ImportValidationError
from .models import Department, Person
from .operation_limits import MAX_ASSET_MODEL_IMPORT_BYTES, MAX_ASSET_MODEL_IMPORT_ROWS
from .serializers import PersonSerializer


PEOPLE_IMPORT_COLUMNS = (
    ("employee_no", "员工编号", True, "稳定人员标识；创建和更新都必须填写，请按文本填写以保留前导零。"),
    ("name", "姓名", True, "创建新人员时必填；更新时留空表示保留原姓名。"),
    ("department_code", "部门编码", False, "仅匹配已存在的部门编码，不会自动创建部门；更新时留空表示保留原部门。"),
    ("email", "邮箱", False, "邮箱地址；更新时留空表示保留原邮箱。"),
    ("organization", "单位", False, "所属单位；更新时留空表示保留原值。"),
    ("contact", "联系方式", False, "电话或其他联系方式；更新时留空表示保留原值。"),
    ("is_active", "状态", False, "可填 true/false、启用/停用；创建时留空默认启用，更新时留空表示保留原状态。"),
)
PEOPLE_IMPORT_FIELD_LABELS = {key: label for key, label, _required, _description in PEOPLE_IMPORT_COLUMNS}
PEOPLE_IMPORT_FIELDS = frozenset(PEOPLE_IMPORT_FIELD_LABELS)
PEOPLE_IMPORT_REQUIRED_HEADERS = frozenset({"employee_no"})
PEOPLE_UPDATE_FIELDS = ("name", "department", "email", "organization", "contact", "is_active")

MAX_PEOPLE_IMPORT_BYTES = MAX_ASSET_MODEL_IMPORT_BYTES
MAX_PEOPLE_IMPORT_ROWS = MAX_ASSET_MODEL_IMPORT_ROWS


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, Decimal)):
        return str(value)
    if isinstance(value, float):
        try:
            decimal_value = Decimal(str(value))
        except InvalidOperation:
            return str(value)
        if not decimal_value.is_finite():
            return str(value)
        return format(decimal_value, "f").rstrip("0").rstrip(".") or "0"
    return str(value).strip()


def _is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _identity_key(value):
    return str(value or "").strip().lower()


def _normalize_headers(cells):
    headers = [_cell_text(cell.value).lstrip("\ufeff").strip() for cell in cells]
    while headers and not headers[-1]:
        headers.pop()
    if not headers or any(not header for header in headers):
        raise ImportFileError("人员导入文件第一行必须包含完整的字段名")
    if len(headers) != len(set(headers)):
        raise ImportFileError("人员导入文件包含重复字段名，请保留每个字段一列")
    missing = sorted(PEOPLE_IMPORT_REQUIRED_HEADERS.difference(headers))
    if missing:
        labels = "、".join(PEOPLE_IMPORT_FIELD_LABELS[field] for field in missing)
        raise ImportFileError(f"人员导入文件缺少必填列：{labels}（employee_no）")
    ignored = [header for header in headers if header not in PEOPLE_IMPORT_FIELDS]
    return tuple(headers), tuple(ignored)


def _read_upload(upload):
    if not upload:
        raise ImportFileError("请上传 .xlsx 文件")
    size = int(getattr(upload, "size", 0) or 0)
    if size > MAX_PEOPLE_IMPORT_BYTES:
        raise ImportFileError(f"人员导入文件不能超过 {MAX_PEOPLE_IMPORT_BYTES // (1024 * 1024)} MB")
    filename = str(getattr(upload, "name", "") or "people-import.xlsx")
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix != "xlsx":
        raise ImportFileError("人员导入仅支持 .xlsx 文件")

    upload.seek(0)
    try:
        workbook = load_workbook(
            upload,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:  # openpyxl exposes several parser-specific exceptions
        raise ImportFileError("Excel 文件无法解析，请确认文件为有效的 .xlsx 工作簿") from exc

    try:
        if not workbook.worksheets:
            raise ImportFileError("Excel 文件没有可读取的工作表")
        sheet = next(
            (item for item in workbook.worksheets if item.title == "人员导入"),
            workbook.worksheets[0],
        )
        iterator = sheet.iter_rows()
        header_cells = next(iterator, None)
        if not header_cells:
            raise ImportFileError("人员导入文件第一行必须包含字段名")
        headers, ignored_columns = _normalize_headers(header_cells)
        rows = []
        for line, cells in enumerate(iterator, start=2):
            if not any(not _is_blank(cell.value) for cell in cells):
                continue
            if any(cell.data_type == "f" for cell in cells):
                raise ImportFileError(f"第 {line} 行包含公式，请将公式结果复制为值后再导入")
            if len(cells) > len(headers) and any(
                not _is_blank(cell.value) for cell in cells[len(headers):]
            ):
                raise ImportFileError(f"第 {line} 行包含未命名的多余数据列")
            values = [cell.value for cell in cells[:len(headers)]]
            values += [""] * max(0, len(headers) - len(values))
            rows.append((line, {header: _cell_text(value) for header, value in zip(headers, values)}))
            if len(rows) > MAX_PEOPLE_IMPORT_ROWS:
                raise ImportFileError(
                    f"单次人员导入最多支持 {MAX_PEOPLE_IMPORT_ROWS} 条业务记录，请拆分文件后重试"
                )
        if not rows:
            raise ImportFileError("人员导入文件中没有人员数据")
        return filename, headers, ignored_columns, tuple(rows)
    finally:
        workbook.close()


def _error_items(detail):
    if isinstance(detail, DjangoValidationError):
        detail = getattr(detail, "message_dict", None) or detail.messages
    elif isinstance(detail, DRFValidationError):
        detail = detail.detail
    result = []

    def visit(value, field="row"):
        if isinstance(value, dict):
            for key, nested in value.items():
                visit(nested, str(key))
            return
        if isinstance(value, (list, tuple)):
            for nested in value:
                visit(nested, field)
            return
        result.append({
            "field": field,
            "label": PEOPLE_IMPORT_FIELD_LABELS.get(field, field),
            "message": str(value or "数据格式不正确"),
        })

    visit(detail)
    return result or [{"field": "row", "label": "整行", "message": "数据格式不正确"}]


def _resolve_people(employee_numbers, *, lock=False):
    keys = {_identity_key(value) for value in employee_numbers if value}
    if not keys:
        return {}
    queryset = Person.objects.select_related("department")
    if lock:
        queryset = queryset.select_for_update()
    queryset = queryset.annotate(_employee_key=Lower("employee_no")).filter(_employee_key__in=keys)
    matches = defaultdict(list)
    for person in queryset:
        matches[_identity_key(person.employee_no)].append(person)
    return matches


def _resolve_departments(codes, *, lock=False):
    keys = {_identity_key(value) for value in codes if value}
    if not keys:
        return {}
    queryset = Department.objects.all()
    if lock:
        queryset = queryset.select_for_update()
    queryset = queryset.annotate(_department_key=Lower("code")).filter(_department_key__in=keys)
    matches = defaultdict(list)
    for department in queryset:
        matches[_identity_key(department.code)].append(department)
    return matches


def _parse_active(value):
    normalized = str(value or "").strip().lower()
    if normalized in {"true", "1", "yes", "启用", "active"}:
        return True
    if normalized in {"false", "0", "no", "停用", "inactive"}:
        return False
    raise DjangoValidationError({"is_active": "状态只能填写 true/false、启用/停用"})


def _row_payload(row, headers, department, *, existing):
    payload = {}
    for field in ("name", "email", "organization", "contact"):
        if field in headers and row.get(field, "").strip():
            payload[field] = row[field].strip()
    if "department_code" in headers and row.get("department_code", "").strip():
        payload["department"] = department.pk
    if "is_active" in headers and row.get("is_active", "").strip():
        payload["is_active"] = _parse_active(row["is_active"])
    if existing is None:
        payload["employee_no"] = row["employee_no"].strip()
    return payload


def _changed_fields(person, validated_data):
    changes = []
    for field in PEOPLE_UPDATE_FIELDS:
        if field not in validated_data:
            continue
        value = validated_data[field]
        current = person.department_id if field == "department" else getattr(person, field)
        value = value.pk if field == "department" else value
        if value != current:
            changes.append(field)
    return changes


def _base_row(line, row, person=None):
    return {
        "line": line,
        "employee_no": row.get("employee_no", "").strip(),
        "name": row.get("name", "").strip(),
        "department": (
            person.department.name
            if person is not None and person.department_id and person.department is not None
            else row.get("department_code", "").strip()
        ),
        "operation": "error",
        "valid": False,
        "changes": [],
        "errors": [],
    }


def _analyze_rows(parsed, *, lock=False):
    filename, headers, ignored_columns, parsed_rows = parsed
    employee_counts = Counter(
        _identity_key(row.get("employee_no"))
        for _line, row in parsed_rows
        if row.get("employee_no", "").strip()
    )
    people = _resolve_people(
        (row.get("employee_no", "").strip() for _line, row in parsed_rows),
        lock=lock,
    )
    departments = _resolve_departments(
        (row.get("department_code", "").strip() for _line, row in parsed_rows),
        lock=lock,
    )
    results = []
    internal = []
    for line, row in parsed_rows:
        key = _identity_key(row.get("employee_no"))
        matches = people.get(key, []) if key else []
        person = matches[0] if len(matches) == 1 else None
        result = _base_row(line, row, person)
        try:
            if not key:
                raise DjangoValidationError({"employee_no": "员工编号不能为空，且必须作为人员导入的稳定标识"})
            if employee_counts[key] > 1:
                raise DjangoValidationError({"employee_no": "文件内重复的员工编号"})
            if len(matches) > 1:
                raise DjangoValidationError({"employee_no": "系统中存在多个相同员工编号，无法安全判断更新对象"})

            department = None
            department_code = row.get("department_code", "").strip()
            if department_code:
                department_matches = departments.get(_identity_key(department_code), [])
                if not department_matches:
                    raise DjangoValidationError({"department_code": f"未找到部门编码“{department_code}”"})
                if len(department_matches) > 1:
                    raise DjangoValidationError({"department_code": f"部门编码“{department_code}”匹配到多个部门"})
                department = department_matches[0]
                result["department"] = department.name

            payload = _row_payload(row, headers, department, existing=person)
            serializer = PersonSerializer(
                instance=person,
                data=payload,
                partial=person is not None,
            )
            serializer.is_valid(raise_exception=True)
            if person is None:
                result["operation"] = "create"
                result["valid"] = True
                internal.append({"result": result, "serializer": serializer})
                results.append(result)
                continue

            result["changes"] = _changed_fields(person, serializer.validated_data)
            result["operation"] = "update" if result["changes"] else "unchanged"
            result["valid"] = True
            internal.append({"result": result, "serializer": serializer})
        except (DjangoValidationError, DRFValidationError) as exc:
            result["errors"] = _error_items(exc)
        results.append(result)

    valid = sum(1 for result in results if result["valid"])
    error = len(results) - valid
    preview = {
        "filename": filename,
        "total": len(results),
        "create": sum(1 for result in results if result["operation"] == "create"),
        "update": sum(1 for result in results if result["operation"] == "update"),
        "unchanged": sum(1 for result in results if result["operation"] == "unchanged"),
        "error": error,
        "valid": valid,
        "invalid": error,
        "ignored_columns": list(ignored_columns),
        "rows": results,
    }
    return preview, internal


def _parse(upload):
    return _read_upload(upload)


def preview_people_import(upload):
    return _analyze_rows(_parse(upload))[0]


def commit_people_import(upload, request):
    parsed = _parse(upload)
    try:
        with transaction.atomic():
            preview, internal = _analyze_rows(parsed, lock=True)
            if preview["error"]:
                raise ImportValidationError(preview)

            created = updated = unchanged = 0
            for item in internal:
                operation = item["result"]["operation"]
                if operation == "create":
                    item["serializer"].save()
                    created += 1
                elif operation == "update":
                    item["serializer"].save()
                    updated += 1
                else:
                    unchanged += 1

            write_audit_log(
                request,
                action="import",
                resource_type="person",
                resource_id="bulk",
                extra={
                    "source": "people_import",
                    "filename": parsed[0][:160],
                    "total_count": len(internal),
                    "created_count": created,
                    "updated_count": updated,
                    "unchanged_count": unchanged,
                    "ignored_columns": list(parsed[2])[:20],
                },
            )
            return {
                "created": created,
                "updated": updated,
                "unchanged": unchanged,
                "total": len(internal),
                "errors": [],
            }
    except ImportValidationError:
        raise
    except (DjangoValidationError, DRFValidationError, IntegrityError) as exc:
        latest, _ = _analyze_rows(parsed)
        raise ImportValidationError(latest, concurrent=True) from exc


def build_people_import_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "人员导入"
    headers = [key for key, _label, _required, _description in PEOPLE_IMPORT_COLUMNS]
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for index in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 22

    guide = workbook.create_sheet("填写说明")
    guide.append(["字段", "显示名称", "创建必填", "更新行为", "填写说明"])
    for key, label, required, description in PEOPLE_IMPORT_COLUMNS:
        update_behavior = "空白保留原值" if key != "employee_no" else "只用于匹配，不修改员工编号"
        guide.append([key, label, "是" if required else "否", update_behavior, description])
    for cell in guide[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 18
    guide.column_dimensions["C"].width = 14
    guide.column_dimensions["D"].width = 24
    guide.column_dimensions["E"].width = 72
    guide.freeze_panes = "A2"
    return workbook
